"""Regression tests for the bug-hunt findings on the Altium writers (dblib-sync, altium_lib, altium_native, parity).
Run:  python3 -m tests.test_bughunt_expa"""
import copy, json, tempfile
from pathlib import Path

from ecadlib import altium_dblib_sync as S, altium_lib as AL, altium_native as AN
from tests.fixtures_lite import lite
from tests.test_dblib_sync import _actions, _base_items, _item, _models, _plan

ROOT = Path(__file__).parent.parent


def _spec_fp(name="R_0603_1608Metric"):
    return AL.footprint_from_spec(json.loads((ROOT / "ecadlib/package_specs" / f"{name}.json").read_text()))[0]


# ---------------------------------------------------------------- EXP-12: recipe BODY paths stay inside the library
@lite
def test_exp12_recipe_body_path_cannot_leave_the_library_folder(tmp_path):
    src = tmp_path / "src" / "a" / "PcbLib"
    src.mkdir(parents=True)
    AL.write_pcblib(src / "R.PcbLib", [_spec_fp()])
    (tmp_path / "src" / "secret.txt").write_text("key")
    rp = src / "R.PcbLib.ecadbuild"
    evil = "BODY\t69\t100\t0\t{}\t0\t0\t0\t0\t4\t0\t0\t1\t0\t1\t1\t0\t1"
    rp.write_text(rp.read_text().replace("\nEND", "\n" + evil.format("../../secret.txt") + "\n"
                                         + evil.format("ecadlib_3d/../../../x.bat") + "\nEND"))
    dst = tmp_path / "share" / "lib" / "PcbLib"
    dst.mkdir(parents=True)
    (dst / "R.PcbLib").write_bytes((src / "R.PcbLib").read_bytes())
    AN.copy_recipe(src / "R.PcbLib", dst / "R.PcbLib")
    files = sorted(str(p.relative_to(tmp_path)) for p in tmp_path.rglob("*") if p.is_file())
    assert not any(f.endswith(("secret.txt", "x.bat")) and not f.startswith("src") for f in files), files
    assert "secret" not in (dst / "R.PcbLib.ecadbuild").read_text()
    assert AN.recipe_steps(src / "R.PcbLib") == []                     # never names a file outside ecadlib_3d/
    bt = AN.build_text(tmp_path / "src" / "a")
    assert "secret" not in bt and "x.bat" not in bt, bt
    assert AN.safe_step("ecadlib_3d/R_0/R.step") and not AN.safe_step("ecadlib_3d/../R.step")
    assert not AN.safe_step("C:/x.step") and not AN.safe_step("/etc/passwd") and not AN.safe_step("ecadlib_3d\\..\\x")


# ---------------------------------------------------------------- EXP-13: unnumbered pads (NPTH holes) in compare
@lite
def test_exp13_compare_handles_unnumbered_pads(tmp_path):
    fp = _spec_fp()
    for x in (-2.0, 2.0):
        fp["pads"].append(dict(fp["pads"][0], designator="", x=x, y=1.0, w=1.0, h=1.0, shape="circle", layer=AL.MULTI,
                               hole=1.0, plated=False))
    lib = tmp_path / "USB.PcbLib"
    AL.write_pcblib(lib, [fp])
    assert AN.compare_one(lib, lib.with_name(lib.name + AN.EXT)) == []
    assert AN.compare(tmp_path)["checked"] == 1


# ---------------------------------------------------------------- EXP-75: an Altium edit of pad rotation / mask shows
@lite
def test_exp75_pad_rotation_and_mask_edit_is_an_edit(tmp_path):
    lib = tmp_path / "R.PcbLib"
    fp = _spec_fp()
    AL.write_pcblib(lib, [fp])
    rp = lib.with_name(lib.name + AN.EXT)
    AN.record(lib, AN._sha(lib.read_bytes()))
    built = tmp_path / "built.PcbLib"                                    # a rebuild: other bytes, same pads
    built.write_bytes(AL.pcblib_bytes([fp], lib.name, ("02/02/2030", "01:01:01")))
    assert AN.compare_one(built, rp) == []
    edited = [dict(fp, pads=[dict(p, rot=90.0, mask="none") for p in fp["pads"]])]
    lib.write_bytes(AL.pcblib_bytes(edited, lib.name, ("02/02/2030", "01:01:01")))
    rp.write_text(rp.read_text())                                         # the recipe is still ecadlib's
    probs = AN.compare_one(lib, rp)
    assert probs and "rot" in probs[0] and "mask" in probs[0], probs
    assert AN.lib_state(lib) == "edited"
    assert AN.replace_action(b"new bytes", lib) == "edited"


# ---------------------------------------------------------------- dblib-sync
def _res(out, name):
    return next(r for r in out["results"] if r["name"] == name)


@lite
def test_exp1_control_character_in_text_does_not_stop_the_sync(tmp_path):
    _models(tmp_path)
    items = _base_items()
    items[0]["description"] = "Resistor\x0b10k"
    out = S.sync(_plan(tmp_path, items), tmp_path / "lib")
    assert _res(out, "RES-0603-10k")["action"] == "added", out
    row = S.read_tables(tmp_path / "lib/ecadlib.xlsx")["Resistors"][1][0]
    assert row["Description"] == "Resistor 10k", row
    assert "\x0b" not in (tmp_path / "lib/ecadlib.csv").read_text()


@lite
def test_exp2_header_cell_is_text_not_a_formula(tmp_path):
    _models(tmp_path)
    items = _base_items()
    items[0]["parameters"].append({"name": '=HYPERLINK("http://evil.example/x","Open")', "value": "x"})
    S.sync(_plan(tmp_path, items), tmp_path / "lib")
    import openpyxl
    ws = openpyxl.load_workbook(tmp_path / "lib/ecadlib.xlsx")["Resistors"]
    assert all(c.data_type == "s" for c in ws[1] if c.value), [(c.value, c.data_type) for c in ws[1]]


@lite
def test_exp4_sync_keeps_plan_paths_relative(tmp_path):
    _models(tmp_path)
    plan = _plan(tmp_path, _base_items())
    assert S.load_plan(plan)[0]["symbol"]["file"] == "models/SchLib/Resistor.SchLib"
    out = S.sync(plan, tmp_path / "lib")
    assert _res(out, "RES-0603-10k")["action"] == "added", out
    assert str(tmp_path) not in plan.read_text()


@lite
def test_exp5_refs_differing_only_in_case_are_an_error(tmp_path):
    _models(tmp_path)
    items = _base_items()[:1]
    b = _item("RES-B", "Resistors", "Resistor", "R_0603_1608Metric")
    b["footprint"]["ref"] = "r_0603_1608metric"
    items.append(b)
    out = S.sync(_plan(tmp_path, items), tmp_path / "lib")
    r = _res(out, "RES-B")
    assert r["action"] == "error" and "differ only in case" in r["detail"], r
    assert [p.name for p in (tmp_path / "lib/PcbLib").glob("*.PcbLib")] == ["R_0603_1608Metric.PcbLib"]


@lite
def test_exp9_svn_gets_the_recipes_too(tmp_path):
    _models(tmp_path)
    out = S.sync(_plan(tmp_path, _base_items()), tmp_path / "lib", mode="svndblib")
    w = set(out["written"])
    assert "PcbLib/R_0603_1608Metric.PcbLib.ecadbuild" in w and "SchLib/Resistor.SchLib.ecadbuild" in w, w
    assert "PcbLib/R_0603_1608Metric.PcbLib.ecadstate" in w, w
    add = S.svn_commands(out["written"], [], "m")[0]
    assert "PcbLib/R_0603_1608Metric.PcbLib.ecadbuild" in add
    st = json.loads((tmp_path / "lib" / S.STATE).read_text())
    assert not any(f.endswith((AN.EXT, AN.STATE_EXT)) for f in st["files"]), st


@lite
def test_exp15_parameter_names_cannot_break_the_dblib(tmp_path):
    _models(tmp_path)
    items = _base_items()
    items[0]["parameters"] += [{"name": "Power|Voltage", "value": "1"}, {"name": "Temp\nRange", "value": "2"}]
    plan = _plan(tmp_path, items)
    S.sync(plan, tmp_path / "lib")
    ini = (tmp_path / "lib/ecadlib.DbLib").read_text()
    assert "FieldNameOnly=Power; Voltage|" in ini and "FieldNameOnly=Temp Range|" in ini
    out = S.sync(plan, tmp_path / "lib")
    assert out["written"] == [], out["written"]


@lite
def test_exp17_files_from_a_failed_run_are_still_ecadlibs(tmp_path, monkeypatch):
    _models(tmp_path)
    lib = tmp_path / "lib"
    plan = _plan(tmp_path, _base_items())

    def boom(*a, **k):
        raise OSError("disk full")
    monkeypatch.setattr(S, "write_xlsx", boom)
    try:
        S.sync(plan, lib)
    except OSError:
        pass
    monkeypatch.undo()
    assert not (lib / S.STATE).exists() and (lib / "SchLib/Resistor.SchLib").exists()
    S.sync(plan, lib)
    out = S.sync(_plan(tmp_path, [i for i in _base_items() if i["name"] != "RES-0603-10k"]), lib, prune=True)
    assert _res(out, "SchLib/Resistor.SchLib")["action"] == "removed", out["results"]
    assert not (lib / "SchLib/Resistor.SchLib").exists()


def test_exp23_long_category_names_get_their_own_sheet():
    a = {"name": "A", "folder": "Components/Optoelectronics and Light Emitting Diodes"}
    b = {"name": "B", "folder": "Components/Optoelectronics and Light Emitters"}
    c = {"name": "C", "folder": "Components/optoelectronics and light emitters"}
    m = S.sheet_names([a, b, c])
    sheets = {it["name"]: m.get(id(it)) or S.table_of(it) for it in (a, b, c)}
    assert sheets["A"] != sheets["B"] and S.table_key(sheets["B"]) == S.table_key(sheets["C"]), sheets
    assert all(len(v) <= 31 for v in sheets.values()), sheets
    assert S.sheet_name("History") == "History_"


@lite
def test_exp26_name_with_a_path_is_refused(tmp_path):
    _models(tmp_path)
    for bad in ("../Company", "Libs/Company", "..", ""):
        try:
            S.sync(_plan(tmp_path, _base_items()), tmp_path / "lib", name=bad)
        except ValueError as e:
            assert "file name" in str(e)
        else:
            raise AssertionError(bad)
    assert not (tmp_path / "lib").exists()


def test_exp31_canonical_value_wins_over_a_house_named_parameter():
    from ecadlib import rules as RU
    r = RU.resolve(None, ".")
    r = dict(r, parameters=dict(r["parameters"], names={"Manufacturer Part Number": "Mfr PN"}))
    row = RU.house_row({"Part Number": "X", "Manufacturer Part Number": "RC0603-10K", "Mfr PN": "STALE"}, r)
    assert row["Mfr PN"] == "RC0603-10K" and "Manufacturer Part Number" not in row, row
    row = RU.house_row({"Part Number": "X", "Manufacturer Part Number": "", "Mfr PN": "FROM-BOM"}, r)
    assert row["Mfr PN"] == "FROM-BOM", row


@lite
def test_exp37_hand_numbers_dates_and_formulas_survive(tmp_path):
    import datetime, openpyxl
    _models(tmp_path)
    lib = tmp_path / "lib"
    S.sync(_plan(tmp_path, _base_items()), lib)
    wb = openpyxl.load_workbook(lib / "ecadlib.xlsx")
    ws = wb["Resistors"]
    c0 = ws.max_column
    for i, (h, v) in enumerate((("Stock", 1500), ("Unit Price", 0.0042), ("Stock Value", f"=A2*B2"),
                                ("Last Order", datetime.datetime(2026, 9, 1))), 1):
        ws.cell(1, c0 + i, h)
        ws.cell(2, c0 + i, v)
    wb.save(lib / "ecadlib.xlsx")
    items = _base_items()
    items[0]["description"] = "rev B"
    out = S.sync(_plan(tmp_path, items), lib)
    assert _res(out, "RES-0603-10k")["action"] == "updated"
    ws = openpyxl.load_workbook(lib / "ecadlib.xlsx")["Resistors"]
    got = {ws.cell(1, c).value: ws.cell(2, c) for c in range(1, ws.max_column + 1)}
    assert got["Stock"].value == 1500 and got["Unit Price"].value == 0.0042, got
    assert got["Stock Value"].data_type == "f" and got["Stock Value"].value == "=A2*B2"
    assert got["Last Order"].value == datetime.datetime(2026, 9, 1)
    assert got["Description"].value == "rev B" and got["Comment"].data_type == "s"     # "=Value" stays text


@lite
def test_exp38_duplicate_part_name_in_the_plan_is_an_error(tmp_path):
    _models(tmp_path)
    items = _base_items()
    dup = copy.deepcopy(items[0])
    dup.update(folder="Components/Passives", description="other part")
    items.append(dup)
    out = S.sync(_plan(tmp_path, items), tmp_path / "lib")
    acts = [r["action"] for r in out["results"] if r["name"] == "RES-0603-10k"]
    assert acts == ["added", "error"], out["results"]
    tabs = S.read_tables(tmp_path / "lib/ecadlib.xlsx")
    assert "Passives" not in tabs and tabs["Resistors"][1][0]["Description"] == "RES-0603-10k desc"


@lite
def test_exp39_ohm_sign_in_a_ref(tmp_path):
    _models(tmp_path)
    AL.write_schlib(tmp_path / "models/SchLib/Shunt.SchLib", [AL.resistor_symbol(name="Shunt 10mΩ")])
    it = _item("SHUNT", "Resistors", "Shunt", "R_0603_1608Metric")
    it["symbol"]["ref"] = "Shunt 10mΩ"
    it["symbol"]["file"] = "models/SchLib/Shunt.SchLib"
    out = S.sync(_plan(tmp_path, [it]), tmp_path / "lib")
    assert _res(out, "SHUNT")["action"] == "added", out
    row = S.read_tables(tmp_path / "lib/ecadlib.xlsx")["Resistors"][1][0]
    assert row["Library Ref"] == "Shunt 10mOhm" and row["Library Path"] == "SchLib\\Shunt 10mOhm.SchLib", row


@lite
def test_exp48_a_notes_sheet_is_not_a_component_table(tmp_path):
    import openpyxl
    _models(tmp_path)
    lib = tmp_path / "lib"
    S.sync(_plan(tmp_path, _base_items()), lib)
    wb = openpyxl.load_workbook(lib / "ecadlib.xlsx")
    ws = wb.create_sheet("Notes")
    ws.append(["Who", "What"]); ws.append(["Ann", "Check C12"])
    wb.save(lib / "ecadlib.xlsx")
    items = _base_items()
    items[0]["description"] = "rev B"
    out = S.sync(_plan(tmp_path, items), lib)
    assert not any(r["table"] == "Notes" for r in out["results"]), out["results"]
    assert "Notes$" not in (lib / "ecadlib.DbLib").read_text()
    ws = openpyxl.load_workbook(lib / "ecadlib.xlsx")["Notes"]
    assert [[c.value for c in r] for r in ws.iter_rows()] == [["Who", "What"], ["Ann", "Check C12"]]


@lite
def test_exp53_parameters_never_replace_the_key_columns(tmp_path):
    _models(tmp_path)
    items = _base_items()
    items[0]["parameters"] += [{"name": "Part Number", "value": "10K-0603"}, {"name": "Footprint Path", "value": "X.PcbLib"},
                               {"name": "Footprint Ref", "value": "RESC1608"}]
    plan = _plan(tmp_path, items)
    S.sync(plan, tmp_path / "lib")
    out = S.sync(plan, tmp_path / "lib")
    assert _res(out, "RES-0603-10k")["action"] == "up_to_date", out
    rows = S.read_tables(tmp_path / "lib/ecadlib.xlsx")["Resistors"][1]
    assert [(r["Part Number"], r["Footprint Ref"], r["Footprint Path"]) for r in rows] == \
        [("RES-0603-10k", "R_0603_1608Metric", "PcbLib\\R_0603_1608Metric.PcbLib")], rows


def test_exp59_blank_folder_is_misc():
    assert S.table_of({"folder": "Components/  "}) == "Misc"


@lite
def test_exp66_ready_item_without_model_file_is_skipped(tmp_path):
    _models(tmp_path)
    items = _base_items()
    items.append(_item("NOWHERE1", "ICs", None, None))
    items[-1]["symbol"]["ref"] = items[-1]["footprint"]["ref"] = "NOWHERE1"
    out = S.sync(_plan(tmp_path, items), tmp_path / "lib")
    assert _res(out, "NOWHERE1")["action"] == "skipped" and _res(out, "RES-0603-10k")["action"] == "added", out


@lite
def test_exp68_long_value_is_cut_once_and_resync_is_quiet(tmp_path):
    _models(tmp_path)
    items = _base_items()
    items[0]["parameters"].append({"name": "Notes", "value": "x" * 40000})
    plan = _plan(tmp_path, items)
    S.sync(plan, tmp_path / "lib")
    out = S.sync(plan, tmp_path / "lib")
    assert _res(out, "RES-0603-10k")["action"] == "up_to_date" and out["written"] == [], out


@lite
def test_exp69_legacy_workbook_not_merged_is_kept(tmp_path):
    _models(tmp_path)
    lib = tmp_path / "lib"
    plan = _plan(tmp_path, _base_items())
    S.sync(plan, lib, name="Bibliothèque")
    legacy = lib / "Bibliothèque.xlsx"
    legacy.write_bytes((lib / "Bibliotheque.xlsx").read_bytes())       # both there: the ASCII one is what is read
    out = S.sync(plan, lib, name="Bibliothèque")
    assert legacy.exists(), "not merged, so not deleted"
    assert _res(out, legacy.name)["action"] == "not_in_plan" and legacy.name not in out["removed"]


@lite
def test_exp77_a_library_ecadlib_did_not_write_is_not_replaced(tmp_path):
    _models(tmp_path)
    lib = tmp_path / "lib"
    (lib / "PcbLib").mkdir(parents=True)
    company = AL.pcblib_bytes([dict(_spec_fp(), description="COMPANY OWN")], "R_0603_1608Metric.PcbLib")
    (lib / "PcbLib/R_0603_1608Metric.PcbLib").write_bytes(company)
    out = S.sync(_plan(tmp_path, _base_items()), lib)
    r = _res(out, "RES-0603-10k")
    assert r["action"] == "error" and "not written by ecadlib" in r["detail"], r
    assert (lib / "PcbLib/R_0603_1608Metric.PcbLib").read_bytes() == company
    out = S.sync(_plan(tmp_path, _base_items()), lib, overwrite_edits=True)
    assert _res(out, "RES-0603-10k")["action"] == "added", out


@lite
def test_exp78_access_dblib_is_refused(tmp_path):
    _models(tmp_path)
    lib = tmp_path / "lib"
    lib.mkdir()
    ini = ("[OutputDatabaseLinkFile]\r\nVersion=1.1\r\n[DatabaseLinks]\r\nConnectionString=Provider=Microsoft.Jet.OLEDB.4.0;"
           "Data Source=\\\\server\\db\\Company.mdb\r\nLibraryDatabaseType=Microsoft Access\r\n"
           "LibraryDatabasePath=\\\\server\\db\\Company.mdb\r\n[Table1]\r\nTableName=Resistors\r\n")
    (lib / "ecadlib.DbLib").write_bytes(ini.encode())
    try:
        S.sync(_plan(tmp_path, _base_items()), lib)
    except ValueError as e:
        assert "Microsoft Access" in str(e)
    else:
        raise AssertionError("merged into an Access DbLib")
    assert (lib / "ecadlib.DbLib").read_bytes() == ini.encode() and not (lib / "SchLib").exists()


@lite
def test_exp79_duplicate_workbook_rows_are_reported(tmp_path):
    import openpyxl
    _models(tmp_path)
    lib = tmp_path / "lib"
    plan = _plan(tmp_path, _base_items())
    S.sync(plan, lib)
    wb = openpyxl.load_workbook(lib / "ecadlib.xlsx")
    ws = wb["Resistors"]
    ws.append([c.value for c in ws[2]])
    wb.save(lib / "ecadlib.xlsx")
    out = S.sync(plan, lib)
    assert [r["action"] for r in out["results"] if r["name"] == "RES-0603-10k"] == ["up_to_date", "duplicate"], out
    out = S.sync(plan, lib, prune=True)
    assert len(S.read_tables(lib / "ecadlib.xlsx")["Resistors"][1]) == 1


if __name__ == "__main__":
    import sys
    from tests import run_all
    sys.exit(run_all.main([Path(__file__).stem]))
