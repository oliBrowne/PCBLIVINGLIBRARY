"""Bug-hunt regressions (BOM reading, group bomb). Plain functions: python3 -m tests.run_all test_bughunt_bomb"""
import contextlib, io, shutil, tempfile, uuid, zipfile
from pathlib import Path
from unittest import mock

import openpyxl

from ecadlib import bom_inputs as BI, bom_triage as T, reports as R
from ecadlib.kicad_index import KiCadIndex

FIX = Path(__file__).parent / "fixtures"


def _write(text, suffix=".csv", enc="utf-8"):
    f = Path(tempfile.mkdtemp()) / f"bom{suffix}"
    f.write_bytes(text.encode(enc) if isinstance(text, str) else text)
    return f


def _xlsx(build, name="bom.xlsx"):
    wb = openpyxl.Workbook()
    build(wb)
    f = Path(tempfile.mkdtemp()) / name
    wb.save(f)
    return f


def _sym(ref, value, fields=None, unit=1, dnp=False, root="11111111-1111-1111-1111-111111111111"):
    props = {"Reference": ref, "Value": value, "Footprint": "Package_SO:SOIC-8_3.9x4.9mm_P1.27mm", **(fields or {})}
    p = "\n".join(f'    (property "{k}" "{v}" (at 0 0 0))' for k, v in props.items())
    return (f'  (symbol (lib_id "Device:U") (at 0 0 0) (unit {unit}) (in_bom yes) (on_board yes) '
            f'(dnp {"yes" if dnp else "no"})\n    (uuid {uuid.uuid4()})\n{p}\n'
            f'    (instances (project "p" (path "/{root}" (reference "{ref}") (unit {unit}))))\n  )')


def _sch(body, root="11111111-1111-1111-1111-111111111111"):
    return f'(kicad_sch (version 20231120) (generator eeschema) (uuid {root}) (paper "A4") (lib_symbols)\n{body}\n)\n'


def _item(name, fp_ref, value="10k", mpn=""):
    pars = [{"name": "Value", "value": value}] + ([{"name": "Manufacturer Part Number", "value": mpn}] if mpn else [])
    return {"name": name, "parameters": pars, "footprint": {"ref": fp_ref}}


def test_bom59_dnp_line_not_built_from_fitted_line_of_other_package():
    from ecadlib.librarian import _line_for
    for mpn in ("RC0603FR-0710KL", ""):
        lines = [{"designator": "R5", "value": "10k", "footprint": "R_0402_1005Metric", "mpn": mpn, "dnp": True},
                 {"designator": "R1", "value": "10k", "footprint": "R_0603_1608Metric", "mpn": mpn, "dnp": False}]
        rows = [{"status": "dnp (not fitted)", "local_footprint": ""}, {"status": "needs_footprint_generation",
                                                                          "local_footprint": ""}]
        ln, row = _line_for(_item("RES-0402-10k", "R_0402_1005Metric", mpn=mpn), lines, rows)
        assert ln["designator"] == "R5", (mpn, ln)
        ln, _ = _line_for(_item("RES-0603-10k", "R_0603_1608Metric", mpn=mpn), lines, rows)
        assert ln["designator"] == "R1"
    # same footprint on both lines: the fitted one still wins
    lines = [{"designator": "R5", "value": "10k", "footprint": "R_0603_1608Metric", "dnp": True},
             {"designator": "R1", "value": "10k", "footprint": "R_0603_1608Metric", "dnp": False}]
    rows = [{"status": "dnp (not fitted)"}, {"status": "needs_footprint_generation"}]
    assert _line_for(_item("RES-0603-10k", "R_0603_1608Metric"), lines, rows)[0]["designator"] == "R1"


def test_bom30_cell_over_128k_is_read():
    lines, _ = T.read_bom(_write("Designator,MPN,Description\nR1,RC0603FR-0710KL," + "x" * 200_000 + "\n"))
    assert lines[0]["mpn"] == "RC0603FR-0710KL" and len(lines[0]["description"]) == 200_000


def test_bom50_unicode_line_separators_do_not_split_a_row():
    text = ("Designator,Qty,Value,MPN,Description\nR1,1,10k,RC0603FR-0710KL,Thick film resistor\n"
            "R2,1,1k,RC0603FR-071KL,Chip\x0cresistor\nC1,1,100n,GRM155R71C104KA88D,Cap\x85x\n")
    lines, _ = T.read_bom(_write(text))
    assert [l["designator"] for l in lines] == ["R1", "R2", "C1"], lines
    assert lines[0]["description"] == "Thick film resistor" and lines[1]["description"] == "Chip resistor", lines


def test_bom54_fractional_numbers_without_float_noise():
    assert T._plain_number(123456.789) == "123456.789"
    assert T._plain_number(12345678.9) == "12345678.9"
    assert T._plain_number(0.1) == "0.1" and T._plain_number(3.0) == "3" and T._plain_number(1e-07) == "100n"


def test_bom33_quantity_zero_is_not_fitted():
    lines, _ = T.read_bom(_write("Designator,Quantity,Comment,MPN\nR1,1,10k,RC0603FR-0710KL\nR2,0,1k,RC0603FR-071KL\n"))
    assert [bool(l.get("dnp")) for l in lines] == [False, True], lines
    rows = T.triage(lines, KiCadIndex([]), [], None)
    assert rows[1]["status"].startswith("dnp"), rows


def test_bom43_repeated_header_row_is_skipped():
    text = "Designator,Value,MPN\nR1,10k,RC0603FR-0710KL\nDesignator,Value,MPN\nC1,100n,GRM155R71C104KA88D\n"
    lines, _ = T.read_bom(_write(text))
    assert [l["designator"] for l in lines] == ["R1", "C1"], lines


def test_bom46_ragged_rows():
    text = ("Designator,Qty,Value,Footprint,MPN\nR1,1,10k,R_0603_1608Metric,RC0603FR-0710KL\n"
            "R2,1,4,7k,R_0603_1608Metric,RC0603FR-074K7L\nU1,1,LM358,SOIC-8,LM358DR,extra,cells\n")
    lines, _ = T.read_bom(_write(text))
    assert lines[1]["value"] == "4,7k" and lines[1]["footprint"] == "R_0603_1608Metric", lines
    assert lines[1]["mpn"] == "RC0603FR-074K7L" and "decimal comma" in lines[1]["bom_note"], lines
    assert "mpn" not in lines[2] and "footprint" not in lines[2] and "header has 5" in lines[2]["bom_note"], lines
    rows = T.triage(lines, KiCadIndex([]), [], None)
    assert "header has 5" in rows[2]["notes"], rows[2]


def test_bom40_fuzzy_header_mismaps():
    m = T.map_headers(["Designator", "Total Qty", "Qty per board", "MPN"])
    assert m["qty"] == "Qty per board", m
    assert "mpn" not in T.map_headers(["Designator", "Internal PN", "Description", "Qty"])
    assert "footprint" not in T.map_headers(["Designator", "Value", "MPN", "Package Qty"])
    assert "qty" not in T.map_headers(["Designator", "Mount", "Value", "MPN"])
    assert T.map_headers(["Refs", "Qty", "MPN"])["designator"] == "Refs"     # short plurals still match


def test_bom44_unmatched_rules_column_is_reported():
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        lines, m = T.read_bom(_write("Pos,Wert,Herst-Nr\nR1,10k,RC0603FR-0710KL\n"),
                              {"designator": "Pos", "value": "Wert", "mpn": "Hersteller-Nr"})
    assert "mpn = 'Hersteller-Nr' is not a column" in err.getvalue(), err.getvalue()
    assert m["designator"] == "Pos"


def test_bom58_altium_supplier_lcsc_pair():
    text = ("Designator,Quantity,Comment,Manufacturer Part Number 1,Supplier 1,Supplier Part Number 1\n"
            "R1,1,10k,RC0603FR-0710KL,LCSC,C25804\nC1,1,100n,CL05B104KO5NNNC,JLCPCB,C1525\n"
            "U1,1,LM358,LM358DR,Mouser,595-LM358DR\n")
    lines, m = T.read_bom(_write(text))
    assert [l.get("lcsc") for l in lines] == ["C25804", "C1525", None], lines
    assert "lcsc" in m and m["mpn"] == "Manufacturer Part Number 1", m


def test_bom57_mpn_alternates_and_or_equivalent():
    text = ("Designator,Value,Footprint,MPN\nR1,10k,R_0603,RC0603FR-0710KL / ERJ-3EKF1002V\n"
            'R2,10k,R_0603,"RC0603FR-0710KL, ERJ-3EKF1002V"\nR3,10k,R_0603,RC0603FR-0710KL or equivalent\n'
            "R4,10k,R_0603,RC0603FR-0710KL (or eq.)\nU1,LM358,SOIC-8,LM358DR/NOPB\n")
    lines, _ = T.read_bom(_write(text))
    assert [l["mpn"] for l in lines] == ["RC0603FR-0710KL"] * 4 + ["LM358DR/NOPB"], lines
    assert "ERJ-3EKF1002V" in lines[0]["bom_note"] and "equivalent" in lines[2]["bom_note"], lines
    assert "bom_note" not in lines[4]


def test_bom53_designator_on_two_lines_is_flagged():
    lines, _ = T.read_bom(_write("Designator,Value,MPN\nR1,10k,RC0603FR-0710KL\nR2-R3,10k,RC0603FR-0710KL\n"
                                 "R2,1k,RC0603FR-071KL\n"))
    rows = T.triage(lines, KiCadIndex([]), [], None)
    assert "R2 is also on line 3" in rows[1]["notes"] and "R2 is also on line 2" in rows[2]["notes"], rows
    assert "designator" not in rows[0]["notes"]


def test_bom36_kicad_generic_symbol_names_are_not_part_numbers():
    assert T.part_numbers({"value": "Conn_01x04"}) == []
    assert T.part_numbers({"value": "LED_Small"}) == [] and T.part_numbers({"value": "TestPoint"}) == []
    assert T.part_numbers({"value": "STM32F103C8T6"}) == ["STM32F103C8T6"]


def test_bom49_easyeda_skip_is_noted_and_empty_providers_means_none():
    from ecadlib import providers as P
    prov = P.EasyEDA.__new__(P.EasyEDA)
    rows = T.triage([{"designator": "U1", "mpn": "TPS62162DSGR"}], KiCadIndex([]), [prov], None)
    assert "easyeda: skipped (needs an LCSC number)" in rows[0]["notes"], rows
    d = Path(tempfile.mkdtemp())
    bom = _write("Designator,MPN\nU1,TPS62162DSGR\n")
    with contextlib.redirect_stdout(io.StringIO()):
        T.main([str(bom), "--providers", "", "--offline", "-o", str(d / "t"), "--cache", str(d / "c"), "--json"])
    assert "## Providers" in (d / "t.md").read_text() and "| easyeda |" not in (d / "t.md").read_text()


def test_bom38_pipe_in_a_cell_keeps_the_table():
    lines, m = T.read_bom(_write("Designator,Value,Footprint,MPN\nR1,10k,R_0603,A|B|C\n"))
    rows = T.triage(lines, KiCadIndex([]), [], None)
    d = Path(tempfile.mkdtemp())

    class H:
        log = []
    T.write_reports(rows, m, [], H(), d / "t", "x.csv")
    row = next(ln for ln in (d / "t.md").read_text().splitlines() if ln.startswith("| 1 "))
    assert "A\\|B\\|C" in row
    html = R.md_to_html([ln for ln in (d / "t.md").read_text().splitlines()])
    tr = html[html.index("<tr><td>1</td>"):]
    tr = tr[:tr.index("</tr>")]
    assert tr.count("<td>") == 10 and "<td>A|B|C</td>" in tr, tr


def test_bom56_html_xls_is_refused_and_prose_is_not_a_header():
    page = "<html><body><table><tr><td>Designator</td><td>MPN</td></tr><tr><td>R1</td><td>X</td></tr></table></body></html>"
    for f in (_write(page, ".xls"), _write('<?xml version="1.0"?>\n<Workbook xmlns="urn:schemas-microsoft-com:office:'
                                           'spreadsheet"></Workbook>', ".xls")):
        try:
            T.read_bom(f)
            raise AssertionError("read")
        except ValueError as e:
            assert "HTML page or Excel 2003 XML" in str(e), e
    prose = ("This file lists the designator of each part and its MPN (manufacturer part number) for the build.\n"
             "Please check every line before ordering.\n")
    try:
        T.read_bom(_write(prose, ".txt"))
        raise AssertionError("read")
    except ValueError as e:
        assert "no BOM header" in str(e)


def test_bom48_kicad_project_reads_its_schematic():
    d = Path(tempfile.mkdtemp())
    for f in FIX.joinpath("kicad_sch").iterdir():
        (shutil.copytree if f.is_dir() else shutil.copy)(f, d / f.name)
    a = T.read_bom(d / "board.kicad_pro")
    assert a == T.read_bom(d / "board.kicad_sch") and a[0], a
    (d / "board.kicad_sch").unlink()
    try:
        T.read_bom(d / "board.kicad_pro")
        raise AssertionError("read")
    except ValueError as e:
        assert ".kicad_sch" in str(e) and "no BOM header" not in str(e)


def test_bom34_two_row_header_with_merged_group_cells():
    def build(wb):
        ws = wb.active
        ws.append(["Designator", "Qty", "Value", "Manufacturer", None, "Supplier", None])
        ws.append([None, None, None, "Name", "Part Number", "Name", "Part Number"])
        ws.merge_cells("D1:E1")
        ws.merge_cells("F1:G1")
        ws.append(["R1", 1, "10k", "Yageo", "RC0603FR-0710KL", "DigiKey", "311-10.0KHRCT-ND"])
        ws.append(["U1", 1, "LM358", "TI", "LM358DR", "Mouser", "595-LM358DR"])
    lines, m = T.read_bom(_xlsx(build))
    assert [l["mpn"] for l in lines] == ["RC0603FR-0710KL", "LM358DR"], (m, lines)
    assert [l["manufacturer"] for l in lines] == ["Yageo", "TI"], lines


def test_bom32_house_headers_behind_a_cover_sheet():
    def build(wb):
        ws = wb.active
        ws.title = "Deckblatt"
        for r in (["Projekt", "X"], ["Revision", "B"], ["Bearbeiter", "Max"]):
            ws.append(r)
        ws2 = wb.create_sheet("Stueckliste")
        ws2.append(["Pos", "Bauteilwert", "Gehaeuse", "Herst.-Nr.", "Hersteller", "Stk"])
        ws2.append(["R1", "10k", "0603", "RC0603FR-0710KL", "Yageo", 1])
    cols = {"designator": "Pos", "value": "Bauteilwert", "footprint": "Gehaeuse", "mpn": "Herst.-Nr.",
            "manufacturer": "Hersteller", "qty": "Stk"}
    lines, _ = T.read_bom(_xlsx(build), cols)
    assert lines[0]["mpn"] == "RC0603FR-0710KL" and lines[0]["qty"] == "1", lines


def test_bom39_uncached_formula_cells_are_refused():
    def build(wb):
        ws = wb.active
        ws.append(["Designator", "MPN", "Value", "Qty"])
        ws.append(["U1", '=CONCATENATE("LM358","DR")', "LM358", "=1+0"])
    f = _xlsx(build)
    for patch in (contextlib.nullcontext(), mock.patch("pandas.read_excel", side_effect=ImportError("no openpyxl"))):
        with patch:
            try:
                T.read_bom(f)
                raise AssertionError("read")
            except ValueError as e:
                assert "with no saved result" in str(e) and "B2" in str(e), e
    g = f.with_name("saved.xlsx")                       # a saved empty-text result (t="str", <v></v>) is a value
    with zipfile.ZipFile(f) as zi, zipfile.ZipFile(g, "w") as zo:
        for i in zi.infolist():
            d = zi.read(i)
            if i.filename == "xl/worksheets/sheet1.xml":
                d = d.replace(b'<c r="B2"><f>CONCATENATE("LM358","DR")</f><v></v></c>',
                              b'<c r="B2" t="str"><f>CONCATENATE("LM358","DR")</f><v>LM358DR</v></c>')
                d = d.replace(b'<c r="D2"><f>1+0</f><v></v></c>', b'<c r="D2" t="str"><f>""</f><v></v></c>')
            zo.writestr(i, d)
    lines, _ = T.read_bom(g)
    assert lines[0]["mpn"] == "LM358DR", lines


def test_bom41_huge_sparse_sheet_and_zip_bomb():
    def build(wb):
        ws = wb.active
        ws.append(["Designator", "Qty", "Value", "MPN"])
        ws.append(["R1", 1, "10k", "RC0603FR-0710KL"])
        ws["XFD1048576"] = "x"
    f = _xlsx(build)
    lines, _ = T.read_bom(f)                                              # read sparsely, quickly
    assert lines == [{"designator": "R1", "qty": "1", "value": "10k", "mpn": "RC0603FR-0710KL"}], lines
    try:
        BI.xlsx_sheets(f)                                                 # the no-openpyxl reader refuses it
        raise AssertionError("read")
    except ValueError as e:
        assert "is this really a BOM" in str(e)
    with mock.patch.object(T, "MAX_UNZIPPED", 1000):
        for fn in (T.read_bom, BI.xlsx_sheets):
            try:
                fn(f)
                raise AssertionError("read")
            except ValueError as e:
                assert "unpacks to" in str(e), e


def test_bom35_multi_unit_fields_from_any_unit():
    d = Path(tempfile.mkdtemp())
    (d / "mu.kicad_sch").write_text(_sch(_sym("U1", "LM358", {"MPN": "", "Manufacturer": ""}, unit=1)
                                         + _sym("U1", "LM358", {"MPN": "LM358DR", "Manufacturer": "TI"}, unit=2)
                                         + _sym("U2", "LM358", {"MPN": "LM358DR"}, unit=2)
                                         + _sym("U2", "LM358", {"MPN": "LM358DR"}, unit=1, dnp=True)))
    lines, _ = T.read_bom(d / "mu.kicad_sch")
    u1 = next(l for l in lines if l["designator"] == "U1")
    u2 = next(l for l in lines if l["designator"] == "U2")
    assert u1["mpn"] == "LM358DR" and u1["manufacturer"] == "TI" and u1["qty"] == "1", lines
    assert u2["dnp"] is True, lines


def test_bom55_sub_sheet_path_case_insensitive():
    d = Path(tempfile.mkdtemp())
    (d / "Sub").mkdir()
    (d / "Sub" / "Leaf.kicad_sch").write_text(_sch(_sym("R5", "10k"), root="33333333-3333-3333-3333-333333333333"))
    sheet = ('  (sheet (at 0 0) (size 1 1) (uuid a0000000-0000-0000-0000-000000000001)\n'
             '    (property "Sheetname" "A" (at 0 0 0))\n    (property "Sheetfile" "sub/leaf.kicad_sch" (at 0 0 0)))')
    (d / "root.kicad_sch").write_text(_sch(_sym("R1", "1k") + "\n" + sheet))
    lines, _ = T.read_bom(d / "root.kicad_sch")
    assert sorted(l["designator"] for l in lines) == ["R1", "R5"], lines
