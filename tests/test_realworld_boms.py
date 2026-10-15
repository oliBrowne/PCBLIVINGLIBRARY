"""BOMs the way engineers actually export them, through every reader, plus the error contracts the fuzzers rely on.
Offline. Run:  python3 -m tests.test_realworld_boms"""
import contextlib, io, json, tempfile, time
from pathlib import Path

import openpyxl

from ecadlib import altium_lib as AL, bom_triage as T, cli, footprint_gen as FG, kicad_import as KI, lifecycle as L
from ecadlib import mcp_server as M, sexpr, upkeep as UK
from ecadlib.app import Workspace


def _write(name, data, d=None):
    f = Path(d or tempfile.mkdtemp()) / name
    f.write_bytes(data if isinstance(data, bytes) else data.encode())
    return f


def _cli(*argv):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
        rc = cli.main(list(argv))
    return rc, buf.getvalue()


ACTIVEBOM_TOTALS = '''"Bill of Materials","","","",""
"Source Data From:","Board.PrjPcb","","",""
"","","","",""
"Name","Description","Designator","Quantity","Manufacturer Part Number 1"
"STM32F103C8T6","ARM MCU","U1","1","STM32F103C8T6"
"10k","RES 10K 0603","R1, R2","2","RC0603FR-0710KL"
"","","","",""
"","","Total","3",""
'''


def test_totals_row_is_not_a_part():
    lines, _ = T.read_bom(_write("bom.csv", ACTIVEBOM_TOTALS))
    assert [l["designator"] for l in lines] == ["U1", "R1, R2"], lines
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ref Des", "Qty", "Value", "Mfr P/N"])
    ws.append(["R1-R10", 10, "10k", "RC0603FR-0710KL"])
    ws.append([None, "Grand Total:", None, None])
    ws.append(["TOTAL", 10])
    f = Path(tempfile.mkdtemp()) / "bom.xlsx"
    wb.save(f)
    lines, _ = T.read_bom(f)
    assert [l["designator"] for l in lines] == ["R1-R10"], lines


def test_digikey_part_number_is_not_the_mpn():
    text = ("Index,Quantity,Part Number,Manufacturer Part Number,Description,Customer Reference\n"
            "1,10,311-10.0KHRCT-ND,RC0603FR-0710KL,RES 10K 0603,\"R1,R2\"\n")
    lines, m = T.read_bom(_write("digikey.csv", text))
    assert m["mpn"] == "Manufacturer Part Number" and "lcsc" not in m, m
    assert lines[0]["mpn"] == "RC0603FR-0710KL" and "lcsc_invalid" not in lines[0], lines
    lines, m = T.read_bom(_write("plain.csv", "Designator,Part Number\nU1,LM358DR\n"))   # alone it is still the MPN
    assert m["mpn"] == "Part Number" and lines[0]["mpn"] == "LM358DR", m


def test_quoted_cell_spanning_lines():
    text = 'Designator,MPN,Description\nR1,RC0603FR-0710KL,"RES 10k\n0603 1%"\nU1,LM358DR,Op amp\n'
    lines, _ = T.read_bom(_write("bom.csv", text))
    assert lines[0]["description"] == "RES 10k 0603 1%" and lines[1]["mpn"] == "LM358DR", lines


def test_workbook_with_cover_sheet_first():
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.active["A1"] = "Widget rev B"
    ws = wb.create_sheet("BOM")
    ws["A1"] = "Bill of Materials"
    ws.merge_cells("A1:D2")
    ws.append(["Comment", "Designator", "Footprint", "Quantity"])
    ws.append(["STM32F103C8T6", "U1", "LQFP-48", 1])
    f = Path(tempfile.mkdtemp()) / "bom.xlsx"
    wb.save(f)
    lines, m = T.read_bom(f)
    assert m["designator"] == "Designator" and lines[0]["value"] == "STM32F103C8T6", (m, lines)


def test_csv_saved_as_xlsx_and_corrupt_workbooks():
    lines, _ = T.read_bom(_write("bom.xlsx", "Designator,MPN\nU1,LM358DR\n"))        # a CSV with an Excel name
    assert lines[0]["mpn"] == "LM358DR"
    for name, data in (("cut.xlsx", b"PK\x03\x04" + b"\0" * 40), ("old.xls", b"\xd0\xcf\x11\xe0" + b"\0" * 600)):
        try:
            T.read_bom(_write(name, data))
        except ValueError:
            continue
        raise AssertionError(f"{name}: expected ValueError")


def test_en_dash_ranges_count():
    assert T.count_refs("R1–R4") == 4 and T.count_refs("C1 — C3, C7") == 4 and T.count_refs("R5-8") == 4


def test_utf16_upload_through_workspace():
    """Excel 'Unicode Text' (UTF-16 + tabs) reads on the CLI; the web UI / MCP refused it as binary."""
    with tempfile.TemporaryDirectory() as d:
        ws = Workspace(d, [], providers=["digikey"], offline=True)
        t = "Designator\tComment\tQuantity\tManufacturer Part Number\r\nU1\tSTM32F103C8T6\t1\tSTM32F103C8T6\r\n"
        rel = ws.save_upload("bom.txt", t.encode("utf-16"))
        with contextlib.redirect_stdout(io.StringIO()):
            r = ws.triage(rel)
        assert len(r["rows"]) == 1 and r["rows"][0]["mpn"] == "STM32F103C8T6", r["rows"]
        for bad in (Path(d) / "uploads", Path(d) / ("a" * 300 + ".csv")):   # a folder; a name too long for the OS
            try:
                ws._read_bom(bad)                         # a message, not IsADirectoryError / OSError
            except ValueError as e:
                assert "not a file" in str(e)
            else:
                raise AssertionError(f"no error for {bad}")


HOUSE_BOM = "Pos;Bauteil;Sachnummer;Menge\n1;R1;RC0603FR-0710KL;1\n2;U1;STM32F103C8T6;2\n"
HOUSE_RULES = '[bom]\ncolumns = { designator = "Bauteil", mpn = "Sachnummer", qty = "Menge" }\n'


def test_rules_bom_columns_reach_sourcing_upkeep_and_lifecycle():
    """[bom] columns worked for triage/librarian but sourcing, upkeep and the lifecycle reader ignored them."""
    with tempfile.TemporaryDirectory() as d:
        bom = _write("bom.csv", HOUSE_BOM, d)
        _write("ecadlib.toml", HOUSE_RULES, d)
        rc, text = _cli("sourcing", str(bom), "--offline", "--cache", f"{d}/c", "-o", f"{d}/out/s", "--providers", "jlcpcb_search")
        assert "cannot read BOM" not in text and rc in (0, 1, 3), text
        assert "STM32F103C8T6" in Path(f"{d}/out/s.md").read_text()
        cols = {"designator": "Bauteil", "mpn": "Sachnummer", "qty": "Menge"}
        assert [l["mpn"] for l in UK.Source(bom, columns=cols).lines()] == ["RC0603FR-0710KL", "STM32F103C8T6"]
        assert len(L.read_inputs([bom], cols)) == 2
        rules = {"bom": {"columns": cols}, "upkeep": {}}
        with contextlib.redirect_stdout(io.StringIO()):
            rep = UK.run(bom, ["sourcing"], f"{d}/st", f"{d}/out/u", rules=rules, ctx=UK.Ctx(cache=f"{d}/c", offline=True),
                         save=False)
        assert rep["checks"] and all(s["status"] != "error" for s in rep["checks"]), rep["checks"]


def test_unreadable_bom_is_a_message_not_a_traceback():
    with tempfile.TemporaryDirectory() as d:
        bad = _write("bom.csv", "Pos;Bezeichner;Wert\n1;R1;10k\n", d)
        header_only = _write("empty.csv", "Designator,Value,MPN\n", d)
        for argv in (["triage", str(bad), "-o", f"{d}/t", "--offline", "--cache", f"{d}/c"],
                     ["dblib", str(bad), "-o", f"{d}/x"], ["dblib", str(header_only), "-o", f"{d}/y"],
                     ["librarian", str(bad), "-o", f"{d}/l", "--offline", "--cache", f"{d}/c"]):
            rc, text = _cli(*argv)
            assert rc == 2 and "Traceback" not in text and ("BOM" in text or "header" in text), (argv, rc, text)


def test_bad_rules_file_is_a_message_not_a_traceback():
    with tempfile.TemporaryDirectory() as d:
        _write("ecadlib.toml", '[agent.watch]\nsources = ["folder:boms"]\n', d)
        rc, text = _cli("agent", "--workspace", d, "status")
        assert rc == 2 and "agent.watch.sources[0]" in text and "Traceback" not in text, text


FP = Path("tests/fixtures/kicad/USB_C_Receptacle_HRO_TYPE-C-31-M-12.kicad_mod")


def test_malformed_kicad_files_raise_valueerror():
    text = FP.read_text()
    tree = sexpr.parse(text)
    pad = next(sexpr.find(tree, "pad"))
    pad[:] = [x for x in pad if not (isinstance(x, list) and x and x[0] == "size")]     # a pad without (size ...)
    for bad in (sexpr.dumps(tree), "(footprint X (pad 1 smd rect (at)))"):
        try:
            KI.convert_footprint(bad, "X")
        except ValueError:
            continue
        raise AssertionError(f"no ValueError for {bad[:60]!r}")
    sym = sexpr.parse('(symbol "X" (symbol "X_1_1" (pin input line (length 2.54) (name "A") (number "1"))))')
    try:
        KI.convert_symbol([sym], "X")                     # a pin without (at ...)
    except ValueError as e:
        assert "malformed KiCad symbol" in str(e)
    else:
        raise AssertionError("no ValueError for a pin without (at)")


def test_damaged_pcblib_raises_valueerror():
    spec = json.loads(Path("ecadlib/package_specs/R_0603_1608Metric.json").read_text())
    fp, _ = KI.convert_footprint(FG.build(spec)[0], "R0603")
    with tempfile.TemporaryDirectory() as d:
        f = Path(d) / "a.PcbLib"
        AL.write_pcblib(f, [fp])
        import olefile
        with olefile.OleFileIO(str(f)) as o:
            data = o.openstream("Library/Data").read()
        b = bytearray(f.read_bytes())
        i = bytes(b).find(data[:64])
        b[i + 1:i + 5] = b"\x00\x00\xff\x7f"                # first block length points past the stream
        f.write_bytes(bytes(b))
        try:
            AL.read_pcblib(f)
        except ValueError as e:
            assert "damaged" in str(e)
        else:
            raise AssertionError("no ValueError for a damaged PcbLib")


def test_mcp_wrong_id_is_a_user_error():
    with tempfile.TemporaryDirectory() as d:
        ws = Workspace(d, [], offline=True)
        err = io.StringIO()
        with contextlib.redirect_stderr(err):
            r = M.handle(ws, {"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                              "params": {"name": "agent_queue_get", "arguments": {"id": "nope"}}})
        assert r["result"]["isError"] and "Traceback" not in err.getvalue(), (r, err.getvalue())
        assert r["result"]["content"][0]["text"].startswith("no queue item"), r
        with contextlib.redirect_stderr(err):
            r = M.handle(ws, {"jsonrpc": "2.0", "id": 2, "method": "tools/call", "params": {"name": "where_used",
                              "arguments": {"projects": [], "plan": "."}}})   # a folder as the plan
        assert r["result"]["isError"] and "plan not found" in r["result"]["content"][0]["text"], r
        assert "Traceback" not in err.getvalue(), err.getvalue()


def test_spec_validation_is_fast_and_still_strict():
    import jsonschema
    specs = [json.loads(p.read_text()) for p in sorted(Path("ecadlib/package_specs").rglob("*.json"))]
    t0 = time.time()
    for s in specs * 3:
        FG.validate(s)
    assert time.time() - t0 < 3, "validate rebuilt the schema validator on every call"
    try:
        FG.validate({"name": "x"})
    except jsonschema.ValidationError as e:
        assert "required" in e.message
    else:
        raise AssertionError("schema not enforced")


def test_european_decimal_values():
    """Semicolon CSVs from European Excel write 4,7k / 2,2µF: they parsed as 4 ohm and 2 farad."""
    from ecadlib.kicad_index import parse_value
    assert parse_value("4,7k") == 4700 and abs(parse_value("2,2µF") - 2.2e-6) < 1e-15
    assert abs(parse_value("0,1µF") - 1e-7) < 1e-15 and parse_value("1,5 kΩ") == 1500 and parse_value("R1,R2") is None
    lines, _ = T.read_bom(_write("eu.csv", "Designator;Value;Qty\nR1;4,7k;1\nC1;2,2µF;1\n".encode("latin-1")))
    assert [parse_value(l["value"]) for l in lines][0] == 4700 and T.part_numbers(lines[1]) == [], lines


def test_dblib_xlsx_keeps_altium_parameter_references_as_text():
    """parts.xlsx (the Excel DbLib) wrote Comment '=Value' as an Excel formula: Excel shows #NAME?, a reader gets None."""
    with tempfile.TemporaryDirectory() as d:
        bom = _write("bom.csv", "Designator,Qty,Value,Footprint,MPN\nR1,1,10k,R_0603_1608Metric,RC0603FR-0710KL\n"
                                "X1,1,,,=HYPERLINK(\"http://x\")\n", d)
        rc, text = _cli("dblib", str(bom), "-o", f"{d}/out")
        assert rc in (0, None), text
        wb = openpyxl.load_workbook(f"{d}/out/parts.xlsx")
        cells = [c for ws in wb for row in ws.iter_rows() for c in row]
        assert not [c.coordinate for c in cells if c.data_type == "f"], "formulas in parts.xlsx"
        assert "=Value" in [c.value for c in cells]
