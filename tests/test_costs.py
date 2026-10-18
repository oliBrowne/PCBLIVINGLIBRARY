"""Offline tests for ecadlib.costs (the per-model cost sheet) and the built-in prices of non-Claude models.
    python3 -m tests.test_costs"""
import io, tempfile
from pathlib import Path
from unittest import mock

from ecadlib import costs as C, datasheet as D, llm as LL, rules as RU
from ecadlib.llm import profiles as PR, run as LR


def test_builtin_prices_and_tiers_for_2026_models():
    for m, tier, pr in [("gpt-6-astra", "frontier", (10.0, 50.0)), ("gpt-6-luna", "small", (0.10, 0.50)),
                        ("gpt-5.6-terra", "mid", (2.0, 12.0)), ("openai/gpt-5-mini", "mid", (0.25, 2.0)),
                        ("gemini-4-argon", "frontier", (2.0, 10.0)), ("gemini-3.8-flash", "mid", (0.75, 3.75)),
                        ("deepseek-flash", "mid", (0.30, 1.20))]:
        p = PR.profile(m, "openai-compatible")
        assert p["tier"] == tier and (p["price_in"], p["price_out"]) == pr, (m, p)
    assert PR.profile("deepseek-v4-pro", "openai-compatible")["vision"] is False
    assert PR.profile("gpt-4o", "openai-compatible")["price_in"] is None          # not in the table: unknown, as before
    # [ai.profiles] still wins over the table
    r = {"ai": {"profiles": {"gpt-6-luna": {"price_in": 0.2, "price_out": 0.9}}}}
    assert (PR.profile("gpt-6-luna", "openai-compatible", r)["price_in"]) == 0.2


def test_estimate_knows_the_price_of_a_direct_openai_model():
    """`ecadlib datasheet estimate` priced only Claude, OpenRouter and local models: a direct OpenAI / Gemini model
    said 'no price known' (and a max_cost_usd cap refused to send)."""
    r = RU.load({"ai": {"provider": "openai", "model": "gpt-6-luna", "detect": False}})
    LR.info("gpt-6-luna", r)
    assert LL.price("gpt-6-luna") == (0.10, 0.50) and D.price_of("gpt-6-luna") == (0.10, 0.50)


def test_job_costs():
    rin, rout = C.read_tokens()
    assert 20_000 < rin < 35_000 and rout == D.OUT_TOKENS
    opus = C.jobs("claude-opus-5-5")
    assert abs(opus["read"] - (rin * 1.25 * 4 + rout * 20) / 1e6) < 1e-9
    assert abs(opus["read_batch"] - (rin * 4 + rout * 20) / 1e6 / 2) < 1e-9
    assert abs(opus["bom50"] - 10 * opus["read"]) < 1e-9 and abs(opus["lib500"] - 100 * opus["read"]) < 1e-9
    assert 0.15 < opus["check"] < 0.3 and abs(opus["lib500_check_all"] - 500 * opus["check"]) < 1e-9
    gpt = C.jobs("gpt-6-luna", "openai")
    assert gpt["read_batch"] is None and gpt["lib500_batch"] is None           # no Message Batches outside Claude
    assert C.jobs("deepseek-v4-pro", "openai-compatible")["check"] is None      # no image input: no visual check
    local = C.jobs("local", "ollama")
    assert local["read"] == 0 and local["lib500_check_all"] == 0
    unknown = C.jobs("never-seen-model", "openai-compatible")
    assert unknown["read"] is None and unknown["usd_in"] is None


def test_sheet_files_and_cli():
    rows = C.sheet()
    assert len(rows) == len(C.SHEET) and rows[1]["model"].startswith("Claude Opus 5.5")
    assert all(r["read"] is not None for r in rows)                             # every model in the sheet has a price
    with tempfile.TemporaryDirectory() as d:
        md, cs = Path(d) / "costs.md", Path(d) / "costs.csv"
        assert C.main(["--all", "--md", str(md), "--csv", str(cs)]) == 0
        text = md.read_text()
        assert "your own API key" in text and "| Claude Opus 5.5 (default) |" in text and "estimates" in text
        assert cs.read_text().splitlines()[0].startswith("Model,Model id,Tier")
        try:
            import openpyxl
        except ImportError:
            openpyxl = None
        if openpyxl:
            xl = Path(d) / "costs.xlsx"
            assert C.main(["--all", "--xlsx", str(xl)]) == 0
            wb = openpyxl.load_workbook(xl)
            assert wb.sheetnames == ["Cost per job", "How it is estimated"] and wb["Cost per job"].max_row == len(C.SHEET) + 1
    buf = io.StringIO()
    with mock.patch("sys.stdout", buf):
        assert C.main(["--model", "gemini-3.8-flash"]) == 0
    assert "gemini-3.8-flash: $0.75 in" in buf.getvalue() and "500-part library" in buf.getvalue()
    buf = io.StringIO()
    with mock.patch("sys.stdout", buf):
        assert C.main(["--model", "never-seen-model"]) == 1
    assert "no price known" in buf.getvalue()


if __name__ == "__main__":
    import sys
    from tests.run_all import main
    sys.exit(main(["test_costs"]))
