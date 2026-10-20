"""Regression tests for the bug hunt's KiCad-side writer reports (EXP-*): kicad_export, kicad_dblib, parts_export,
sexpr. Plain functions (tests.run_all, no pytest)."""
import io, json, re, sys, tempfile, types
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

from ecadlib import altium_lib as AL, fsutil, kicad_dblib as KD, kicad_export as X, parts_export as PE, sexpr
from tests import test_kicad_export as TX

QUIET = lambda *a, **k: None


def _plan(d):
    """The test_librarian plan (4 ready parts with .kicad_mod, .PcbLib and .SchLib), as a plain folder: no change
    control history (ledger, lock, usage record, stored released files, or the schema stamp that says there was a
    ledger), so a hand edit to the plan in a test is not settled back. TPS1234DSGR (symbol matched with an ordering
    suffix) and EEPART1 (EasyEDA conversion), held for a person to verify, come released anyway (TX.CHECKED)."""
    import shutil
    plan, items, http = TX._plan(d)
    for f in ("library.ledger.jsonl", "library.lock.json", "library.usage.json", "ecadlib_schema.json"):
        (plan.parent / f).unlink(missing_ok=True)
    shutil.rmtree(plan.parent / ".ecadlib", ignore_errors=True)
    ready = {i["name"] for i in json.loads(plan.read_text(encoding="utf-8")) if i["status"] == "ready"}
    assert ready == {"RES-0603-10k", "LED-RED-0603", "TPS1234DSGR", "EEPART1"}, ready   # the parts the tests export
    return plan, items, http


def _edit(plan, fn):
    items = json.loads(plan.read_text(encoding="utf-8"))
    fn({i["name"]: i for i in items})
    plan.write_text(json.dumps(items, indent=1), encoding="utf-8")


def _files(o):
    return {str(f.relative_to(o)): f.read_bytes() for f in sorted(Path(o).rglob("*")) if f.is_file()}


def test_exp11_non_plan_json_leaves_the_library_alone():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        out = Path(d) / "kx"
        X.export(plan, out, "ecadlib", http=http, log=QUIET)
        before = _files(out)
        for bad in ("{}", '{"parts": []}', "[]", "[1, 2]"):
            wrong = plan.parent / "settings.json"
            wrong.write_text(bad, encoding="utf-8")
            err = io.StringIO()
            with redirect_stderr(err), redirect_stdout(io.StringIO()):
                rc = X.main([str(wrong), "-o", str(out), "--offline", "--cache", str(Path(d) / "c")])
            assert rc == 2 and "kicad-export" in err.getvalue(), (bad, rc, err.getvalue())
            assert _files(out) == before, bad                        # nothing deleted, nothing rewritten


def test_exp57_malformed_entry_is_named():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        _edit(plan, lambda it: it["RES-0603-10k"]["parameters"].append({"value": "x"}))
        try:
            X.export(plan, Path(d) / "kx", "ecadlib", http=http, log=QUIET)
            raise AssertionError("a parameter without a name was accepted")
        except X.PlanError as e:
            assert "RES-0603-10k" in str(e) and "has no name" in str(e), e
        _edit(plan, lambda it: (it["RES-0603-10k"]["parameters"].pop(), it["EEPART1"].update(parameters=None)))
        rows = PE.rows(plan, include_held=True)                     # export-parts: "parameters": null is no parameters
        assert any(r["Part ID"] == "EEPART1" for r in rows)
        assert X.plan_fields({"name": "P", "parameters": [{"value": "1"}, {"name": "Value", "value": "2"}]}, "l", None, {})[1] == ("Value", "2")


def test_exp27_two_libraries_in_one_folder_keep_their_records():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        out = Path(d) / "proj"
        X.export(plan, out, "passives", http=http, log=QUIET)
        X.export(plan, out, "ics", http=http, log=QUIET)
        f = out / "passives.pretty" / "R_0603_1608Metric.kicad_mod"
        f.write_text(f.read_text(encoding="utf-8").replace("(layer", "(layer", 1) + "\n; hand edit\n", encoding="utf-8")
        rep = X.export(plan, out, "passives", http=http, log=QUIET)
        assert any(w.startswith("passives.pretty/R_0603_1608Metric.kicad_mod: edited since the last export") for w in rep["warnings"]), rep["warnings"]
        assert "; hand edit" in (out / "ecadlib-hand-edits" / "passives.pretty" / "R_0603_1608Metric.kicad_mod").read_text(encoding="utf-8")
        assert set(json.loads((out / "kicad_export.json").read_text(encoding="utf-8"))["libraries"]) == {"passives", "ics"}
        X.export(plan, out, "ics", http=http, log=QUIET)            # the report's top level is library ics again
        full = json.loads((out / "kicad_export.json").read_text(encoding="utf-8"))
        full["libraries"]["passives"]["footprints"].remove("SOIC-8_3.9x4.9mm_P1.27mm")   # passives' export left it out
        (out / "kicad_export.json").write_text(json.dumps(full), encoding="utf-8")
        r = KD.build(plan, out, "passives")                          # checked against passives' record, not ics'
        assert [w for w in r["warnings"] if "not in passives.pretty" in w] == \
            ["TPS1234DSGR: footprint passives:SOIC-8_3.9x4.9mm_P1.27mm is not in passives.pretty"], r["warnings"]


def test_exp76_existing_library_of_that_name_is_kept_aside():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        out = Path(d) / "proj"
        (out / "mylib.pretty").mkdir(parents=True)
        own_sym = '(kicad_symbol_lib (version 20231120) (generator "kicad") (symbol "OUR_OWN_PART"))\n'
        (out / "mylib.kicad_sym").write_text(own_sym, encoding="utf-8")
        own_fp = '(footprint "R_0603_1608Metric" (layer "F.Cu") (descr "our own"))\n'
        (out / "mylib.pretty" / "R_0603_1608Metric.kicad_mod").write_text(own_fp, encoding="utf-8")
        rep = X.export(plan, out, "mylib", http=http, log=QUIET)
        kept = out / "ecadlib-hand-edits"
        assert (kept / "mylib.kicad_sym").read_text(encoding="utf-8") == own_sym
        assert (kept / "mylib.pretty" / "R_0603_1608Metric.kicad_mod").read_text(encoding="utf-8") == own_fp
        assert sum("not written by ecadlib" in w for w in rep["warnings"]) == 2, rep["warnings"]
        rep = X.export(plan, out, "mylib", http=http, log=QUIET)     # its own files now: no more warnings
        assert not any("not written by ecadlib" in w for w in rep["warnings"]), rep["warnings"]


def test_exp6_every_kept_version_stays():
    with tempfile.TemporaryDirectory() as d:
        out, f = Path(d), Path(d) / "lib.pretty" / "R.kicad_mod"
        f.parent.mkdir()
        for text in ("EDIT ONE", "EDIT TWO", "EDIT TWO"):
            f.write_text(text, encoding="utf-8")
            X._keep_aside(out, f, "lib.pretty/R.kicad_mod")
        kept = sorted(p.read_text(encoding="utf-8") for p in (out / "ecadlib-hand-edits").rglob("*") if p.is_file())
        assert kept == ["EDIT ONE", "EDIT TWO"], kept


def test_exp65_crlf_line_ends_are_no_hand_edit():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        out = Path(d) / "kx"
        X.export(plan, out, "ecadlib", http=http, log=QUIET)
        for f in list(out.glob("*.kicad_sym")) + list(out.glob("ecadlib.pretty/*.kicad_mod")):
            f.write_bytes(f.read_bytes().replace(b"\n", b"\r\n"))   # git core.autocrlf on Windows
        rep = X.export(plan, out, "ecadlib", http=http, log=QUIET)
        assert not any("edited since" in w for w in rep["warnings"]) and not (out / "ecadlib-hand-edits").exists(), rep["warnings"]


def test_exp28_multi_entry_libraries_use_the_plan_ref():
    from ecadlib import symbol_layout as SL
    with tempfile.TemporaryDirectory() as d:
        small = SL.build("Resistor", [{"num": "1", "name": "A", "type": "passive"}, {"num": "2", "name": "B", "type": "passive"}])
        big = SL.build("TPS1234DSGR", [{"num": str(i), "name": f"P{i}", "type": "io"} for i in range(1, 9)])
        AL.write_schlib(Path(d) / "Company.SchLib", [small, big])
        k = X.symbol_from_schlib(Path(d) / "Company.SchLib", "TPS1234DSGR", [], pick="TPS1234DSGR")
        assert X._pins(k) == {str(i) for i in range(1, 9)}
        try:
            X.symbol_from_schlib(Path(d) / "Company.SchLib", "X", [], pick="NOT_THERE")
            raise AssertionError("a ref not in a multi-symbol SchLib fell back to the first symbol")
        except ValueError as e:
            assert "NOT_THERE" in str(e)
        specs = {n: json.loads((Path(X.__file__).parent / "package_specs" / f"{n}.json").read_text()) for n in
                 ("R_0603_1608Metric", "C_0805_2012Metric")}
        fps = [AL.footprint_from_spec(s)[0] for s in specs.values()]
        AL.write_pcblib(Path(d) / "Chips.PcbLib", fps)
        for n in specs:
            want = sexpr.parse(X.footprint_from_pcblib(Path(d) / "Chips.PcbLib", n, [], pick=n))
            alone = Path(d) / f"{n}.PcbLib"
            AL.write_pcblib(alone, [next(f for f in fps if f["name"] == n)])
            assert list(sexpr.find(want, "pad")) == list(sexpr.find(sexpr.parse(X.footprint_from_pcblib(alone, n, [])), "pad")), n


def test_exp28_plan_ref_picks_the_symbol_in_an_export():
    from ecadlib import symbol_layout as SL
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        small = SL.build("Resistor", [{"num": "1", "name": "A", "type": "passive"}, {"num": "2", "name": "B", "type": "passive"}])
        big = SL.build("TPS1234DSGR", [{"num": str(i), "name": f"P{i}", "type": "io"} for i in range(1, 9)])
        AL.write_schlib(plan.parent / "models" / "SchLib" / "Company.SchLib", [small, big])
        _edit(plan, lambda it: it["TPS1234DSGR"].update(symbol={"ref": "TPS1234DSGR", "file": "models/SchLib/Company.SchLib"},
                                                       source=dict(it["TPS1234DSGR"]["source"], symbol=None)))
        out = Path(d) / "kx"
        rep = X.export(plan, out, "ecadlib", http=http, log=QUIET)
        assert X._pins(TX._syms(out / "ecadlib.kicad_sym")["TPS1234DSGR"]) == {str(i) for i in range(1, 9)}, rep["errors"]


def test_exp55_same_footprint_ref_with_other_content_is_an_error():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        _edit(plan, lambda it: it["LED-RED-0603"]["footprint"].update(ref="R_0603_1608Metric"))   # LED's own files
        rep = X.export(plan, Path(d) / "kx", "ecadlib", http=http, log=QUIET)
        assert any("differs from the footprint" in e for e in rep["errors"]), rep["errors"]
        rows = {r["part"]: r["footprint"] for r in rep["parts"]}
        assert sorted(v for v in rows.values() if v == "R_0603_1608Metric") == ["R_0603_1608Metric"], rows


def test_exp14_symbol_pins_without_a_pad_fail_the_export():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        _edit(plan, lambda it: it["TPS1234DSGR"].update(footprint=dict(it["RES-0603-10k"]["footprint"])))
        rep = X.export(plan, Path(d) / "kx", "ecadlib", http=http, log=QUIET)   # the release check, run again, holds it
        assert any(s["part"] == "TPS1234DSGR" and "symbol pins without a pad" in s["why"] for s in rep["skipped"]), rep["skipped"]
        assert "TPS1234DSGR" not in {r["part"] for r in rep["parts"]} and "TPS1234DSGR" not in TX._syms(Path(d) / "kx" / "ecadlib.kicad_sym")
        rep = X.export(plan, Path(d) / "kx2", "ecadlib", include_held=True, http=http, log=QUIET)   # the export's own check
        assert any("TPS1234DSGR: symbol pins without a pad" in e for e in rep["errors"]), rep["errors"]


def test_exp22_exp58_model_files_by_source():
    with tempfile.TemporaryDirectory() as d:
        plan, items, http = _plan(d)
        libs = Path(d) / "mylibs"
        for v, body in (("Vendor_A", "RESISTOR"), ("Vendor_B", "SOIC")):
            (libs / f"{v}.3dshapes").mkdir(parents=True)
            (libs / f"{v}.3dshapes" / "Body.step").write_text(f"ISO-10303-21;\n/* {body} */\nEND-ISO-10303-21;\n")
        (libs / "Vendor_C.3dshapes").mkdir()
        (libs / "Vendor_C.3dshapes" / "LED_TI.stp").write_text("ISO-10303-21;\n/* LED stp */\nEND-ISO-10303-21;\n")
        for part, path in (("RES-0603-10k", "Vendor_A.3dshapes/Body.step"), ("TPS1234DSGR", "Vendor_B.3dshapes/Body.step"),
                           ("LED-RED-0603", "Vendor_C.3dshapes/LED_TI.stp")):
            f = plan.parent / items[part]["footprint"]["kicad_mod"]
            t = re.sub(r'\(model "[^"]*".*?\n\t?\)\n', "", f.read_text(encoding="utf-8"), flags=re.S)
            f.write_text(t.rstrip().rstrip(")") + f'\n\t(model "${{KICAD8_3DMODEL_DIR}}/{path}" (offset (xyz 0 0 0)) '
                         f'(scale (xyz 1 1 1)) (rotate (xyz 0 0 0)))\n)\n', encoding="utf-8")
        out = Path(d) / "kx"
        X.export(plan, out, "ecadlib", libs=[libs], http=http, log=QUIET)
        model = lambda fp: [str(m[1]) for m in sexpr.find(sexpr.parse((out / "ecadlib.pretty" / f"{fp}.kicad_mod").read_text()), "model")]
        r, s = model("R_0603_1608Metric"), model("SOIC-8_3.9x4.9mm_P1.27mm")
        assert len(r) == len(s) == 1 and r != s, (r, s)
        for m, body in ((r[0], "RESISTOR"), (s[0], "SOIC")):
            assert body in (out / "ecadlib.3dshapes" / Path(m).name).read_text(), (m, body)
        assert model("LED_0603_1608Metric") == ["${KIPRJMOD}/ecadlib.3dshapes/LED_TI.stp"]
        assert "LED stp" in (out / "ecadlib.3dshapes" / "LED_TI.stp").read_text()


def test_exp19_no_body_note_when_the_pcblib_model_is_written():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        _edit(plan, lambda it: [i["footprint"].pop("kicad_mod", None) for i in it.values()])
        (plan.parent / "models" / "KiCad").rename(plan.parent / "models" / "KiCad.off")
        rep = X.export(plan, Path(d) / "kx", "ecadlib", http=http, log=QUIET)
        f = rep["footprints"]["R_0603_1608Metric"]
        assert rep["models"].get("R_0603_1608Metric.step") and not any(n.startswith("3D body") for n in f["notes"]), f


def test_exp35_every_3d_body_of_a_footprint():
    with tempfile.TemporaryDirectory() as d:
        sq = [(-1, -1), (1, -1), (1, 1), (-1, 1)]
        fp = {"name": "F", "description": "", "pads": [{"designator": "1", "x": 0, "y": 0, "w": 1, "h": 1, "shape": "rect"}],
              "tracks": [], "fills": [], "regions": [], "texts": [],
              "bodies": [{"pts": sq, "height": 1, "step": {"name": "body.step", "data": b"ISO-10303-21; BODY", "rot": (0, 0, 0)}},
                         {"pts": sq, "height": 1, "step": {"name": "pins.step", "data": b"ISO-10303-21; PINS", "rot": (0, 0, 0)}}]}
        AL.write_pcblib(Path(d) / "F.PcbLib", [fp])
        nodes, files = X.altium_model(Path(d) / "F.PcbLib", "F", "lib", Path(d) / "shapes", {}, [])
        assert len(nodes) == 2 and sorted((Path(d) / "shapes" / f).read_bytes() for f, _ in files) == \
            [b"ISO-10303-21; BODY", b"ISO-10303-21; PINS"], files


def test_exp20_altium_part_zero_is_kicad_unit_zero():
    with tempfile.TemporaryDirectory() as d:
        prims, parts, modes = [], [], []
        for part in (1, 2):
            prims.append(("rect", 0, 0, 40, -40)); parts.append(part); modes.append(0)
            prims.append(("pin", {"num": str(part * 10), "name": "IN", "type": "input"}, "left", 0, -10, 30))
            parts.append(part); modes.append(0)
        prims.append(("pin", {"num": "99", "name": "VCC", "type": "power"}, "top", 20, 0, 30)); parts.append(0); modes.append(0)
        sym = {"name": "DUAL", "description": "", "designator": "U?", "comment": "*", "footprint": None, "prims": prims,
               "prim_parts": parts, "prim_modes": modes, "parts": 2, "modes": 1, "label_at": ((0, 2), (0, -48))}
        (Path(d) / "D.SchLib").write_bytes(AL.schlib_bytes([sym]))
        k = X.symbol_from_schlib(Path(d) / "D.SchLib", "DUAL", [])
        units = {str(u[1]): {str(sexpr.first(p, "number")[1]) for p in sexpr.find(u, "pin")} for u in sexpr.find(k, "symbol")}
        assert units == {"DUAL_0_1": {"99"}, "DUAL_1_1": {"10"}, "DUAL_2_1": {"20"}}, units


def test_exp29_exp40_pcblib_texts_and_rule_layers():
    from ecadlib import kicad_import as K, rules as RU
    kmod = Path(TX.__file__).parent / "fixtures"
    with tempfile.TemporaryDirectory() as d:
        plan, items, _ = _plan(d)
        src = plan.parent / items["TPS1234DSGR"]["footprint"]["kicad_mod"]
        pcb = plan.parent / items["TPS1234DSGR"]["footprint"]["file"]
        t = sexpr.parse(X.footprint_from_pcblib(pcb, "S", []))
        refs = [n for n in sexpr.find(t, "property") if n[1] == "Reference"]
        assert len(refs) == 1 and sexpr.first(refs[0], "layer")[1] == "F.SilkS", refs
        want = [n for n in sexpr.parse(src.read_text()) if isinstance(n, list) and n[:2] in (["fp_text", "reference"], ["property", "Reference"])][0]
        got_at = [round(float(v), 2) for v in sexpr.first(refs[0], "at")[1:3]]
        assert got_at == [round(float(v), 2) for v in sexpr.first(want, "at")[1:3]], (got_at, sexpr.first(want, "at"))
        assert any(n[:3] == ["fp_text", "user", "${REFERENCE}"] and sexpr.first(n, "layer")[1] == "F.Fab"
                   for n in sexpr.find(t, "fp_text")), sexpr.find(t, "fp_text")
        rules = RU.resolve({"layers": {"courtyard": "Mechanical1", "assembly": "Mechanical2"}}, Path(d))
        spec = json.loads((Path(X.__file__).parent / "package_specs" / "SOIC-8_3.9x4.9mm_P1.27mm.json").read_text())
        AL.write_pcblib(Path(d) / "m.PcbLib", [AL.footprint_from_spec(spec)[0]], RU.layers(rules))
        layers = {str(sexpr.first(n, "layer")[1]) for n in sexpr.parse(X.footprint_from_pcblib(Path(d) / "m.PcbLib", "S", [], rules=rules))
                  if isinstance(n, list) and sexpr.first(n, "layer")}
        assert {"F.CrtYd", "F.Fab"} <= layers, layers


def test_exp44_footprint_level_mask_and_paste_margins_reach_every_pad():
    from ecadlib import kicad_import as K
    txt = '''(footprint "QFN_vendor" (version 20240108) (layer "F.Cu")
 (solder_mask_margin 0.07) (solder_paste_margin -0.04)
 (fp_rect (start -3 -2) (end 3 2) (stroke (width 0.05) (type default)) (fill none) (layer "F.CrtYd"))
 (pad "1" smd rect (at -1 0) (size 1 2) (layers "F.Cu" "F.Paste" "F.Mask"))
 (pad "2" smd rect (at 1 0) (size 1 2) (layers "F.Cu" "F.Paste" "F.Mask") (solder_mask_margin 0.02))
)'''
    for conv in (lambda: K.convert_footprint(txt, "Q", "", []), lambda: AL.footprint_from_kicad(txt)):
        fp = conv()
        fp = fp[0] if isinstance(fp, tuple) else fp
        got = {p["designator"]: (p.get("mask"), p.get("paste")) for p in fp["pads"]}
        assert got == {"1": (0.07, -0.04), "2": (0.02, -0.04)}, got
    with tempfile.TemporaryDirectory() as d:
        from ecadlib import parity as PA
        fp = K.convert_footprint(txt, "Q", "", [])
        AL.write_pcblib(Path(d) / "q.PcbLib", [fp[0] if isinstance(fp, tuple) else fp])
        assert not [x for x in PA.footprint_pair(txt, AL.read_pcblib(Path(d) / "q.PcbLib")[0]) if x["level"] == "error"]


def test_exp24_bom_is_ignored():
    assert sexpr.parse("﻿(footprint \"R\" (pad 1))") == ["footprint", "R", ["pad", "1"]]


def test_exp3_exp51_lib_table_rows():
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "sym-lib-table"
        p.write_text('﻿(sym_lib_table\n  (version 7)\n  (lib (name "other")(type "KiCad")(uri "x")(options "")(descr ""))\n'
                     '  (lib\n    (name "mylib")\n    (type "KiCad")\n    (uri "${KIPRJMOD}/old.kicad_sym")\n'
                     '    (options "")\n    (descr "")\n  )\n)\n', encoding="utf-8")
        warn = []
        text = KD.add_lib_row(p, "mylib", "${KIPRJMOD}/mylib.kicad_sym", "d", "KiCad", warnings=warn)
        libs = {str(sexpr.first(l, "name")[1]): str(sexpr.first(l, "uri")[1]) for l in sexpr.find(sexpr.parse(text), "lib")}
        assert libs == {"other": "x", "mylib": "${KIPRJMOD}/mylib.kicad_sym"} and warn, (libs, warn)
        try:
            KD.add_lib_row(p, "mylib", "${KIPRJMOD}/mylib.kicad_dbl", "d", "Database")
            raise AssertionError("a KiCad row was replaced by a database row")
        except ValueError as e:
            assert "pass another --name" in str(e)
        assert "mylib.kicad_sym" in p.read_text(encoding="utf-8")


def test_exp61_project_notice_is_kept():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        out = Path(d) / "proj"
        out.mkdir()
        (out / "NOTICE.txt").write_text("ACME third-party notices\n", encoding="utf-8")
        X.export(plan, out, "ecadlib", http=http, log=QUIET)
        text = (out / "NOTICE.txt").read_text(encoding="utf-8")
        assert text.startswith("ACME third-party notices") and X.KICAD_NOTICE in text, text
        X.export(plan, out, "ecadlib", http=http, log=QUIET)
        assert (out / "NOTICE.txt").read_text(encoding="utf-8") == text


def test_exp43_a_failed_write_leaves_the_old_library():
    with tempfile.TemporaryDirectory() as d:
        plan, _, http = _plan(d)
        out = Path(d) / "kx"
        X.export(plan, out, "ecadlib", http=http, log=QUIET)
        good = (out / "ecadlib.kicad_sym").read_bytes()
        _edit(plan, lambda it: it["RES-0603-10k"].update(description="changed"))
        real = fsutil.replace

        def full(src, dst):
            if str(dst).endswith(".kicad_sym"):
                raise OSError(28, "No space left on device")
            return real(src, dst)
        with mock.patch.object(fsutil, "replace", full):
            try:
                X.export(plan, out, "ecadlib", http=http, log=QUIET)
                raise AssertionError("no error")
            except OSError:
                pass
        assert (out / "ecadlib.kicad_sym").read_bytes() == good
        assert not [f for f in out.iterdir() if f.name.endswith(".tmp")]


def test_exp18_printed_paths_are_quoted():
    assert X._shell_quote("out/kicad") == "out/kicad"
    q = X._shell_quote("My Project/kicad libs")
    assert q in ("'My Project/kicad libs'", '"My Project/kicad libs"'), q


def test_exp50_exp64_plan_fields():
    it = {"name": "P", "parameters": [{"name": "Supplier 1", "value": "Mouser"}, {"name": "Supplier Part Number 1", "value": "A"},
                                      {"name": "Supplier 2", "value": "Mouser"}, {"name": "Supplier Part Number 2", "value": "B"},
                                      {"name": "ComponentLink1Description", "value": "Datasheet"},
                                      {"name": "ComponentLink1URL", "value": "https://x/d.pdf"}]}
    f = dict(X.plan_fields(it, "l", None, {}))
    assert f["Mouser"] == "A" and f["Mouser 2"] == "B" and f["Datasheet"] == "https://x/d.pdf" and "Datasheet URL" not in f, f
    from ecadlib import parity as PA
    assert PA.altium_fields({p["name"]: p["value"] for p in it["parameters"]})["Datasheet"] == "https://x/d.pdf"


def test_exp21_exp41_exp70_parts_table():
    items = [{"name": "U1", "status": "ready", "description": "-40..125C LDO", "symbol": {"ref": "U1", "file": "s"},
              "footprint": {"ref": "SOT", "file": "f"},
              "parameters": [{"name": "Value", "value": "+3.3V"}, {"name": "Lifecycle", "value": "Active"},
                             {"name": "Lifecycle Status", "value": "NRND"}]}]
    rows = PE.rows(items, "My Lib")
    assert rows[0]["KiCad Symbol"] == "My_Lib:U1" and rows[0]["KiCad Footprint"] == "My_Lib:SOT", rows
    assert rows[0]["Lifecycle"] == "Active" and rows[0]["Other Parameters"] == {"Lifecycle Status": "NRND"}, rows
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import load_workbook
        PE.write(rows, Path(d) / "p.xlsx", "xlsx")
        ws = load_workbook(Path(d) / "p.xlsx").active
        cells = {ws.cell(1, c).value: ws.cell(2, c) for c in range(1, ws.max_column + 1)}
        assert cells["Value"].value == "+3.3V" and cells["Description"].value == "-40..125C LDO", cells["Value"].value
        assert cells["Value"].data_type == "s"


def test_exp42_library_name_safe_for_odbc():
    lib = X.libname("R&D;Parts={x}")
    assert not set(";{}=") & set(lib), lib
    s = KD.dbl(lib, f"{lib}.sqlite", {})["source"]["connection_string"]
    assert s.count(";") == 1, s


def test_exp62_dbl_settings_are_kept():
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "l.kicad_dbl"
        tables = {"Resistors": [{"Part ID": "R1", "Symbol": "l:R1", "Footprint": "", "Datasheet": "x"}]}
        old = KD.dbl("l", "l.sqlite", tables)
        old["source"]["timeout_seconds"], old["cache"]["max_age"] = 10, 600
        old["libraries"][0]["fields"][0]["visible_in_chooser"] = True
        p.write_text(json.dumps(old), encoding="utf-8")
        new = KD.keep_settings(p, KD.dbl("l", "l.sqlite", tables))
        assert new["source"]["timeout_seconds"] == 10 and new["cache"]["max_age"] == 600
        assert new["libraries"][0]["fields"][0]["visible_in_chooser"] is True


def test_exp67_sqlite_reserved_table_names():
    assert KD.table_name({"folder": "Components/SQLite_Drivers"}) == "_SQLite_Drivers"
    with tempfile.TemporaryDirectory() as d:
        try:
            KD.write_sqlite(Path(d) / "l.sqlite", {"t": [{"Part ID": "A"}, {"Part ID": "A"}]})   # duplicate key: fails
        except Exception:
            pass
        assert not list(Path(d).iterdir()), list(Path(d).iterdir())


def test_exp32_skips_are_skips_under_pytest():
    from tests import test_kicad_real as KR

    class Skipped(BaseException):
        pass

    def skip(msg):
        raise Skipped(msg)
    with mock.patch.dict(sys.modules, {"pytest": types.SimpleNamespace(skip=skip)}):
        try:
            KR._skip("kicad-cli")
            raise AssertionError("not skipped")
        except Skipped:
            pass
