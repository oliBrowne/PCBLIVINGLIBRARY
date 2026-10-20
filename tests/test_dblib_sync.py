"""Offline tests for ecadlib.altium_dblib_sync (plan.json -> .DbLib / .SVNDbLib folder). Run:
    python3 -m tests.test_dblib_sync

Models are written with altium_lib in a temp dir; svn is never run: a fake runner records the commands."""
import copy, json, tempfile
from pathlib import Path
from types import SimpleNamespace

from ecadlib import altium_dblib_sync as S, altium_export as X, altium_lib as AL, cli
from tests.fixtures_lite import released_anyway

ROOT = Path(__file__).parent.parent


def _models(d):
    (d / "models" / "SchLib").mkdir(parents=True)
    (d / "models" / "PcbLib").mkdir(parents=True)
    AL.write_schlib(d / "models/SchLib/Resistor.SchLib", [AL.resistor_symbol()])
    AL.write_schlib(d / "models/SchLib/Capacitor.SchLib", [AL.capacitor_symbol()])
    for n in ("R_0603_1608Metric", "C_0402_1005Metric"):
        spec = json.loads((ROOT / "ecadlib/package_specs" / f"{n}.json").read_text())
        AL.write_pcblib(d / f"models/PcbLib/{n}.PcbLib", [AL.footprint_from_spec(spec)[0]])


def _item(name, cat, sym, fp, status="ready", **params):
    return {"name": name, "folder": f"Components/{cat}", "description": f"{name} desc", "comment": "=Value",
            "parameters": [{"name": k.replace("_", " "), "value": v} for k, v in params.items()],
            "cleared": ["Supplier 2"], "status": status,
            "symbol": {"ref": sym, "file": f"models/SchLib/{sym}.SchLib" if sym else None},
            "footprint": {"ref": fp, "file": f"models/PcbLib/{fp}.PcbLib" if fp else None},
            "hold_reasons": ["files: no symbol model (.SchLib)"] if status == "held" else []}


def _plan(d, items):
    released_anyway(items, d)                  # the sync checks every part again: these were released anyway
    (d / "plan.json").write_text(json.dumps(items, indent=1))
    return d / "plan.json"


def _base_items():
    return [_item("RES-0603-10k", "Resistors", "Resistor", "R_0603_1608Metric", Value="10k", Manufacturer="Yageo"),
            _item("CAP-0402-100nF", "Capacitors", "Capacitor", "C_0402_1005Metric", Value="100nF"),
            _item("X1", "Crystals", None, "C_0402_1005Metric", status="held"),
            _item("U9", "ICs", "Nope", "Nope", status="needs_altium_models")]


def _actions(out):
    return {r["name"]: r["action"] for r in out["results"]}


def test_fresh_sync_layout_and_readback():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        out = S.sync(_plan(d, _base_items()), d / "lib")
        assert _actions(out) == {"RES-0603-10k": "added", "CAP-0402-100nF": "added", "X1": "held", "U9": "skipped"}, out
        lib = d / "lib"
        for f in ("ecadlib.DbLib", "ecadlib.xlsx", "ecadlib.csv", S.STATE, "SchLib/Resistor.SchLib", "PcbLib/R_0603_1608Metric.PcbLib"):
            assert (lib / f).exists(), f
        # models read back with altium_lib and hold the refs the rows name
        assert [c["name"] for c in AL.read_schlib(lib / "SchLib/Capacitor.SchLib")[1]] == ["Capacitor"]
        assert [f["name"] for f in AL.read_pcblib(lib / "PcbLib/R_0603_1608Metric.PcbLib")] == ["R_0603_1608Metric"]
        tabs = S.read_tables(lib / "ecadlib.xlsx")
        assert set(tabs) == {"Resistors", "Capacitors"}, tabs              # held / skipped parts get no row
        row = tabs["Resistors"][1][0]
        assert row["Library Path"] == "SchLib\\Resistor.SchLib" and row["Footprint Path"] == "PcbLib\\R_0603_1608Metric.PcbLib"
        assert row["Comment"] == "=Value" and row["Value"] == "10k" and row["Manufacturer"] == "Yageo", row
        assert tabs["Resistors"][0][:len(X.COLUMNS)] == X.COLUMNS
        ini = (lib / "ecadlib.DbLib").read_bytes().decode()
        assert ini.startswith("[OutputDatabaseLinkFile]\r\nVersion=1.1\r\n[DatabaseLinks]\r\n") and "\n" not in ini.replace("\r\n", "")
        for k in ("LibraryDatabaseType=Microsoft Excel", "LibraryDatabasePath=ecadlib.xlsx", "DatabasePathRelative=1",
                  "Data Source=C:\\ecadlib\\library\\ecadlib.xlsx", "TableName=Resistors$", "LastFocusedTable=Resistors$",
                  "Options=FieldName=Resistors$.Part Number|TableNameOnly=Resistors$|FieldNameOnly=Part Number|FieldType=0|ParameterName=Part Number|",
                  "FieldNameOnly=Library Path|FieldType=1|ParameterName=[Library Path]|"):
            assert k in ini, k
        assert "SVNRepositoryURL" not in ini
        assert "Crystals" not in ini and "Nope" not in (lib / "ecadlib.csv").read_text()


def test_resync_is_incremental():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        plan = _plan(d, _base_items())
        S.sync(plan, d / "lib")
        lib = d / "lib"
        stamp = {p: p.stat().st_mtime_ns for p in lib.rglob("*") if p.is_file()}
        out = S.sync(plan, lib)
        assert _actions(out)["RES-0603-10k"] == "up_to_date" and out["written"] == [], out
        assert {p: p.stat().st_mtime_ns for p in lib.rglob("*") if p.is_file()} == stamp     # nothing rewritten

        # a value change touches only that row; a hand-added column survives; a cleared field is blanked
        import openpyxl
        wb = openpyxl.load_workbook(lib / "ecadlib.xlsx")
        ws = wb["Resistors"]
        ws.cell(1, ws.max_column + 1, "Stock Location"); ws.cell(2, ws.max_column, "Drawer 4")
        ws.cell(2, X.COLUMNS.index("Supplier 2") + 1, "Mouser")
        wb.save(lib / "ecadlib.xlsx")
        items = _base_items()
        items[0]["parameters"][0]["value"] = "10k 1%"
        out = S.sync(_plan(d, items), lib)
        a = _actions(out)
        assert a["RES-0603-10k"] == "updated" and a["CAP-0402-100nF"] == "up_to_date", out
        res = next(r for r in out["results"] if r["name"] == "RES-0603-10k")
        assert res["detail"] == "row: Value, Supplier 2", res
        row = S.read_tables(lib / "ecadlib.xlsx")["Resistors"][1][0]
        assert row["Value"] == "10k 1%" and row["Stock Location"] == "Drawer 4" and row["Supplier 2"] == "", row
        assert set(out["written"]) == {"ecadlib.xlsx", "ecadlib.csv", "ecadlib.DbLib"}, out["written"]   # new column -> FieldMap
        assert "FieldNameOnly=Stock Location|" in (lib / "ecadlib.DbLib").read_text()

        # a changed model file is copied (and only that one)
        AL.write_schlib(d / "models/SchLib/Resistor.SchLib", [AL.resistor_symbol(description="Resistor, new")])
        out = S.sync(plan, lib)
        assert _actions(out)["RES-0603-10k"] == "updated" and "SchLib/Resistor.SchLib" in out["written"], out
        assert "SchLib/Capacitor.SchLib" not in out["written"]


def test_held_rows_are_left_alone_and_category_moves():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        items = _base_items()
        S.sync(_plan(d, items), d / "lib")
        held = copy.deepcopy(items)
        held[0].update(status="held", hold_reasons=["footprint: courtyard"], description="changed while held")
        out = S.sync(_plan(d, held), d / "lib")
        r = next(r for r in out["results"] if r["name"] == "RES-0603-10k")
        assert r["action"] == "held" and r["detail"] == "footprint: courtyard" and r["table"] == "Resistors", r
        assert S.read_tables(d / "lib/ecadlib.xlsx")["Resistors"][1][0]["Description"] == "RES-0603-10k desc"
        moved = copy.deepcopy(items)
        moved[1]["folder"] = "Components/Passives"
        out = S.sync(_plan(d, moved), d / "lib")
        r = next(r for r in out["results"] if r["name"] == "CAP-0402-100nF")
        assert r["action"] == "updated" and r["detail"].startswith("moved from Capacitors"), r
        tabs = S.read_tables(d / "lib/ecadlib.xlsx")
        assert tabs["Capacitors"][1] == [] and tabs["Passives"][1][0]["Part Number"] == "CAP-0402-100nF"
        assert "TableName=Passives$" in (d / "lib/ecadlib.DbLib").read_text()


def test_removed_items_reported_then_pruned():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        lib = d / "lib"
        S.sync(_plan(d, _base_items()), lib)
        (lib / "PcbLib/Hand_Made.PcbLib").write_bytes((lib / "PcbLib/R_0603_1608Metric.PcbLib").read_bytes())
        only_cap = [i for i in _base_items() if i["name"] != "RES-0603-10k"]
        out = S.sync(_plan(d, only_cap), lib)
        rep = {r["name"]: (r["action"], r.get("kind")) for r in out["results"]}
        assert rep["RES-0603-10k"] == ("not_in_plan", "row") and rep["PcbLib/Hand_Made.PcbLib"] == ("not_in_plan", "file"), rep
        assert "SchLib/Resistor.SchLib" not in rep and (lib / "SchLib/Resistor.SchLib").exists()   # the kept row still uses it
        assert S.read_tables(lib / "ecadlib.xlsx")["Resistors"][1], "kept without --prune"
        out = S.sync(_plan(d, only_cap), lib, prune=True)
        rep = {r["name"]: (r["action"], r["detail"]) for r in out["results"]}
        assert rep["RES-0603-10k"][0] == "removed" and rep["SchLib/Resistor.SchLib"][0] == "removed", rep
        assert rep["PcbLib/Hand_Made.PcbLib"] == ("not_in_plan", "not written by ecadlib: kept"), rep
        assert not (lib / "SchLib/Resistor.SchLib").exists() and (lib / "PcbLib/Hand_Made.PcbLib").exists()
        assert S.read_tables(lib / "ecadlib.xlsx")["Resistors"][1] == []
        assert sorted(out["removed"]) == ["PcbLib/R_0603_1608Metric.PcbLib", "SchLib/Resistor.SchLib"]


def test_bad_model_is_an_error_and_never_reaches_the_table():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        (d / "models/SchLib/Resistor.SchLib").write_bytes(b"not an OLE file")
        out = S.sync(_plan(d, _base_items()), d / "lib")
        r = next(r for r in out["results"] if r["name"] == "RES-0603-10k")   # the release check, run again, holds it
        assert r["action"] == "held" and "does not read back" in r["detail"], r
        assert not (d / "lib/SchLib/Resistor.SchLib").exists()
        assert "Resistors" not in S.read_tables(d / "lib/ecadlib.xlsx")
        AL.write_schlib(d / "models/SchLib/Resistor.SchLib", [AL.resistor_symbol(name="Wrong")])
        r = next(r for r in S.sync(_plan(d, _base_items()), d / "lib")["results"] if r["name"] == "RES-0603-10k")
        assert r["action"] == "error" and "not 'Resistor'" in r["detail"], r


def test_existing_dblib_settings_survive():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        lib = d / "lib"
        S.sync(_plan(d, _base_items()), lib)
        f = lib / "ecadlib.DbLib"
        txt = f.read_bytes().decode().replace("C:\\ecadlib\\library", "E:\\mine").replace("ViewMode=0", "ViewMode=1")
        txt += "[FieldMap999]\r\nOptions=FieldName=Resistors$.Value|TableNameOnly=Resistors$|FieldNameOnly=Value|FieldType=1|ParameterName=Value|VisibleOnAdd=True|AddMode=0|RemoveMode=0|UpdateMode=0\r\n"
        txt = txt.replace("[FieldMap1]\r\n", "[FieldMap1000]\r\n", 1)
        f.write_bytes(txt.encode())
        S.sync(_plan(d, _base_items()), lib)                     # no --win-dir: the user's connection is kept
        new = f.read_bytes().decode()
        assert "E:\\mine\\ecadlib.xlsx" in new and "ViewMode=1" in new and new.count("FieldNameOnly=Value|") == 3, new[:400]
        assert "[FieldMap1000]" in new and "[FieldMap1]\r\n" not in new           # edited sections kept as they were
        S.sync(_plan(d, _base_items()), lib, win_dir=r"D:\lib")
        assert "Data Source=D:\\lib\\ecadlib.xlsx" in f.read_text()


def test_dblib_text_refactor_unchanged():
    t = X.dblib_text("CS", ["A", "B"], excel=True)
    assert t.count("[FieldMap") == 2 * len(X.COLUMNS) and "LastFocusedTable=A$" in t and "LibraryDatabaseType=\r\n" in t
    assert "Options=FieldName=B$.Library Ref|TableNameOnly=B$|FieldNameOnly=Library Ref|FieldType=1|ParameterName=[Library Ref]|" in t
    assert X.dblib_text("CS", {"A": ["Part Number", "X"]}).count("[FieldMap") == 2


class FakeSvn:
    def __init__(self, wc=True, url="https://svn.example.com/repo/trunk/lib", fail=None):
        self.wc, self.url, self.fail, self.calls = wc, url, fail, []

    def __call__(self, cmd, cwd=None):
        self.calls.append((cmd, cwd))
        if cmd[1] == "info":
            ok = self.wc
            return SimpleNamespace(returncode=0 if ok else 1, stdout=(self.url + "\n") if ok else "", stderr="" if ok else "E155007: not a working copy")
        rc = 1 if cmd[1] == self.fail else 0
        return SimpleNamespace(returncode=rc, stdout=f"{cmd[1]} ok\n", stderr="")


def test_svndblib_and_commit_with_fake_svn():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        svn = FakeSvn()
        out = S.sync(_plan(d, _base_items()), d / "wc", mode="svndblib", runner=svn, which=lambda _: "/usr/bin/svn")
        ini = (d / "wc/ecadlib.SVNDbLib").read_bytes().decode()
        assert "SVNRepositoryURL=https://svn.example.com/repo/trunk/lib\r\n" in ini and "FileRepositoryIsRelative=0\r\n" in ini
        assert "LibrarySearchPath=https://svn.example.com/repo/trunk/lib/SchLib;https://svn.example.com/repo/trunk/lib/PcbLib\r\n" in ini
        assert not (d / "wc/ecadlib.DbLib").exists() and out["svn_url"] == svn.url
        s = S.svn_step(d / "wc", out["written"], out["removed"], commit=True, message="sync", runner=svn, which=lambda _: "x")
        assert s["ran"], s
        verbs = [c[1] for c, _ in svn.calls]
        assert verbs == ["info", "info", "add", "commit"], verbs
        add, commit = svn.calls[2][0], svn.calls[3][0]
        assert add[:4] == ["svn", "add", "--force", "--parents"] and "SchLib/Resistor.SchLib" in add and "ecadlib.SVNDbLib" in add
        assert commit[:6] == ["svn", "commit", "--depth", "empty", "-m", "sync"] and {"SchLib", "PcbLib"} <= set(commit)
        assert all(cwd == str(d / "wc") for _, cwd in svn.calls[2:])
        # prune -> svn delete
        S.sync(_plan(d, _base_items()[1:]), d / "wc", mode="svndblib", runner=svn, which=lambda _: "x")
        out = S.sync(_plan(d, _base_items()[1:]), d / "wc", mode="svndblib", prune=True, runner=svn, which=lambda _: "x")
        cmds = S.svn_commands(out["written"], out["removed"], "m")
        assert ["svn", "delete", "--force", "PcbLib/R_0603_1608Metric.PcbLib", "SchLib/Resistor.SchLib"] in cmds, cmds


def test_svn_without_svn_or_working_copy_prints_commands():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        out = S.sync(_plan(d, _base_items()), d / "wc", mode="svndblib", which=lambda _: None)
        assert out["svn_url"] is None and "no SVN repository URL" in out["warning"]
        assert "SVNRepositoryURL=\r\n" in (d / "wc/ecadlib.SVNDbLib").read_bytes().decode()
        s = S.svn_step(d / "wc", out["written"], [], commit=True, which=lambda _: None)
        assert not s["ran"] and "not on PATH" in s["note"] and s["commands"][0].startswith("svn add --force --parents ")
        svn = FakeSvn(wc=False)
        s = S.svn_step(d / "wc", out["written"], [], commit=True, runner=svn, which=lambda _: "x")
        assert not s["ran"] and "not an svn working copy" in s["note"] and [c[1] for c, _ in svn.calls] == ["info"]
        svn = FakeSvn(fail="commit")
        s = S.svn_step(d / "wc", out["written"], [], commit=True, runner=svn, which=lambda _: "x")
        assert not s["ran"] and s["error"] == "commit failed (exit 1)"
        assert S.svn_step(d / "wc", [], [], commit=True)["note"] == "nothing to commit"
        # an svn URL given later is written into the existing file
        S.sync(_plan(d, _base_items()), d / "wc", mode="svndblib", svn_url="file:///r/lib", which=lambda _: None)
        assert "SVNRepositoryURL=file:///r/lib\r\n" in (d / "wc/ecadlib.SVNDbLib").read_bytes().decode()


def test_cli_dry_run_and_registration():
    import contextlib, io
    assert cli.COMMANDS["dblib-sync"][0] == "ecadlib.altium_dblib_sync"
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            S.main([str(_plan(d, _base_items())), "--to", str(d / "lib"), "--dry-run", "-o", str(d / "r.json")])
        assert not (d / "lib").exists(), "dry run writes nothing"
        assert "would write SchLib/Resistor.SchLib" in buf.getvalue() and "2 added, 1 held, 1 skipped" in buf.getvalue()
        rep = json.loads((d / "r.json").read_text())
        assert rep["dry_run"] and len(rep["results"]) == 4



def test_blank_header_column_keeps_hand_columns_aligned():
    """A blank header cell between columns must not shift the hand-added columns right of it onto the wrong data."""
    import openpyxl
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        plan = _plan(d, _base_items())
        S.sync(plan, d / "lib")
        x = d / "lib/ecadlib.xlsx"
        wb = openpyxl.load_workbook(x); ws = wb["Resistors"]
        n = ws.max_column
        ws.cell(1, n + 2, "Stock"); ws.cell(2, n + 2, "500")          # column n+1 left without a header
        ws.append(["HAND-1"] + [""] * n + ["42"])
        wb.save(x)
        S.sync(plan, d / "lib")
        rows = {r["Part Number"]: r for r in S.read_tables(x)["Resistors"][1]}
        assert rows["RES-0603-10k"]["Stock"] == "500" and rows["HAND-1"]["Stock"] == "42", rows


def test_shared_model_failing_readback_fails_every_item_using_it():
    """Two items sharing a symbol whose file does not hold the ref: both are errors, no row points at a missing file."""
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        a = _base_items()[0]
        b = copy.deepcopy(a); b["name"] = "RES-0603-22k"
        (d / "models/SchLib/Resistor.SchLib").write_bytes((d / "models/SchLib/Capacitor.SchLib").read_bytes())
        for dry in (True, False):
            out = S.sync(_plan(d, [a, b]), d / "lib", dry_run=dry)
            assert _actions(out) == {"RES-0603-10k": "error", "RES-0603-22k": "error"}, out
        assert not (d / "lib/SchLib/Resistor.SchLib").exists() and not (d / "lib/ecadlib.xlsx").exists()


def test_csv_resaved_in_cp1252_does_not_crash():
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        plan = _plan(d, _base_items())
        S.sync(plan, d / "lib")
        (d / "lib/ecadlib.csv").write_bytes("Table,Part Number\nR\xe9sistors,X\n".encode("cp1252"))
        out = S.sync(plan, d / "lib")
        assert "ecadlib.csv" in out["written"]


def test_connection_is_portable():
    """On every OS (the Windows CI job runs this too) the Data Source is --win-dir or the fixed placeholder, never the
    library's own absolute path: Altium opens the workbook through LibraryDatabasePath + DatabasePathRelative=1."""
    with tempfile.TemporaryDirectory() as t:
        d = Path(t); _models(d)
        lib = d / "lib"
        S.sync(_plan(d, _base_items()), lib)
        f = lib / "ecadlib.DbLib"
        text = f.read_text()
        assert "Data Source=C:\\ecadlib\\library\\ecadlib.xlsx;" in text, text[:300]
        assert str(S.PureWindowsPath(S.os.path.abspath(lib))) not in text and str(lib) not in text and t not in text
        assert "LibraryDatabasePath=ecadlib.xlsx" in text and "DatabasePathRelative=1" in text
        assert S.win_folder() == S.PureWindowsPath(S.PLACEHOLDER_DIR)
        assert S.win_folder("D:\\libs") == S.PureWindowsPath("D:\\libs")
        # a connection set in Altium's DbLib editor stays; --win-dir replaces it
        f.write_bytes(f.read_bytes().replace(b"C:\\ecadlib\\library", b"E:\\mine"))
        S.sync(_plan(d, _base_items()), lib)
        assert "Data Source=E:\\mine\\ecadlib.xlsx" in f.read_text()
        S.sync(_plan(d, _base_items()), lib, win_dir="D:\\libs")
        assert "Data Source=D:\\libs\\ecadlib.xlsx" in f.read_text()


def test_menu_part_goes_to_the_table_its_source_names():
    # a real "Add part from MPN" run filed RC0603FR-0710KL (EasyEDA C98220, "Resistors, R0603") under ICs: the menu
    # writes designator U1 for every part, and the designator was the only thing that chose the table
    import pandas as pd
    from ecadlib import altium365 as A365, providers as P
    lines = [{"designator": "U1", "mpn": "RC0603FR-0710KL", "value": "RC0603FR-0710KL", "lcsc": "C98220"},
             {"designator": "U2", "mpn": "CL05B104KO5NNNC", "value": "CL05B104KO5NNNC", "lcsc": "C1525"},
             {"designator": "U3", "mpn": "STM32F103C8T6", "value": "STM32F103C8T6", "lcsc": "C8734"},
             {"designator": "U4", "mpn": "IRS2003", "value": "IRS2003", "description": "MOSFET driver"}]
    tri = pd.DataFrame([{"mpn": "RC0603FR-0710KL", "package": "R0603", "local_footprint": "",
                         "src_description": P.describe("", "Resistors", "R0603")},
                        {"mpn": "CL05B104KO5NNNC", "package": "", "local_footprint": "Capacitor_SMD:C_0402_1005Metric",
                         "src_description": ""},
                        {"mpn": "STM32F103C8T6", "package": "LQFP-48", "local_footprint": "",
                         "src_description": "ARM Cortex-M3 MCU 64KB Flash"}], dtype=str)
    got = {i["name"]: S.table_of(i) for i in A365.plan(lines, tri)}
    assert got["RC0603FR-0710KL"] == "Resistors", got           # named by its MPN (the menu sets Value = MPN)
    assert got["CL05B104KO5NNNC"] == "Capacitors", got
    assert got["STM32F103C8T6"] == "ICs" and got["IRS2003"] == "ICs", got
    assert X.kind_of(["Chip Resistor - Surface Mount"]) == "Resistors"
    assert X.kind_of(["Resistor Networks & Arrays", "Capacitive touch sensor"]) is None
    # what the real run's triage held for C98220: LCSC's long text and a bare "0603" package
    lcsc = "'-55℃~+155℃ 100mW 10kΩ 75V Thick Film Resistor ±1% ±100ppm/℃ 0603 Chip Resistor - Surface Mount ROHS"
    assert X.kind_of(["", lcsc], ["", "", "0603"]) == "Resistors"
    assert X.kind_of(["Gate driver IC, 600V half bridge - SOIC-8"]) is None
    assert X.kind_of([lcsc]) is None                    # the noun alone is not enough: an LDO "stable with ceramic
    assert X.kind_of(["LDO regulator, stable with ceramic output capacitors"], ["", "", "SOT-223"]) is None  # capacitors"


def test_lcsc_number_footprint_is_not_a_chip_capacitor():
    # an EasyEDA part's footprint is named by its LCSC number: easyeda:C7593 is an NE555, not "C" + a 4-digit size
    for fp in ("easyeda:C7593", "easyeda:C6186", "LCSC:C1525", "C2002", "easyeda:C1206"):
        assert X.kind_of([], [fp]) is None, fp
    rows = X.to_rows([{"designator": "U1", "mpn": "NE555DR", "value": "NE555DR", "footprint": "easyeda:C7593"}])
    assert list(rows) == ["ICs"], list(rows)
    for fp, cat in (("C0603", "Capacitors"), ("R0402", "Resistors"), ("L1008", "Inductors"),
                    ("Capacitor_SMD:C_0805_2012Metric", "Capacitors"), ("R_01005_0402Metric", "Resistors")):
        assert X.kind_of([], [fp]) == cat, fp
    # an IC package outranks the words: a digipot / matched pair / ESD array is an IC
    assert X.kind_of(["Precision Matched Resistors"], ["SOT-23"]) is None
    assert X.kind_of(["Programmable Resistor"], ["", "MSOP-10"]) is None
    assert X.kind_of(["ESD Protection for Connector"], ["Package_TO_SOT_SMD:SOT-23-6"]) is None
    assert X.kind_of(["USB Connectors"], ["SMD"]) == "Connectors"


if __name__ == "__main__":
    from tests.run_all import main
    raise SystemExit(main(["test_dblib_sync"]))
