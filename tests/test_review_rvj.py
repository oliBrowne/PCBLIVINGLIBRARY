"""Review rvJ regressions: BOM inputs, KiCad database library, parts export, Altium %UTF8% twins, plan lock and
atomic writes on Windows (simulated: a fake msvcrt, a rename refused with PermissionError)."""
import contextlib, json, os, shutil, sqlite3, sys, tempfile, threading, time, types
from pathlib import Path

from ecadlib import agent_queue as AQ, altium_lib as AL, bom_inputs as BI, fsutil, kicad_dblib as KD, parts_export as PE

FIX = Path(__file__).parent / "fixtures" / "kicad_sch"


def _plan(d, items):
    (Path(d) / "plan.json").write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
    return Path(d) / "plan.json"


def _item(name, folder, params=(), description=""):
    return {"name": name, "status": "ready", "folder": folder, "description": description,
            "symbol": {"ref": name, "file": "s.SchLib"}, "footprint": {"ref": "R0603", "file": "f.PcbLib"},
            "parameters": [{"name": k, "value": v} for k, v in params]}


def test_unannotated_references_are_not_merged():
    """Two unannotated 'R?' resistors on one sheet counted as one part (qty 1 instead of 2)."""
    with tempfile.TemporaryDirectory() as d:
        shutil.copytree(FIX, Path(d) / "k")
        f = Path(d) / "k" / "board.kicad_sch"
        f.write_text(f.read_text(encoding="utf-8").replace('"R1"', '"R?"').replace('"R2"', '"R?"'), encoding="utf-8")
        with contextlib.redirect_stderr(open(os.devnull, "w")):
            lines, _ = BI.read_kicad_sch(f)
        want, _ = BI.read_kicad_sch(FIX / "board.kicad_sch")
        assert sum(int(l["qty"]) for l in lines) == sum(int(l["qty"]) for l in want) == 9
        r10k = next(l for l in lines if "R?" in l["designator"])
        assert (r10k["designator"], r10k["qty"]) == ("R30, R?", "3"), r10k


def test_kicad_dblib_case_insensitive_tables_and_columns():
    """Folders 'Resistors' / 'resistors' or parameters 'Tolerance' / 'tolerance' crashed the build with sqlite3
    OperationalError (SQLite names ignore case)."""
    with tempfile.TemporaryDirectory() as d:
        plan = _plan(d, [_item("R1", "Passives/Resistors", [("Tolerance", "1%"), ("tolerance", "5%")]),
                         _item("R2", "Other/resistors", [("tolerance", "2%")])])
        tables, _ = KD.part_rows(plan, "lib")
        assert list(tables) == ["Resistors"] and len(tables["Resistors"]) == 2
        KD.write_sqlite(Path(d) / "x.sqlite", tables)
        with contextlib.closing(sqlite3.connect(Path(d) / "x.sqlite")) as con:
            got = dict(con.execute('SELECT "Part ID", "Tolerance" FROM "Resistors"').fetchall())
        assert got == {"R1": "1%", "R2": "2%"}, got


def test_kicad_dblib_rewrites_open_database_on_windows():
    """On Windows the .sqlite KiCad has open can't be renamed over (PermissionError); the tables are then replaced in
    one transaction inside the existing file."""
    with tempfile.TemporaryDirectory() as d:
        db = Path(d) / "x.sqlite"
        KD.write_sqlite(db, {"Old": [{"Part ID": "A", "Symbol": "l:A", "Footprint": ""}]})
        reader = sqlite3.connect(db)                       # KiCad's ODBC connection
        real, retry = os.replace, getattr(fsutil, "RETRY", False)

        def refuse(src, dst):
            raise PermissionError(13, "The process cannot access the file", str(dst))
        os.replace, fsutil.RETRY = refuse, False
        try:
            KD.write_sqlite(db, {"Resistors": [{"Part ID": "R1", "Symbol": "l:R1", "Footprint": "l:R0603", "Value": "10k"}]})
        finally:
            os.replace, fsutil.RETRY = real, retry
        assert not db.with_suffix(".sqlite.tmp").exists()
        names = [n for (n,) in reader.execute("SELECT name FROM sqlite_master WHERE type='table'")]
        assert names == ["Resistors"], names
        assert reader.execute('SELECT "Value" FROM "Resistors"').fetchall() == [("10k",)]
        reader.close()


def test_atomic_write_retries_while_windows_holds_the_file():
    """os.replace over a file another process has open fails on Windows for as long as it is open; the write now
    retries instead of failing the approve / release check outright."""
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "plan.json"
        p.write_text("old", encoding="utf-8")
        real, retry, calls = os.replace, getattr(fsutil, "RETRY", False), []

        def busy(src, dst):
            calls.append(1)
            if len(calls) < 3:
                raise PermissionError(13, "Access is denied", str(dst))
            return real(src, dst)
        os.replace, fsutil.RETRY = busy, True
        try:
            fsutil.write_text_atomic(p, "new")
        finally:
            os.replace, fsutil.RETRY = real, retry
        assert p.read_text(encoding="utf-8") == "new" and len(calls) == 3
        assert [f.name for f in Path(d).iterdir()] == ["plan.json"]           # no temp file left behind


def test_plan_lock_waits_on_windows_past_msvcrt_ten_seconds():
    """msvcrt.LK_LOCK gives up with OSError after 10 one-second tries, so `ecadlib release-check` waiting for a
    librarian run of minutes crashed on Windows. A fake msvcrt (held for 12 'tries') must not make wait=True fail,
    and the lock is released (LK_UNLCK) before the file is closed."""
    state = {"held": 12, "calls": []}

    def locking(fd, mode, n):
        state["calls"].append(mode)
        if mode == fake.LK_LOCK:                         # the real LK_LOCK: 10 tries, then OSError
            if state["held"] > 10:
                raise OSError(36, "Resource deadlock avoided")
            return
        if mode == fake.LK_NBLCK:
            if state["held"] > 0:
                state["held"] -= 1
                raise OSError(13, "Permission denied")
    fake = types.SimpleNamespace(LK_LOCK=1, LK_NBLCK=2, LK_UNLCK=0, locking=locking)
    real_import = __builtins__["__import__"] if isinstance(__builtins__, dict) else __builtins__.__import__

    def imp(name, *a, **k):
        if name == "fcntl":
            raise ImportError("no fcntl on Windows")
        if name == "msvcrt":
            return fake
        return real_import(name, *a, **k)
    import builtins
    sleep = time.sleep
    builtins.__import__, time.sleep = imp, (lambda s: None)
    try:
        with tempfile.TemporaryDirectory() as d:
            with AQ.plan_lock(d, wait=True) as got:
                assert got
            state["held"] = 1
            with AQ.plan_lock(d, wait=False) as got:
                assert not got
    finally:
        builtins.__import__, time.sleep = real_import, sleep
    assert fake.LK_UNLCK in state["calls"]
    assert state["calls"].count(fake.LK_UNLCK) == 1                     # only the lock that was taken is released


def test_parts_export_xlsx_with_control_characters():
    """A description with a vertical tab (as PDF text extraction gives) made openpyxl refuse the whole workbook."""
    try:
        import openpyxl
    except ImportError:
        print("skip: openpyxl not installed")
        return
    with tempfile.TemporaryDirectory() as d:
        rs = PE.rows([_item("U1", "ICs", [("Note", "=1+1")], description="LDO\x0b3.3V\x1f")])
        out = PE.write(rs, Path(d) / "p.xlsx", "xlsx")
        ws = openpyxl.load_workbook(out).active
        hdr = [c.value for c in ws[1]]
        row = {h: c.value for h, c in zip(hdr, ws[2])}
        assert row["Description"] == "LDO 3.3V ", row["Description"]
        assert row["Other Parameters"] == "'Note=1+1" or not str(row["Other Parameters"]).startswith("=")


def test_exports_read_and_write_utf8_whatever_the_locale():
    """plan.json / sym-lib-table were read in the locale's encoding: on Windows (cp1252) or a C locale a UTF-8
    'Ω' or 'Bibliothèque' failed with UnicodeDecodeError or came out garbled."""
    import subprocess
    code = r'''
import json, sys, tempfile
from pathlib import Path
from ecadlib import kicad_dblib as KD, parts_export as PE
d = Path(tempfile.mkdtemp())
(d / "sym-lib-table").write_bytes('(sym_lib_table\n\t(version 7)\n\t(lib (name "Biblioth\u00e8que")(type "KiCad")(uri "x")(options "")(descr "\u03a9"))\n)\n'.encode())
KD.add_lib_row(d / "sym-lib-table", "x-db", "${KIPRJMOD}/x.kicad_dbl", "d")
assert "Biblioth\u00e8que" in (d / "sym-lib-table").read_bytes().decode("utf-8")
(d / "plan.json").write_bytes(json.dumps([{"name": "R1", "status": "ready", "description": "10k\u03a9", "symbol": {"ref": "R", "file": "s"},
                                          "footprint": {}}], ensure_ascii=False).encode())
assert PE.rows(d / "plan.json")[0]["Description"] == "10k\u03a9"
assert KD.part_rows(d / "plan.json", "x")[0]["Parts"][0]["Description"] == "10k\u03a9"
print("ok")
'''
    env = dict(os.environ, LC_ALL="C", LANG="C", PYTHONCOERCECLOCALE="0", PYTHONUTF8="0", PYTHONIOENCODING="utf-8")
    r = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).parent.parent, env=env, capture_output=True, text=True)
    assert r.stdout.strip() == "ok", r.stderr[-800:]


def test_schlib_utf8_twin_for_every_non_ascii_value():
    """Altium writes a %UTF8% twin for every non-ASCII value (451 of 451 in the Celestial library, '80µA' and '90°'
    included); ours only did for text cp1252 can't hold, so '±1%' had no twin."""
    sym = dict(AL.resistor_symbol("R_10K", "Resistor 10k ±1%"), params=[("Temp", "-40°C")])
    with tempfile.TemporaryDirectory() as d:
        p = AL.write_schlib(Path(d) / "u.SchLib", [sym])
        import olefile
        with olefile.OleFileIO(str(p)) as ole:                                 # closed: Windows can't delete open files
            raw = ole.openstream("R_10K/Data").read()
        assert b"|%UTF8%Text=-40\xc2\xb0C|||Text=-40\xb0C|" in raw
        assert b"|%UTF8%ComponentDescription=Resistor 10k \xc2\xb11%|||ComponentDescription=Resistor 10k \xb11%|" in raw
        h, syms = AL.read_schlib(p)
        assert h["COMPDESCR0"] == "Resistor 10k ±1%"
        assert [r["TEXT"] for r in syms[0]["records"] if r.get("NAME") == "Temp"] == ["-40°C"]
    assert AL.schblock("|RECORD=41|Text=10k|Name=Value") == AL.pblock("|RECORD=41|Text=10k|Name=Value")   # ASCII unchanged
