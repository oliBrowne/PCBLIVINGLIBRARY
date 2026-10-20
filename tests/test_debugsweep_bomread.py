"""Debug-sweep regressions (BOM reading, group bomread). Plain functions: python -m tests.run_all test_debugsweep_bomread"""
import contextlib, io

from ecadlib import bom_triage as T
from tests.fixtures_lite import lite


@lite
def test_ds2_03_trailing_unnamed_notes_column_keeps_mpn(tmp_path):
    # Excel writes every row to the sheet's width: a notes column with no header gives 'Footprint,' in the header
    f = tmp_path / "notes_col.csv"
    f.write_bytes(b"Designator,Qty,Value,MPN,Footprint,\r\nR1,1,10k,RC0603FR-0710KL,0603,check stock\r\n"
                  b"C1,1,100n,GRM188R71H104KA93D,0603,\r\nU1,1,LDO,AMS1117-3.3,SOT-223,alt: LM1117\r\n"
                  b"R2,1,4,7k,RC0603FR-074K7L,0603,\r\n")
    lines, _ = T.read_bom(f)
    got = [(l["designator"], l["value"], l.get("mpn"), l.get("footprint"), l.get("bom_note")) for l in lines]
    assert got[:3] == [("R1", "10k", "RC0603FR-0710KL", "0603", None),
                       ("C1", "100n", "GRM188R71H104KA93D", "0603", None),
                       ("U1", "LDO", "AMS1117-3.3", "SOT-223", None)], got
    # a row one cell wider than the header row is still an unquoted decimal comma
    assert got[3][:4] == ("R2", "4,7k", "RC0603FR-074K7L", "0603") and "decimal comma" in got[3][4], got


@lite
def test_ds2_03_a_large_csv_reads_in_linear_time(tmp_path):
    # the per-row cell counts must not ride in df.attrs: pandas deep-copies attrs into every row iterrows() makes,
    # so 8000 rows took ~15 s (the web UI and MCP triage read uploads with read_bom); linear, it is under a second
    import time
    f = tmp_path / "big.csv"
    f.write_text("Designator,Qty,Value,MPN,Footprint,\n"
                 + "".join(f"R{i},1,10k,RC0603FR-0710KL,0603,\n" for i in range(1, 8001)), encoding="utf-8")
    t0 = time.perf_counter()
    lines, _ = T.read_bom(f)
    took = time.perf_counter() - t0
    assert len(lines) == 8000 and lines[-1]["designator"] == "R8000" and lines[-1]["mpn"] == "RC0603FR-0710KL", lines[-1]
    assert took < 6, f"read_bom took {took:.1f}s for 8000 rows"


@lite
def test_ds2_04_hidden_column_in_a_sparse_workbook_is_the_one_left_out(tmp_path):
    import openpyxl
    from openpyxl.styles import PatternFill

    def book(name, first_col, rows, hide):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "BOM"
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row, first_col):
                if v:
                    ws.cell(r, c, v)
        ws.column_dimensions[hide].hidden = True
        ws["Z100000"].fill = PatternFill("solid", fgColor="FFFF00")   # far, empty, styled: the sheet is read sparsely
        wb.save(tmp_path / name)
        return tmp_path / name

    b = book("from_b.xlsx", 2, [["Designator", "Value", "MPN (old)", "MPN"], ["R1", "10k", "OLD-OBSOLETE-1", "RC0603FR-0710KL"],
                                ["C1", "100n", "OLD-OBSOLETE-2", "GRM188R71H104KA93D"]], "D")
    a = book("spacer.xlsx", 1, [["Designator", "Value", "", "MPN"], ["R1", "10k", "", "RC0603FR-0710KL"],
                                ["C1", "100n", "", "GRM188R71H104KA93D"]], "C")    # an empty hidden spacer column
    for f in (b, a):
        lines, mapping = T.read_bom(f)
        assert mapping["mpn"] == "MPN" and [l.get("mpn") for l in lines] == ["RC0603FR-0710KL", "GRM188R71H104KA93D"], \
            (f.name, mapping, lines)


@lite
def test_ds2_05_triage_reports_of_dotted_bom_names_keep_the_whole_name(tmp_path, monkeypatch):
    from ecadlib.app import Workspace
    (tmp_path / "ecadlib.toml").write_text('[setup]\nworkspace = "ws"\n', encoding="utf-8")
    for rev, val in (("1.1", "10k"), ("1.2", "4k7")):
        (tmp_path / f"PowerBoard_rev{rev}_BOM.csv").write_text(f"Designator,Qty,Value,MPN\nR1,1,{val},RC0603FR-0710KL\n",
                                                               encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        for rev in ("1.1", "1.2"):                                       # CLI, default -o: <workspace>/reports/triage_<BOM>
            T.main([f"PowerBoard_rev{rev}_BOM.csv", "--providers", "", "--offline"])
        reports = tmp_path / "ws" / "reports"
        got = sorted(p.name for p in reports.iterdir())
        assert got == ["triage_PowerBoard_rev1.1_BOM.csv", "triage_PowerBoard_rev1.1_BOM.md",
                       "triage_PowerBoard_rev1.2_BOM.csv", "triage_PowerBoard_rev1.2_BOM.md"], got
        assert "rev1.1" in (reports / "triage_PowerBoard_rev1.1_BOM.md").read_text(encoding="utf-8").splitlines()[0]

        w = Workspace(tmp_path / "web", providers=[], offline=True)      # web / MCP: an upload keyed by its file name
        up = w.save_upload("PowerBoard_rev1.2_BOM.csv", b"Designator,Qty,Value,MPN\nR1,1,4k7,RC0603FR-074K7L\n")
        r = w.triage(up)
        assert r["reports"] == ["reports/triage_PowerBoard_rev1.2_BOM@csv.md",
                                "reports/triage_PowerBoard_rev1.2_BOM@csv.csv"], r["reports"]
        assert w.triage_result(r["id"])["name"] == "PowerBoard_rev1.2_BOM.csv"
        assert w.build_altium_models(up)["skipped"] == []              # reads the triage it just wrote


@lite
def test_ds2_09_integer_value_is_not_glued_to_the_next_column(tmp_path):
    # one unquoted comma in the Description: Value 100 then Qty 2 / Footprint 0603 is not the decimal value 100,2
    q = tmp_path / "qty.csv"
    q.write_text("Designator,Value,Qty,Footprint,MPN,Description\n"
                 "R1 R2,100,2,0603,RC0603FR-07100RL,RES 100 ohm, 1%\n"
                 "R3 R4,10,2,0603,RC0603FR-0710RL,RES 10 ohm, 1%\n"
                 "R5,4,7k,1,0603,RC0603FR-074K7L,RES\n"
                 "R6,22,1,1,0603,RC0603FR-0722R1L,RES\n", encoding="utf-8")
    f = tmp_path / "fp.csv"
    f.write_text("Designator,Value,Footprint,Description\nR1,100,0603,RES 100 ohm, 1%\nR2,4,7k,0603,RES\n", encoding="utf-8")
    got = [(l["designator"], l["value"], l.get("qty"), l.get("footprint"), l.get("mpn"))
           for p in (q, f) for l in T.read_bom(p)[0]]
    assert got == [("R1 R2", "100", "2", None, None), ("R3 R4", "10", "2", None, None),
                   ("R5", "4,7k", "1", "0603", "RC0603FR-074K7L"), ("R6", "22,1", "1", "0603", "RC0603FR-0722R1L"),
                   ("R1", "100", None, None, None), ("R2", "4,7k", None, "0603", None)], got
    notes = [l.get("bom_note", "") for p in (q, f) for l in T.read_bom(p)[0]]
    assert "header has 6" in notes[0] and "header has 6" in notes[1] and "header has 4" in notes[4], notes


@lite
def test_ds2_09_voltage_or_tolerance_after_value_is_its_own_column(tmp_path):
    # Value 100 then Voltage 50V / Tolerance 1%, the stray comma in the Description: not the value '100,50V'
    v = tmp_path / "volt.csv"                       # well-formed rows show what the Voltage column holds
    v.write_text("Designator,Value,Voltage,Footprint,MPN,Description\n"
                 "C10,100n,50V,0603,GRM188R71H104KA93D,CAP CER 100nF 50V X7R\n"
                 "C11,1u,25V,0603,GRM188R61E105KA12D,CAP CER 1uF 25V X5R\n"
                 "C1,100,50V,0603,GRM1885C1H101JA01D,CAP 100pF, C0G\n"
                 "C2,10,16V,0603,GRM188R61C106KA73D,CAP 10uF, X5R\n"
                 "C3,4,7u,25V,0603,GRM188R61E475KE11D,CAP\n", encoding="utf-8")
    t = tmp_path / "tol.csv"                        # no well-formed row: a '1%' is a tolerance, never a value's decimals
    t.write_text("Designator,Value,Tolerance,Footprint,MPN,Description\n"
                 "R1,10,1%,0603,RC0603FR-0710RL,RES 10 ohm, 1%\n"
                 "R2,47,5%,0603,RC0603JR-0747RL,RES 47 ohm, 5%\n"
                 "R3,4,7k,1%,0603,RC0603FR-074K7L,RES\n", encoding="utf-8")
    q = tmp_path / "volt_qty.csv"                   # no well-formed row: the Qty after Value still counts the designators
    q.write_text("Designator,Value,Voltage,Qty,Footprint,MPN,Description\n"
                 "C1 C2,100,50V,2,0603,GRM1885C1H101JA01D,CAP 100pF, C0G\n"
                 "C3 C4,4,7u,25V,2,0603,GRM188R61E475KE11D,CAP\n", encoding="utf-8")
    got = {p.name: [(l["designator"], l["value"], l.get("qty"), l.get("footprint"), l.get("mpn"))
                    for l in T.read_bom(p)[0]] for p in (v, t, q)}
    assert got == {
        "volt.csv": [("C10", "100n", None, "0603", "GRM188R71H104KA93D"), ("C11", "1u", None, "0603", "GRM188R61E105KA12D"),
                     ("C1", "100", None, None, None), ("C2", "10", None, None, None),
                     ("C3", "4,7u", None, "0603", "GRM188R61E475KE11D")],
        "tol.csv": [("R1", "10", None, None, None), ("R2", "47", None, None, None),
                    ("R3", "4,7k", None, "0603", "RC0603FR-074K7L")],
        "volt_qty.csv": [("C1 C2", "100", "2", None, None), ("C3 C4", "4,7u", "2", "0603", "GRM188R61E475KE11D")],
    }, got
    notes = [l.get("bom_note", "") for p in (v, t, q) for l in T.read_bom(p)[0]]
    assert "header has 6" in notes[2] and "header has 6" in notes[5] and "header has 7" in notes[8], notes


@lite
def test_ds2_10_real_manufacturer_and_description_columns_win_whatever_the_order(tmp_path):
    v = tmp_path / "vendor.csv"
    v.write_text("Designator,Qty,Value,Vendor,Vendor PN,Manufacturer,MPN\n"
                 "U1,1,LM358,Digi-Key,296-1395-1-ND,Texas Instruments,LM358DR\n", encoding="utf-8")
    n = tmp_path / "notes.csv"
    n.write_text("Designator,Notes,Value,Description,MPN\n"
                 "C1,check polarity,100uF,Tantalum capacitor 100uF 10V,T491D107K010AT\n", encoding="utf-8")
    only = tmp_path / "only_notes.csv"
    only.write_text("Designator,Value,Notes,MPN\nC1,100uF,Tantalum 10V,T491D107K010AT\n", encoding="utf-8")
    lines, mapping = T.read_bom(v)
    assert lines[0]["manufacturer"] == "Texas Instruments", (mapping, lines)
    lines, mapping = T.read_bom(n)
    assert lines[0]["description"] == "Tantalum capacitor 100uF 10V", (mapping, lines)
    assert T.read_bom(only)[0][0]["description"] == "Tantalum 10V"           # Notes is still the fallback
    # a purchasing BOM with no maker column: 'Vendor' beside 'Vendor Part Number' is a distributor, not the maker
    pur = tmp_path / "purchasing.csv"
    pur.write_text("Designator,Value,Vendor,Vendor Part Number,MPN\nU1,LM358,Digi-Key,296-1395-1-ND,LM358DR\n"
                   "C1,100n,LCSC,C14663,CL10B104KB8NNNC\n", encoding="utf-8")
    lines, mapping = T.read_bom(pur)
    assert "manufacturer" not in mapping, mapping
    assert [(l["designator"], l.get("manufacturer"), l.get("lcsc")) for l in lines] == \
        [("U1", None, None), ("C1", None, "C14663")], lines


@lite
def test_ds2_11_kicad_description_backs_up_a_house_desc_field_and_notes_backs_up_description(tmp_path):
    def sym(ref, value, desc, notes, mpn):
        props = {"Reference": ref, "Value": value, "Footprint": "Resistor_SMD:R_0603_1608Metric", "Description": desc,
                 "Notes": notes, "MPN": mpn}
        p = "\n".join(f'    (property "{k}" "{v}" (at 0 0 0))' for k, v in props.items())
        return (f'  (symbol (lib_id "Device:R") (at 0 0 0) (unit 1) (in_bom yes) (on_board yes) (dnp no)\n'
                f'    (uuid 2222222{ref[-1]}-2222-2222-2222-222222222222)\n{p}\n'
                f'    (instances (project "p" (path "/11111111-1111-1111-1111-111111111111" (reference "{ref}") (unit 1))))'
                '\n  )')
    sch = tmp_path / "board.kicad_sch"
    sch.write_text('(kicad_sch (version 20231120) (generator eeschema) (uuid 11111111-1111-1111-1111-111111111111) '
                   '(paper "A4") (lib_symbols)\n' + sym("C1", "10uF", "Tantalum capacitor 10uF 16V", "", "T491A106K016AT")
                   + "\n" + sym("R1", "10k", "Resistor", "pull-up, check value", "RC0603FR-0710KL") + "\n)\n",
                   encoding="utf-8")
    net = tmp_path / "board.xml"
    net.write_text('<?xml version="1.0" encoding="UTF-8"?>\n<export version="E"><components>\n'
                   '<comp ref="C1"><value>10uF</value><description>Tantalum capacitor 10uF 16V</description>'
                   '<fields><field name="Desc"></field><field name="MPN">T491A106K016AT</field></fields></comp>\n'
                   '<comp ref="R1"><value>10k</value><description>Resistor</description>'
                   '<fields><field name="Desc">RES SMD 10K 1% 0603</field><field name="MPN">RC0603FR-0710KL</field></fields>'
                   '</comp>\n</components></export>\n', encoding="utf-8")
    # a note never replaces the symbols' Description; a house 'Desc' filled per part beats the library's generic
    # text, and an empty one falls back to it
    for f, want in ((sch, [("C1", "Tantalum capacitor 10uF 16V"), ("R1", "Resistor")]),
                    (net, [("C1", "Tantalum capacitor 10uF 16V"), ("R1", "RES SMD 10K 1% 0603")])):
        lines, mapping = T.read_bom(f)
        assert [(l["designator"], l.get("description")) for l in lines] == want, (f.name, mapping, lines)
    lines, _ = T.read_bom(sch, {"description": "Notes"})               # the rules' [bom] columns still decide
    assert [l.get("description") for l in lines] == [None, "pull-up, check value"], lines
