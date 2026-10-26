"""Parts table export for PLM / ERP imports:  ecadlib export-parts plan.json --format csv|json|xlsx -o parts.csv

One row per library part in a librarian plan, with stable columns (COLUMNS; their names and order only ever grow at
the end, so an import mapping set up once keeps working). Parameters that have no column of their own go in
"Other Parameters" (CSV / XLSX: "name=value; ..." sorted by name; JSON: an object). Held parts are left out unless
--include-held. Cells that a spreadsheet would run as a formula get a leading apostrophe in CSV (sheetsafe); XLSX
cells are text cells holding the value as it is (+3.3V), except that a leading '=' keeps the apostrophe; JSON keeps
the plain text.

KiCad names ("KiCad Symbol", "KiCad Footprint") are the lib ids `ecadlib kicad-export --name LIB` writes
(LIB:<part>, LIB:<footprint>), the same ones `ecadlib kicad-dblib` puts in its database."""
import argparse, json, sys
from pathlib import Path

from . import sheetsafe as SS

FORMAT_VERSION = 1
COLUMNS = ["Part ID", "Status", "Category", "Folder", "Description", "Value", "Manufacturer", "Manufacturer Part Number",
           "Datasheet", "Supplier 1", "Supplier Part Number 1", "Supplier 2", "Supplier Part Number 2",
           "Lifecycle", "Symbol Ref", "Symbol File", "Footprint Ref", "Footprint File", "KiCad Symbol", "KiCad Footprint",
           "3D Height (mm)", "Release Check", "Hold Reasons", "Sign-off", "Approved By", "Symbol Source", "Footprint Source",
           "Other Parameters"]
PARAM_COLUMNS = {"Value": "Value", "Manufacturer": "Manufacturer", "Manufacturer Part Number": "Manufacturer Part Number",
                 "ComponentLink2URL": "Datasheet", "Datasheet": "Datasheet", "Supplier 1": "Supplier 1",
                 "Supplier Part Number 1": "Supplier Part Number 1", "Supplier 2": "Supplier 2",
                 "Supplier Part Number 2": "Supplier Part Number 2", "Lifecycle": "Lifecycle", "Lifecycle Status": "Lifecycle"}
SKIP_PARAMS = {"ComponentLink2Description"}
LIBRARY_STATUS = ("ready", "existing")


def kicad_ids(item, lib):
    from .kicad_export import kname
    fp = (item.get("footprint") or {})
    fp = fp if isinstance(fp, dict) else {}
    has_fp = fp.get("ref") and (fp.get("file") or fp.get("kicad_mod"))
    has_sym = (item.get("source") or {}).get("symbol") or (item.get("symbol") or {}).get("file")
    return (f"{lib}:{kname(item['name'])}" if has_sym else "", f"{lib}:{kname(fp['ref'])}" if has_fp else "")


def _signoff(items, base):
    """{part: (sign-off, approved by)} from the approvals policy next to the plan; {} when there is none."""
    from . import approvals as A
    try:
        policy = A.load_policy(base)
    except Exception as e:                       # a broken policy file must not stop an export
        print(f"  warning: approvals policy not read: {e}", file=sys.stderr)
        return {}
    if not policy:
        return {}
    out = {}
    for it in items:
        try:
            ev = A.evaluate(it, policy, base)
        except Exception as e:
            out[it["name"]] = (f"unknown ({type(e).__name__})", "")
            continue
        state = "not required" if not ev["required"] else "approved" if not ev["missing"] and not ev["stale"] else "pending"
        out[it["name"]] = (state, "; ".join(ev["approved_by"]))
    return out


def rows(plan, lib="ecadlib", include_held=False, rules=None):
    """Plan items -> [row dict with exactly COLUMNS], sorted by Part ID. plan: path to plan.json or a list of items.
    rules: the library rules file the release check runs with (default: ecadlib.toml found from the plan)."""
    from .kicad_export import libname
    base = Path(plan).parent if not isinstance(plan, list) else Path(".")
    items = json.loads(Path(plan).read_text(encoding="utf-8")) if not isinstance(plan, list) else plan
    if not isinstance(items, list):
        raise ValueError(f"{plan}: not a librarian plan (expected a list of parts)")
    if not isinstance(plan, list):             # change control (#105): a hand edit of a released part's fields is
        from . import change_gate as CG        # held, and its released fields go out (as kicad-export ships them)
        with CG.publishing(base):
            items = CG.settled_plan(plan, "export-parts", rules=rules)
    if not isinstance(plan, list):             # the release check again, on the files as they are (as every publisher
        from . import release_check as RC, rules as RU      # runs it): a part that fails it is held, not a library part
        checked = RU.resolve(rules, base)
        for it in items:
            why = RC.publish_problems(it, base, checked, writes_files=True) if isinstance(it, dict) and it.get("name") \
                and it.get("status", "ready") in LIBRARY_STATUS else []
            if why:
                it.update(status="held", hold_reasons=why, release_check=dict(it.get("release_check") or {}, result="held"))
    keep = [it for it in items if isinstance(it, dict) and it.get("name")
            and (it.get("status", "ready") in LIBRARY_STATUS or (include_held and it.get("status") == "held"))]
    sign = _signoff(keep, base) if not isinstance(plan, list) else {}
    from .kicad_export import kname
    sym_by, fp_by = {}, {}                     # KiCad names as kicad-export hands them out (items in kname order, the
    for it in sorted(keep, key=lambda i: kname(i["name"])):         # first wins): two part names with one KiCad name
        sym_by.setdefault(kname(it["name"]), it["name"])            # ('A/B', 'A_B') or two footprints whose file names
        ref = (it.get("footprint") or {}).get("ref") if isinstance(it.get("footprint"), dict) else None   # differ only
        if ref:                                                      # in case: the second has no KiCad id, never the
            fp_by.setdefault(kname(ref).casefold(), ref)             # other part's
    out = []
    for it in sorted(keep, key=lambda i: str(i["name"])):
        par = {str(p.get("name")): str(p.get("value", "")) for p in it.get("parameters") or []
               if isinstance(p, dict) and p.get("name") and str(p.get("value", "")).strip()}
        r = dict.fromkeys(COLUMNS, "")
        folder = str(it.get("folder") or "")
        sym, fp = it.get("symbol") or {}, it.get("footprint") or {}
        ks, kf = kicad_ids(it, libname(lib))   # the library name kicad-export / kicad-dblib give it ("My Lib" -> My_Lib)
        if sym_by.get(kname(it["name"])) != it["name"]:
            ks = ""
        if kf and fp_by.get(kname(fp["ref"]).casefold()) != fp["ref"]:
            kf = ""
        rc = it.get("release_check") or {}
        r.update({"Part ID": str(it["name"]), "Status": str(it.get("status", "ready")),
                  "Category": folder.rstrip("/").split("/")[-1] if folder else "", "Folder": folder,
                  "Description": str(it.get("description") or ""), "Symbol Ref": str(sym.get("ref") or ""),
                  "Symbol File": str(sym.get("file") or ""), "Footprint Ref": str(fp.get("ref") or ""),
                  "Footprint File": str(fp.get("file") or ""), "KiCad Symbol": ks, "KiCad Footprint": kf,
                  "3D Height (mm)": f"{float(rc['height']):.2f}" if isinstance(rc.get("height"), (int, float)) else "",
                  "Release Check": str(rc.get("result") or ""), "Hold Reasons": "; ".join(it.get("hold_reasons") or []),
                  "Symbol Source": str((it.get("source") or {}).get("symbol") or ""),
                  "Footprint Source": str((it.get("source") or {}).get("footprint") or "")})
        r["Sign-off"], r["Approved By"] = sign.get(it["name"], ("", ""))
        other = {}
        for k in sorted(par):
            col = PARAM_COLUMNS.get(k)
            if col and not r[col]:
                r[col] = par[k]
            elif k not in SKIP_PARAMS and (not col or r[col] != par[k]):   # a second, different value: kept here
                other[k] = par[k]
        r["Other Parameters"] = other
        out.append(r)
    return out


def _flat(r):
    return {k: ("; ".join(f"{a}={b}" for a, b in v.items()) if isinstance(v, dict) else v) for k, v in r.items()}


def _xlsx(v):
    """An xlsx cell value: the cells are text cells (never run), so '+3.3V' / '-40..85C' / '@x' stay as they are;
    only a leading '=' keeps sheetsafe's apostrophe (defence in depth for readers that re-type text cells)."""
    return SS.cell(v) if v.startswith("=") else v


def write(rs, out, fmt):
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "json":
        out.write_text(json.dumps({"format": "ecadlib-parts", "version": FORMAT_VERSION, "columns": COLUMNS, "parts": rs},
                                  indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    elif fmt == "csv":
        with open(out, "w", newline="", encoding="utf-8-sig") as f:      # utf-8-sig: Excel opens it as UTF-8
            w = SS.DictWriter(f, fieldnames=COLUMNS)
            w.writeheader()
            w.writerows(_flat(r) for r in rs)
    elif fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE      # control characters (a \x0b from a datasheet) that
        wb = Workbook()                                            # XML can't hold: openpyxl refuses the whole file
        ws = wb.active
        ws.title = "Parts"
        ws.append(COLUMNS)
        for r in rs:
            ws.append([_xlsx(ILLEGAL_CHARACTERS_RE.sub(" ", str(_flat(r)[c]))) for c in COLUMNS])   # text cells below
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.data_type, c.number_format = "s", "@"                    # text: '0603' and '1e-7' stay as typed
        ws.freeze_panes = "A2"
        wb.save(out)
    else:
        raise ValueError(f"unknown format {fmt!r}: csv, json or xlsx")
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description="Export the library's parts table (stable columns) for PLM / ERP imports")
    ap.add_argument("plan", help="librarian plan.json")
    ap.add_argument("--format", choices=("csv", "json", "xlsx"), help="default: from the -o suffix, else csv")
    ap.add_argument("-o", "--out", help="output file (default: parts.<format> next to the plan)")
    ap.add_argument("--name", default="ecadlib", help="KiCad library name used by kicad-export (default %(default)s)")
    ap.add_argument("--include-held", action="store_true", help="also export parts the release check held")
    ap.add_argument("--rules", help="library rules file the release check runs with (default: ecadlib.toml found from the plan)")
    a = ap.parse_args(argv)
    fmt = a.format or (Path(a.out).suffix.lower().lstrip(".") if a.out and Path(a.out).suffix.lower() in (".csv", ".json", ".xlsx") else "csv")
    out = Path(a.out) if a.out else Path(a.plan).parent / f"parts.{fmt}"
    try:
        rs = rows(a.plan, a.name, a.include_held, a.rules)
    except (OSError, ValueError) as e:
        print(f"cannot read plan: {e}", file=sys.stderr)
        return 2
    write(rs, out, fmt)
    print(f"wrote {len(rs)} parts to {out} ({fmt})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
