"""What the AI features cost, per job, for your model or for every model we know a price for.

    ecadlib costs                       per-job estimate for the model in your rules ([ai] / [datasheet])
    ecadlib costs --model gpt-6-luna    the same for another model ("local:..." for an [ai.endpoints] entry)
    ecadlib costs --all [--md F] [--csv F] [--xlsx F]   every model in the sheet (docs/costs.md is made with --md)

Only two things call a model: a datasheet read (a part no library, EasyEDA or standard package covers) and the
optional visual check ([agent.proof]). Everything else (KiCad libraries, EasyEDA, stdparts, lifecycle, stock,
upstream checks) is free. These are estimates from the request sizes the code builds, not a bill:

  datasheet read   [datasheet] page_budget pages x datasheet.PAGE_TOKENS + the prompt + the output schema in,
                   datasheet.OUT_TOKENS out (Claude synchronous: input billed as a 1.25x cache write)
  visual check     ([agent.proof] max_pages + 2 drawings of ours) x visual_check.IMG_TOKENS_MAX in, 3000 out
  50-line BOM      SHARE of the lines need a datasheet (the rest resolve free)
  500-part library the same share of 500 parts

Retries cost more: a read that comes back incomplete is re-sent with more pages or (route = "ladder") on the
next model; `ecadlib datasheet estimate BOM` gives the worst case of a real BOM. Prices: datasheet.PRICES for
Claude, ecadlib.llm.profiles.PRICES for the others, $0 for a model on your own computer."""
import argparse, csv, json, sys

from . import datasheet as D

SHARE = 0.2              # lines of a typical BOM that need a datasheet (the bundled sample: 3 of 14 held without one)
BOM_LINES, LIBRARY_PARTS = 50, 500
PROOF_OUT = 3000
# The sheet: (label, model id, provider). Local models are free whatever they are.
SHEET = [
    ("Claude Fable 5.1", "claude-fable-5-1", "anthropic"),
    ("Claude Opus 5.5 (default)", "claude-opus-5-5", "anthropic"),
    ("Claude Sonnet 5.5", "claude-sonnet-5-5", "anthropic"),
    ("Claude Haiku 5.5", "claude-haiku-5-5", "anthropic"),
    ("OpenAI GPT-6 Astra", "gpt-6-astra", "openai"),
    ("OpenAI GPT-6.1 Sol", "gpt-6.1-sol", "openai"),
    ("OpenAI GPT-5.6 Terra", "gpt-5.6-terra", "openai"),
    ("OpenAI GPT-5.4 mini", "gpt-5.4-mini", "openai"),
    ("OpenAI GPT-6 Luna", "gpt-6-luna", "openai"),
    ("OpenAI GPT-5 nano", "gpt-5-nano", "openai"),
    ("Google Gemini 4 Argon", "gemini-4-argon", "gemini"),
    ("Google Gemini 3.1 Pro", "gemini-3.1-pro-preview", "gemini"),
    ("Google Gemini 3.8 Flash", "gemini-3.8-flash", "gemini"),
    ("Google Gemini 3.1 Flash-Lite", "gemini-3.1-flash-lite", "gemini"),
    ("DeepSeek V4 Pro", "deepseek-v4-pro", "openai-compatible"),
    ("DeepSeek V4.1 Flash", "deepseek-flash", "openai-compatible"),
    ("Local model (Ollama, LM Studio, vLLM)", "local", "ollama"),
]
COLS = ["model", "id", "tier", "usd_in", "usd_out", "read", "read_batch", "check", "bom50", "bom50_checked",
        "lib500", "lib500_batch", "lib500_check_all"]
HEAD = {"model": "Model", "id": "Model id", "tier": "Tier", "usd_in": "$ in / MTok", "usd_out": "$ out / MTok",
        "read": "Datasheet read", "read_batch": "Read as batch", "check": "Visual check", "bom50": "50-line BOM",
        "bom50_checked": "50-line BOM + checks", "lib500": "500-part library", "lib500_batch": "500-part library, batch",
        "lib500_check_all": "Visual check on all 500"}


def read_tokens(page_budget=None):
    """(input, output) tokens of one datasheet read at the default page budget."""
    pages = page_budget or D.SETTINGS["page_budget"]
    return pages * D.PAGE_TOKENS + 1500 + len(json.dumps(D.wire_schema())) // 4, D.OUT_TOKENS


def check_tokens(max_pages=None):
    from . import visual_check as V
    pages = max_pages or V.PROOF_DEFAULTS["max_pages"]
    return (pages + 2) * V.IMG_TOKENS_MAX + (len(V.SYSTEM) + len(json.dumps(V.output_schema()))) // 3, PROOF_OUT


def price(model, provider="anthropic", rules=None):
    """(input, output) $ per MTok, or None: Claude's table, the profile's price, $0 for a local server."""
    if model == "local" or provider in ("ollama", "lmstudio", "vllm", "llamacpp"):
        return 0.0, 0.0
    if model in D.PRICES:
        return D.PRICES[model]
    from .llm import profiles as P
    p = P.profile(model, "openai-compatible" if provider != "anthropic" else provider, rules)
    if p.get("price_in") is None or p.get("price_out") is None:
        return None
    return float(p["price_in"]), float(p["price_out"])


def jobs(model, provider="anthropic", rules=None, share=SHARE, vision=None):
    """One row of the sheet: USD per job (None where it does not apply or the price is unknown)."""
    from .llm import profiles as P
    claude = provider == "anthropic"
    prof = P.profile(model, provider if claude else "openai-compatible", rules) if model != "local" else {}
    row = {"id": model, "tier": prof.get("tier", "depends"), "usd_in": None, "usd_out": None}
    pr = price(model, provider, rules)
    if not pr:
        return dict(row, **{k: None for k in COLS if k not in row})
    pin, pout = pr
    rin, rout = read_tokens()
    cin, cout = check_tokens()
    read = (rin * (D.CACHE_WRITE if claude else 1.0) * pin + rout * pout) / 1e6
    batch = (rin * pin + rout * pout) / 1e6 * D.BATCH_DISCOUNT if claude else None   # Message Batches: Claude only
    sees = prof.get("vision", True) if vision is None else vision
    check = (cin * pin + cout * pout) / 1e6 if sees else None
    n_bom, n_lib = round(BOM_LINES * share), round(LIBRARY_PARTS * share)
    return dict(row, usd_in=pin, usd_out=pout, read=read, read_batch=batch, check=check, bom50=n_bom * read,
                bom50_checked=n_bom * (read + check) if check is not None else None, lib500=n_lib * read,
                lib500_batch=n_lib * batch if batch is not None else None,
                lib500_check_all=LIBRARY_PARTS * check if check is not None else None)


def sheet(rules=None, share=SHARE):
    return [dict(jobs(mid, prov, rules, share), model=label) for label, mid, prov in SHEET]


def money(v):
    if v is None:
        return "n/a"
    if v == 0:
        return "$0"
    return f"${v:,.4f}" if v < 0.01 else f"${v:,.3f}" if v < 1 else f"${v:,.2f}"


def markdown(rows, share=SHARE):
    rin, rout = read_tokens()
    cin, cout = check_tokens()
    cols = ["model", "usd_in", "usd_out", "read", "read_batch", "check", "bom50", "bom50_checked", "lib500",
            "lib500_batch", "lib500_check_all"]
    out = ["| " + " | ".join(HEAD[c] for c in cols) + " |", "|" + "---|" * len(cols)]
    for r in rows:
        cells = []
        for c in cols:
            v = r.get(c)
            cells.append(str(v) if c == "model" else ("" if v is None else f"${v:g}") if c in ("usd_in", "usd_out")
                         else money(v))
        out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out), (rin, rout, cin, cout, round(BOM_LINES * share), round(LIBRARY_PARTS * share))


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([HEAD.get(c, c) + (" (USD)" if c not in ("model", "id", "tier", "usd_in", "usd_out") else "")
                    for c in COLS])
        for r in rows:
            w.writerow(["" if r.get(c) is None else (round(r[c], 4) if isinstance(r[c], float) else r[c]) for c in COLS])


def write_xlsx(rows, path, share=SHARE):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Cost per job"
    ws.append([HEAD[c] for c in COLS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c) for c in COLS])
    for row in ws.iter_rows(min_row=2, min_col=COLS.index("read") + 1):
        for c in row:
            c.number_format = '"$"#,##0.000'
    for i, c in enumerate(COLS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 30 if c == "model" else 16
    _, (rin, rout, cin, cout, nb, nl) = markdown(rows, share)
    n = wb.create_sheet("How it is estimated")
    for line in [
            "Estimates, not a bill: request sizes from the code, list prices checked 2026-10-16.",
            f"Datasheet read: ~{rin:,} tokens in, ~{rout:,} out. Visual check: ~{cin:,} in, ~{cout:,} out.",
            f"50-line BOM: {nb} lines need a datasheet ({share:.0%}); 500-part library: {nl} parts. The rest cost $0.",
            "Claude synchronous reads bill input as a 1.25x cache write; Message Batches (Claude only) are half price.",
            "Retries (more pages, or the next model on the ladder) add to this; `ecadlib datasheet estimate BOM` gives the worst case.",
            "Other models are called through ecadlib.llm (openai-compatible): no batch discount there.",
            "Local models cost $0 in API fees. n/a = the model has no image input (no visual check) or no known price.",
            "Prices change: set [ai.profiles.\"<model>\"] price_in / price_out in ecadlib.toml to use your own."]:
        n.append([line])
    n.column_dimensions["A"].width = 120
    wb.save(path)


DOC = """# What it costs to run

ecadlib is free and open source. The AI features use **your own API key** (or a model on your own computer),
so you pay your model provider for what you use, and nothing to this project.

Most parts never touch a model. KiCad library parts, EasyEDA parts and standard packages (resistors, capacitors,
SOIC, QFN and the like) are found or built for free, and so are lifecycle, stock and KiCad upstream checks. A model
is used only to **read a datasheet** for a part nothing else covers, and for the optional **visual check** that
compares our footprint and symbol with the datasheet drawing.

## Estimated cost per job

{table}

These are **estimates**, made from the request sizes the code builds and list prices checked on 2026-10-16. They
are not measured bills.

- **Datasheet read:** about {rin:,} tokens in ({pages} picked datasheet pages, the prompt and the output format)
  and about {rout:,} tokens out. On Claude a synchronous read bills its input as a prompt-cache write (1.25x).
- **Read as batch:** Claude's Message Batches API is half price. ecadlib sends a whole BOM as one batch
  (`ecadlib datasheet batch BOM`). Other providers are called one request at a time.
- **Visual check:** about {cin:,} tokens in ({imgs} images) and {cout:,} out. It needs a model that reads images
  (n/a otherwise). Answers are cached, so re-running it costs nothing.
- **50-line BOM:** {nb} of the 50 lines need a datasheet ({share:.0%}), the rest cost $0.
  **500-part library:** {nl} datasheet reads. "Visual check on all 500" is the extra cost of checking every part.
- **Retries cost more.** An incomplete answer is re-sent with more pages, and with `route = "ladder"` it moves
  to the next model. Budget up to about 2x for hard parts.
- **Smaller models are cheaper but weaker.** ecadlib adds guide rails for them (see `ecadlib ai rails`), and
  their answers are held for review more often.
- **Local models** (Ollama, LM Studio, vLLM, llama.cpp) cost nothing in API fees.

## Your model, your BOM

    ecadlib costs                         # these jobs for the model in your ecadlib.toml
    ecadlib costs --model gemini-3.8-flash
    ecadlib datasheet estimate BOM.csv    # this BOM: tokens counted, cost and worst case, nothing sent

Set `max_cost_usd` under `[datasheet]` and `[agent.proof]` for a hard cap on spending. If a price here is out of
date, set your own in `ecadlib.toml`:

    [ai.profiles."gpt-6-luna"]
    price_in = 0.10      # $ per million input tokens
    price_out = 0.50

Prices come from Anthropic's model table for Claude and from public price trackers for the others (benchlm.ai,
promptunit.ai). Gemini 4 Argon goes to $4/$20 on 2026-12-01 and Gemini 3.8 Flash to $1.50/$7.50 on
2027-01-01. DeepSeek bills half these rates off-peak. The figures in this page come from `ecadlib costs --all --md
docs/costs.md`.
"""


def doc(rows, share=SHARE):
    from . import visual_check as V
    table, (rin, rout, cin, cout, nb, nl) = markdown(rows, share)
    return DOC.format(table=table, rin=rin, rout=rout, cin=cin, cout=cout, nb=nb, nl=nl, share=share,
                      pages=D.SETTINGS["page_budget"], imgs=V.PROOF_DEFAULTS["max_pages"] + 2)


def main(argv=None):
    ap = argparse.ArgumentParser(prog="ecadlib costs", description=__doc__.split("\n\n")[0])
    ap.add_argument("--model", help="model id (default: the one your rules use for datasheets)")
    ap.add_argument("--all", action="store_true", help="every model in the sheet")
    ap.add_argument("--share", type=float, default=SHARE, help=f"share of BOM lines that need a datasheet (default {SHARE})")
    ap.add_argument("--md", help="write the cost page (Markdown) here")
    ap.add_argument("--csv", help="write the sheet as CSV here")
    ap.add_argument("--xlsx", help="write the sheet as an Excel workbook here")
    ap.add_argument("--rules", help="rules file (default: ecadlib.toml found from here)")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args(argv)
    if not 0 <= a.share <= 1:
        ap.error("--share is a fraction from 0 to 1")
    from . import rules as R
    rules = R.load(a.rules)
    if a.all or a.md or a.csv or a.xlsx:
        rows = sheet(rules, a.share)
    else:
        from .llm import endpoint
        ep = endpoint(rules if isinstance(rules, dict) else None, model=a.model)
        prov = "ollama" if ep["local"] else ep["provider"]
        if prov == "anthropic" and not ep["model"].startswith("claude"):   # --model gemini-... with no [ai] table:
            prov = next((p for _, m, p in SHEET if m == ep["model"]), "openai-compatible")   # not priced as Claude
        rows = [dict(jobs(ep["model"], prov, rules, a.share), model=ep["model"])]
    if a.md:
        open(a.md, "w", encoding="utf-8").write(doc(rows, a.share))
    if a.csv:
        write_csv(rows, a.csv)
    if a.xlsx:
        write_xlsx(rows, a.xlsx, a.share)
    if a.json:
        print(json.dumps(rows, indent=1))
    elif not (a.md or a.csv or a.xlsx):
        if len(rows) == 1:
            r = rows[0]
            if r["usd_in"] is None:
                print(f"no price known for {r['id']}: set [ai.profiles.\"{r['id']}\"] price_in / price_out in ecadlib.toml")
                return 1
            print(f"{r['id']}: ${r['usd_in']:g} in / ${r['usd_out']:g} out per million tokens (estimates)")
            for c in COLS[5:]:
                print(f"  {HEAD[c]:<26} {money(r[c])}")
        else:
            print(markdown(rows, a.share)[0])
        print("Estimates from request sizes, not a bill. `ecadlib datasheet estimate BOM` counts a real BOM.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
