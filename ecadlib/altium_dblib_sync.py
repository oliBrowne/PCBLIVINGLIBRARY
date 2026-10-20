"""Altium without 365: sync a librarian / altium365 plan.json into a Database Library (.DbLib) or an
SVN Database Library (.SVNDbLib) folder.

    python -m ecadlib.altium_dblib_sync out/librarian/plan.json --to D:/libs/ecadlib [--svn] [--prune]
        [--name ecadlib] [--win-dir C:\\libs\\ecadlib] [--svn-url URL] [--commit [-m MSG]] [--dry-run] [-o report.json]
    ecadlib dblib-sync ...                                              (same options)

Folder layout (one entry per model file, the altium_lib / altium365 convention):
    <to>/<name>.DbLib | <name>.SVNDbLib     INI file Altium opens (Excel connection, field maps, tables)
    <to>/<file>.xlsx                        parts table: one sheet per category (Capacitors, ICs, ...), all text
    <to>/<file>.csv                         the same rows flat (Table column first), for diffs and review only
                                            (<file> = <name> in ASCII, see data_file_name: Bibliothèque -> Bibliotheque)
    <to>/SchLib/<Library Ref>.SchLib        symbol per file;  [Library Path]   = SchLib\\<Library Ref>.SchLib
    <to>/PcbLib/<Footprint Ref>.PcbLib      footprint per file; [Footprint Path] = PcbLib\\<Footprint Ref>.PcbLib
    <to>/ecadlib_sync.json                  which rows and files this tool manages (only those are --prune-able)

Per plan item, like `altium365 push`: `held` items (release_check) are left exactly as they are; items without
both models are `skipped`; a `ready` item is checked again here (release_check.publish_problems) and is `held` when
its files fail the check; otherwise the row is `added`, `updated` (row fields or a model file changed) or
`up_to_date`. Row updates merge: columns added by hand in the sheet survive, plan values win, the plan's
`cleared` fields are blanked. Rows and model files no longer in the plan are reported (`not_in_plan`) and only
deleted with --prune, and only if this tool wrote them. Unchanged files are not rewritten. Model files written
are read back with altium_lib's readers and must contain the Library Ref / Footprint Ref the row names.

An existing .DbLib/.SVNDbLib is updated, not replaced: settings made in Altium's DbLib editor are kept, missing
tables and field maps are appended, and ConnectionString / SVNRepositoryURL are only rewritten when --win-dir /
--svn-url are given (or the URL is read from `svn info`). Altium opens the workbook through LibraryDatabasePath +
DatabasePathRelative=1 (the .xlsx next to the DbLib), so the library folder can be moved or copied to another PC. The
connection string's Data Source is only what Altium's DbLib editor shows (and uses if the relative option is switched
off): the --win-dir folder, else the fixed C:\\ecadlib\\library placeholder, never this machine's absolute path.

SVN (--svn): <to> is meant to be an SVN working copy; the .SVNDbLib gets SVNRepositoryURL and a LibrarySearchPath
of <url>/SchLib;<url>/PcbLib. With --commit, and `svn` on PATH and <to> a working copy, it runs svn add / delete /
commit for exactly the files it touched; otherwise it prints those commands."""
import argparse, csv, io, json, os, re, shutil, subprocess, sys
from pathlib import Path, PureWindowsPath

from . import altium_lib as AL
from . import altium_native as AN
from . import release_check as RC
from . import sheetsafe as SS
from .fsutil import read_ini_text                     # (text, encoding) of an INI Altium saved
from .altium_export import (COLUMNS, dblib_links, excel_connection, fieldmap_section, table_section)

STATE = "ecadlib_sync.json"
EXT = {"dblib": ".DbLib", "svndblib": ".SVNDbLib"}
KINDS = (("symbol", "SchLib", "Library"), ("footprint", "PcbLib", "Footprint"))


# ------------------------------------------------------------------ plan -> rows
def load_plan(path):
    """plan.json's items as written: model paths stay relative to the plan's folder (sync resolves them against
    `base`), so the change-control bookkeeping that may rewrite plan.json never puts this machine's paths in it."""
    return json.loads(Path(path).read_text())


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")     # what openpyxl refuses in a cell (ILLEGAL_CHARACTERS_RE)
CELL_MAX = 32767                                          # Excel's cell limit: openpyxl cuts longer text on save
LINK_COLUMNS = ("Part Number", "Library Ref", "Library Path", "Footprint Ref", "Footprint Path")   # always ecadlib's


def cell_text(v):
    """A value as the workbook can hold it: control characters (a \\x0b copied from a datasheet) become spaces, and
    it is cut to Excel's 32767 characters, so the .xlsx, the .csv and the next run's compare all see the same text."""
    return _CTRL.sub(" ", str(v))[:CELL_MAX]


def column_name(name):
    """A parameter name as a column and DbLib field map: no '|' or line breaks (they split the INI's Options record),
    no control characters. '' when nothing is left."""
    return _CTRL.sub(" ", AL.clean(name)).strip()


def sheet_name(name):
    """A sheet name Excel accepts that Excel's ACE provider exposes unchanged (plus '$'), so the DbLib TableName
    matches it. ACE drops an apostrophe ("Driver's ICs" -> "Drivers ICs$") and turns a '.' into '#'
    ("Res. Networks" -> "Res# Networks$"): the apostrophe is dropped here as well, and '.' becomes '_'."""
    t = re.sub(r"[\[\]:*?/\\.]", "_", str(name).strip()).replace("'", "")
    t = (t or "Misc")[:31]
    return t + "_" if t.lower() == "history" else t      # Excel reserves the sheet name History


_FOLD = str.maketrans({"ß": "ss", "æ": "ae", "Æ": "AE", "œ": "oe", "Œ": "OE", "ø": "o", "Ø": "O", "ł": "l",
                       "Ł": "L", "đ": "d", "Đ": "D", "þ": "th", "Þ": "Th"})      # letters NFKD doesn't split


def data_file_name(name):
    """Base name of the .xlsx / .csv: the package name folded to ASCII ('Bibliothèque' -> 'Bibliotheque'). Altium
    installs 'Bibliothèque.DbLib' fine, but opening a workbook with a non-ASCII file name through the Excel (ACE)
    connection fails for every table (#82, AD26)."""
    import unicodedata
    s = str(name).translate(_FOLD)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r'[^\x20-\x7e]|[<>:"/\\|?*]', "_", s).strip(" .")
    return s if s.strip("_") else "ecadlib"


def table_of(item):
    """Sheet name from the workspace folder (Components/Capacitors -> Capacitors); see sheet_name."""
    return sheet_name((item.get("folder") or "Misc").rstrip("/").rsplit("/", 1)[-1])


def sheet_names(items, tables=()):
    """{id(item): sheet} for the items whose folder would share a sheet with another folder only because sheet names
    are cut to Excel's 31 characters: the first folder (in name order) keeps the sheet, the others end in ~2, ~3
    ('Optoelectronics and Light Em~2'). Folders differing only in case share a sheet on purpose (Excel ignores
    case). tables: the workbook's sheets (a new name never takes one of them)."""
    raw = lambda it: (it.get("folder") or "Misc").rstrip("/").rsplit("/", 1)[-1].strip() or "Misc"
    full = lambda it: re.sub(r"[\[\]:*?/\\.]", "_", raw(it)).replace("'", "").lower()
    groups, first = {}, {}
    for it in items:
        groups.setdefault(table_key(sheet_name(raw(it))), set()).add(full(it))
        first.setdefault(full(it), raw(it))
    taken = {table_key(t) for t in tables} | set(groups)
    alias = {}
    for names in groups.values():
        n = 2
        for f in sorted(names)[1:]:
            while table_key(cand := f"{sheet_name(first[f])[:28]}~{n}") in taken:
                n += 1
            taken.add(table_key(cand))
            alias[f] = cand
            n += 1
    return {id(it): alias[full(it)] for it in items if full(it) in alias}


def table_key(name):
    """One table under every spelling it has had: a sheet name, a DbLib TableName ('$' at the end), the name ACE
    gave it ('#' for '.', no apostrophe) or an older ecadlib wrote ('.', apostrophe), in any case."""
    name = str(name)
    return sheet_name(name[:-1] if name.endswith("$") else name).replace("#", "_").lower()


def _steps(to, rel):
    """The STEP files (paths relative to the library folder) the recipe of library <to>/<rel> names: the
    ecadlib_3d/<stem>_<k>/ files altium_native writes beside a PcbLib."""
    d = rel.split("/")[0]
    return {f"{d}/{s}" for s in AN.recipe_steps(Path(to) / rel)
            if s.startswith("ecadlib_3d/") and ".." not in s.split("/")}


def plan_row(item, rules=None):
    """The DbLib row for one plan item, columns in altium_export order, then any extra plan parameters; parameter
    columns in the rules' house names ([parameters] names)."""
    row = {c: "" for c in COLUMNS}
    row.update({"Description": cell_text(item.get("description", "")), "Comment": cell_text(item.get("comment", ""))})
    for p in item.get("parameters", []):            # a parameter named like a key / link column never replaces it
        n = column_name(p["name"])
        if n and n not in LINK_COLUMNS:
            row[n] = cell_text(p["value"])
    row["Part Number"] = cell_text(item["name"])
    for k, d, w in KINDS:
        ref = (item.get(k) or {}).get("ref") or ""
        row[f"{w} Ref"] = AL.lib_name(ref) if ref else ""      # the name inside the file ('Ω' is stored as 'Ohm')
        row[f"{w} Path"] = f"{d}\\{AL.file_stem(ref)}.{d}" if ref else ""
    from .rules import house_row
    return house_row(row, rules)


# ------------------------------------------------------------------ parts table (xlsx + csv)
class Kept(str):
    """A cell from the workbook that is not text (a number, date, true/false or formula someone put in a column of
    their own): it compares as its text, and write_xlsx writes it back as it was (raw)."""
    def __new__(cls, text, raw):
        o = str.__new__(cls, text)
        o.raw = raw
        return o


def _cell_in(c, keep):
    v = c.value
    if v is None:
        return ""
    if not keep or (isinstance(v, str) and c.data_type != "f"):
        return str(v)
    return Kept(str(v), v)


def read_tables(xlsx, keep=False):
    """{sheet: (header, [row dicts])} from an existing parts workbook; {} if there is none. Values are text; keep=True
    (dblib-sync) reads formulas as formulas and returns non-text cells as Kept, so they are written back unchanged."""
    if not Path(xlsx).exists():
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=not keep)
    out = {}
    try:
        for ws in wb.worksheets:
            rows = [[_cell_in(c, keep) for c in r] for r in ws.iter_rows()]
            if not rows:
                continue
            cols = [(i, h) for i, h in enumerate(rows[0]) if h]       # by position: a blank header column must not
            hdr = [h for _, h in cols]                                  # shift the hand columns right of it
            out[ws.title] = (hdr, [{h: r[i] if i < len(r) else "" for i, h in cols} for r in rows[1:] if any(r)])
    finally:
        wb.close()
    return out


def write_xlsx(path, tables):
    import openpyxl
    from openpyxl.cell.cell import TYPE_STRING
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (hdr, rows) in tables.items():
        ws = wb.create_sheet(name)
        # all text: "=Value" is a parameter, not a formula; '0402' stays '0402'; the header too (a column named
        # "=HYPERLINK(..)" is text). Only a Kept cell (a number, date or formula read from the workbook) keeps its type.
        for vals in [hdr] + [[r.get(c, "") for c in hdr] for r in rows]:
            ws.append([v.raw if isinstance(v, Kept) else cell_text(v) for v in vals])
            for c, v in zip(ws[ws.max_row], vals):
                if isinstance(c.value, str) and not isinstance(v, Kept):
                    c.data_type = TYPE_STRING
        ws.freeze_panes = "B2"
    tmp = Path(path).with_name(f"{Path(path).name}.{os.getpid()}.tmp")   # its own name: two runs never share one
    wb.save(tmp)
    os.replace(tmp, path)


def csv_text(tables):
    cols = ["Table"]
    for hdr, _ in tables.values():
        cols += [c for c in hdr if c not in cols]
    buf = io.StringIO()
    w = SS.DictWriter(buf, cols, lineterminator="\n", param_refs=True)    # =Value stays; =HYPERLINK(...) does not
    w.writeheader()
    for name, (_, rows) in tables.items():
        for r in rows:
            w.writerow({"Table": name, **r})
    return buf.getvalue()


# ------------------------------------------------------------------ .DbLib / .SVNDbLib INI
def _sections(text):
    secs = []
    for ln in text.replace("\r\n", "\n").split("\n"):
        if re.fullmatch(r"\[[^\]=]+\]", ln.strip()):
            secs.append([ln.strip()[1:-1], []])
        elif secs and ln.strip():
            secs[-1][1].append(ln)
    return secs


PLACEHOLDER_DIR = r"C:\ecadlib\library"      # the Data Source folder unless --win-dir names one


def win_folder(win_dir=None):
    """The Windows folder the connection string names: --win-dir, else PLACEHOLDER_DIR on every OS. Never the library's
    own absolute path: Altium reads the workbook through LibraryDatabasePath + DatabasePathRelative=1 (next to the
    DbLib), and a machine path would stop the folder from being portable."""
    return PureWindowsPath(win_dir or PLACEHOLDER_DIR)


def _kv(lines):
    return dict(ln.split("=", 1) for ln in lines if "=" in ln)


def _same_tables(secs, names, num):
    """-> (secs, names) with each wanted table ({'Drivers ICs$': cols}) held by one [TableN] section. A section naming
    it in another spelling (table_key: "Driver's ICs$" an older ecadlib wrote, "Res# Networks$" Altium saved from
    ACE) is renamed to it, or dropped when another section already has it; its [FieldMapN] sections follow it
    (renamed, or dropped when the kept table maps that column already). A section differing only in case is kept as
    it is, and names takes its spelling. Sections are renumbered only when one was dropped."""
    want = {table_key(tn): tn for tn in names}
    groups = {}
    for s in secs:
        if num(s, "Table") is not None:
            tn = _kv(s[1]).get("TableName", "")
            if tn and table_key(tn) in want:
                groups.setdefault(table_key(tn), []).append(s)
    rename, drop, keep_name = {}, [], {}
    for k, ss in groups.items():
        tn = want[k]
        keep = next((s for s in ss if _kv(s[1])["TableName"].lower() == tn.lower()), ss[0])
        keep_name[k] = _kv(keep[1])["TableName"] if _kv(keep[1])["TableName"].lower() == tn.lower() else tn
        for s in ss:
            old = _kv(s[1])["TableName"]
            if old != keep_name[k]:
                rename[old] = keep_name[k]
            if s is not keep:
                drop.append(s)
    names = {keep_name.get(table_key(tn), tn): cols for tn, cols in names.items()}
    if not rename:
        return secs, names
    dropped = {id(s) for s in drop}
    secs = [s for s in secs if id(s) not in dropped]
    for s in secs:
        if num(s, "Table") is not None and _kv(s[1]).get("TableName") in rename:
            s[1] = [f"TableName={rename[_kv(s[1])['TableName']]}" if ln.split("=", 1)[0] == "TableName" else ln for ln in s[1]]
        elif s[0] == "DatabaseLinks" and _kv(s[1]).get("LastFocusedTable") in rename:
            s[1] = [f"LastFocusedTable={rename[_kv(s[1])['LastFocusedTable']]}" if ln.split("=", 1)[0] == "LastFocusedTable"
                    else ln for ln in s[1]]
    opts = lambda s: dict(p.split("=", 1) for p in _kv(s[1]).get("Options", "").split("|") if "=" in p)
    fm = [s for s in secs if num(s, "FieldMap") is not None]
    have = {opts(s).get("FieldName", "").lower() for s in fm if opts(s).get("TableNameOnly") not in rename}
    gone = set()
    for s in fm:
        o = opts(s)
        t = o.get("TableNameOnly")
        if t not in rename:
            continue
        o["TableNameOnly"], o["FieldName"] = rename[t], f"{rename[t]}.{o.get('FieldNameOnly', '')}"
        if o["FieldName"].lower() in have:
            gone.add(id(s))
            continue
        have.add(o["FieldName"].lower())
        s[1] = [("Options=" + "|".join(f"{a}={b}" for a, b in o.items())) if ln.split("=", 1)[0] == "Options" else ln
                for ln in s[1]]
    secs = [s for s in secs if id(s) not in gone]
    if drop or gone:                          # [Table1].. and [FieldMap1].. stay numbered without gaps
        for p in ("Table", "FieldMap"):
            for i, s in enumerate((s for s in secs if num(s, p) is not None), 1):
                s[0] = f"{p}{i}"
    return secs, names


def dblib_update(existing, tables, links, override=()):
    """New or merged INI text. tables: {sheet: columns} (sheet names without '$'); links: kwargs for
    altium_export.dblib_links; override: [DatabaseLinks] keys whose template value replaces the existing one."""
    names = {t + "$": cols for t, cols in tables.items()}
    tmpl = dblib_links(last_table=next(iter(names), ""), **links)
    if not existing:
        L = list(tmpl)
        for i, tn in enumerate(names, 1):
            L += table_section(i, tn)
        fm = 0
        for tn, cols in names.items():
            for c in cols:
                fm += 1
                L += fieldmap_section(fm, tn, c)
        return "\r\n".join(L) + "\r\n"
    secs = _sections(existing)
    link_tmpl = _kv(tmpl[3:])
    db = next((s for s in secs if s[0] == "DatabaseLinks"), None)
    if db is None:
        db = ["DatabaseLinks", []]
        secs.insert(1 if secs and secs[0][0] == "OutputDatabaseLinkFile" else 0, db)
    have = _kv(db[1])
    for k, v in link_tmpl.items():
        if k not in have:
            db[1].append(f"{k}={v}")
        elif k in override and have[k] != v:
            db[1] = [f"{k}={v}" if ln.split("=", 1)[0] == k else ln for ln in db[1]]
    num = lambda s, p: int(s[0][len(p):]) if s[0].startswith(p) and s[0][len(p):].isdigit() else None
    secs, names = _same_tables(secs, names, num)
    tabs = {_kv(s[1]).get("TableName") for s in secs if num(s, "Table") is not None}
    fmaps = {_kv(s[1]).get("Options", "").split("|", 1)[0].lower() for s in secs if num(s, "FieldMap") is not None}
    ti = max([num(s, "Table") or 0 for s in secs] + [0])
    fi = max([num(s, "FieldMap") or 0 for s in secs] + [0])
    last_t = max([k for k, s in enumerate(secs) if num(s, "Table") is not None] + [secs.index(db)])
    new_t, new_f = [], []
    for tn, cols in names.items():
        if tn not in tabs:
            ti += 1
            sec = table_section(ti, tn)
            new_t.append([sec[0][1:-1], sec[1:]])
        for c in cols:
            if f"FieldName={tn}.{c}".lower() not in fmaps:
                fi += 1
                sec = fieldmap_section(fi, tn, c)
                new_f.append([sec[0][1:-1], sec[1:]])
    secs = secs[:last_t + 1] + new_t + secs[last_t + 1:] + new_f
    return "".join(f"[{n}]\r\n" + "".join(ln + "\r\n" for ln in body) for n, body in secs)


def check_existing_dblib(lib, fname, name):
    """An existing .DbLib / .SVNDbLib at the target must be ecadlib's kind (an Excel workbook next to it): merging
    into a company's Access / SQL Server database library would switch it to ecadlib's workbook and break it."""
    if not Path(lib).is_file():
        return
    db = next((s for s in _sections(read_ini_text(lib)[0]) if s[0] == "DatabaseLinks"), None)
    kv = _kv(db[1]) if db else {}
    typ, path = kv.get("LibraryDatabaseType", "").strip(), kv.get("LibraryDatabasePath", "").strip()
    ours = {f"{fname}.xlsx".lower(), f"{name}.xlsx".lower()}
    if typ and typ.lower() != "microsoft excel":
        raise ValueError(f"{lib} is a {typ} database library, not ecadlib's Excel one: it is left as it is. "
                         "Sync into another folder, or pass another --name")
    if path and PureWindowsPath(path).name.lower() not in ours:
        raise ValueError(f"{lib} reads the workbook {path}, not ecadlib's {fname}.xlsx: it is left as it is. "
                         "Sync into another folder, or pass another --name")


# ------------------------------------------------------------------ svn
def run_cmd(cmd, cwd=None):
    return subprocess.run(cmd, cwd=cwd, capture_output=True, text=True)


def svn_info(target, runner=run_cmd, which=shutil.which):
    """Working-copy URL of target, or None (no svn on PATH, or not a working copy)."""
    if not which("svn"):
        return None
    r = runner(["svn", "info", "--show-item", "url", str(target)])
    return r.stdout.strip() if r.returncode == 0 and r.stdout.strip() else None


def svn_commands(written, removed, message):
    """svn add/delete/commit for exactly the paths (relative to the working copy root) this run touched."""
    written, removed = sorted(set(written)), sorted(set(removed))
    dirs = sorted({p.split("/")[0] for p in written + removed if "/" in p})
    cmds = []
    if written:
        cmds.append(["svn", "add", "--force", "--parents", *written])
    if removed:
        cmds.append(["svn", "delete", "--force", *removed])
    if written or removed:
        cmds.append(["svn", "commit", "--depth", "empty", "-m", message, *dirs, *written, *removed])
    return cmds


def svn_step(target, written, removed, commit=False, message="ecadlib dblib-sync", runner=run_cmd, which=shutil.which):
    """Run (commit=True, svn available, target is a working copy) or just return the svn commands."""
    cmds = svn_commands(written, removed, message)
    res = {"commands": [" ".join(_q(c) for c in cmd) for cmd in cmds], "ran": False, "output": []}
    if not cmds:
        res["note"] = "nothing to commit"
    elif not commit:
        res["note"] = "run these in the working copy (or pass --commit)"
    elif not which("svn"):
        res["note"] = "svn is not on PATH: run these by hand"
    elif runner(["svn", "info", str(target)]).returncode != 0:
        res["note"] = f"{target} is not an svn working copy: check it out first, then run these"
    else:
        for cmd in cmds:
            r = runner(cmd, cwd=str(target))
            res["output"].append((r.stdout or "") + (r.stderr or ""))
            if r.returncode != 0:
                res["error"] = f"{cmd[1]} failed (exit {r.returncode})"
                break
        else:
            res["ran"] = True
    return res


def _q(s):
    return f'"{s}"' if re.search(r"[\s()&|;<>]", s) else s


# ------------------------------------------------------------------ sync
def _readback(path, kind, ref):
    """None if the written model reads back with altium_lib and holds `ref`, else the problem."""
    fname = re.sub(r"(\.\d+)?\.tmp$", "", Path(path).name)       # the file it becomes, not the temporary copy
    try:
        names = ([c["name"] for c in AL.read_schlib(path)[1]] if kind == "SchLib"
                 else [f["name"] for f in AL.read_pcblib(path)])
    except Exception as e:
        return f"{fname} does not read back: {type(e).__name__}: {str(e).replace(Path(path).name, fname)}"
    want = AL.lib_name(ref)                       # as the file stores it ('Ω' -> 'Ohm')
    return None if want in names else f"{fname} holds {names}, not {want!r}"


def _copy_recipe(to, src, dst):
    """AN.copy_recipe, -> the files it wrote (recipe, STEP files, .ecadstate) relative to `to`, for svn."""
    st = dst.with_name(dst.name + AN.STATE_EXT)
    before = st.read_bytes() if st.is_file() else None
    paths = [Path(p) for p in AN.copy_recipe(src, dst)]
    if st.is_file() and st.read_bytes() != before:
        paths.append(st)
    return [p.relative_to(to).as_posix() for p in paths]


def sync(plan, target_dir, mode="dblib", name="ecadlib", prune=False, win_dir=None, svn_url=None,
         dry_run=False, base=None, runner=run_cmd, which=shutil.which, rules=None, overwrite_edits=False):
    """Bring <target_dir> in line with the plan. plan: a plan.json path, or its items (model paths absolute or
    relative to `base`). Returns {"results": [...], "written": [...], "removed": [...], "dblib": path, ...};
    results are {"name", "action", "detail", "table"} with action added/updated/up_to_date/held/skipped/error,
    and not_in_plan/removed for rows (kind "row") and model files (kind "file") the plan no longer has; a pruned
    footprint's STEP files (ecadlib_3d/..., written by ecadlib) are results with kind "3d".
    A model file changed in Altium since ecadlib's copy was built there is left as it is (a note says so) unless
    overwrite_edits (ecadlib.altium_native, <file>.ecadstate). A real sync holds the library lock throughout
    (change control: the model files copied are the settled ones)."""
    if mode not in EXT:
        raise ValueError(f"mode must be one of {sorted(EXT)}")
    lib_dir = Path(base or (Path(plan).parent if isinstance(plan, (str, Path)) else "."))
    if dry_run:
        return _sync(plan, target_dir, mode, name, prune, win_dir, svn_url, dry_run, base, runner, which, rules,
                     overwrite_edits)
    from . import change_gate as CG
    with CG.publishing(lib_dir):
        return _sync(plan, target_dir, mode, name, prune, win_dir, svn_url, dry_run, base, runner, which, rules,
                     overwrite_edits)


def _sync(plan, target_dir, mode, name, prune, win_dir, svn_url, dry_run, base, runner, which, rules, overwrite_edits):
    items = load_plan(plan) if isinstance(plan, (str, Path)) else plan
    base = Path(base or (Path(plan).parent if isinstance(plan, (str, Path)) else "."))
    to = Path(target_dir)
    from .rules import param_name, resolve
    rules = resolve(rules, base)                            # house parameter names: ecadlib.toml found from the plan
    if not name or name in (".", "..") or Path(name).name != name or PureWindowsPath(name).name != name:
        raise ValueError(f"--name {name!r} is a file name, not a path: use --to for the folder")
    check_existing_dblib(to / f"{name}{EXT[mode]}", data_file_name(name), name)
    from . import change_gate as CG                         # change control (#105): hand edits settle first; a held
    if not dry_run:                                         # change syncs the released revision
        CG.publish_guard(plan if isinstance(plan, (str, Path)) else base, items, mode, rules=rules, base=base)
    fname = data_file_name(name)
    xlsx, csvf, lib, statef = to / f"{fname}.xlsx", to / f"{fname}.csv", to / f"{name}{EXT[mode]}", to / STATE
    state = json.loads(statef.read_text()) if statef.exists() else {}
    managed_rows, managed_files = set(state.get("rows", [])), set(state.get("files", []))
    # an older ecadlib named the workbook after the package ('Bibliothèque.xlsx'), which Altium can't open: move to
    # the ASCII name, keeping the rows (and columns added by hand) of the old one
    legacy = [f for f in (to / f"{name}.xlsx", to / f"{name}.csv") if fname != name and f.exists()]
    old_xlsx = to / f"{name}.xlsx"
    from_legacy = old_xlsx in legacy and not xlsx.exists()
    old = read_tables(old_xlsx if from_legacy else xlsx, keep=True)
    # a sheet without a Part Number column (notes, a change log) is not a component table: written back as it is
    others = {t: v for t, v in old.items() if "Part Number" not in v[0]}
    old = {t: v for t, v in old.items() if t not in others}
    tables = {}
    for t, (h, rows) in old.items():                    # a sheet an older ecadlib named "Driver's ICs": ACE's name
        t = next((k for k in tables if table_key(k) == table_key(t)), sheet_name(t))
        hdr, rs = tables.setdefault(t, ([], []))
        hdr += [c for c in h if c not in hdr]
        rs += [dict(r) for r in rows]
    where = {}
    for t, (_, rows) in tables.items():                 # the first row of a Part Number is its row (EXP-79: the rest
        for r in rows:                                  # are reported as duplicates below)
            where.setdefault(r.get("Part Number"), t)
    results, written, removed, copied, failed, prune_steps = [], [], [], {}, {}, set()
    folded = {}                                         # rel.casefold() -> rel: Windows shares ignore case
    sheets = sheet_names(items, tables)
    seen = {}

    def find(pn, t):
        rows = tables.get(t, ([], []))[1]
        return next((r for r in rows if r.get("Part Number") == pn), None)

    for it in items:
        t0 = sheets.get(id(it)) or table_of(it)             # Excel sheet names ignore case: Passives = passives
        res = {"name": it["name"], "action": "", "detail": "", "table": next((k for k in tables if table_key(k) == table_key(t0)), t0)}
        results.append(res)
        if it.get("status") == "held":
            res["action"], res["detail"] = "held", "; ".join(it.get("hold_reasons") or ["release check"])
            if it["name"] in where:
                res["table"] = where[it["name"]]
            continue
        if it.get("status") != "ready":
            res["action"], res["detail"] = "skipped", it.get("status") or "not ready"
            continue
        missing = [k for k, _, _ in KINDS if not ((it.get(k) or {}).get("file") and (it.get(k) or {}).get("ref"))]
        if missing:                                         # "items without both models are skipped"
            res["action"], res["detail"] = "skipped", "no " + " / ".join(missing) + " model"
            continue
        why = RC.publish_problems(it, base, rules)          # the release check again, on the files copied now
        if why:
            res["action"], res["detail"] = "held", "; ".join(why)
            if it["name"] in where:
                res["table"] = where[it["name"]]
            continue
        if _CTRL.search(it["name"]):
            res["action"], res["detail"] = "error", "the part name has a control character (a line break or tab copied " \
                                                    "from somewhere): rename it"
            continue
        if it["name"].casefold() in seen:                   # Altium Part Numbers ignore case
            res["action"] = "error"
            res["detail"] = f"duplicate part name in the plan (also {seen[it['name'].casefold()]!r}): rename one of them"
            continue
        seen[it["name"].casefold()] = it["name"]
        notes, problems = [], []
        for k, d, w in KINDS:                               # model files first: a bad file must not reach the table
            src, ref = it[k]["file"], it[k]["ref"]
            src = Path(src) if Path(src).is_absolute() else base / src
            rel = f"{d}/{AL.file_stem(ref)}.{d}"
            dst = to / rel
            if folded.setdefault(rel.casefold(), rel) != rel:
                problems.append(f"{rel}: same file name as {folded[rel.casefold()]} (the names differ only in case, "
                                "and Windows file names ignore case): rename one of them")
                continue
            if not src.exists():
                problems.append(f"{k} file missing: {src}")
                continue
            data = src.read_bytes()
            if rel in copied:
                if copied[rel] != data:
                    problems.append(f"{rel}: another plan item already wrote different content this run")
                elif rel in failed:                         # same bad file as an earlier item: same error
                    problems.append(failed[rel])
                elif rel in written:
                    notes.append(f"{rel} updated")
                continue
            copied[rel] = data
            n_notes = len(notes)
            try:
                if dst.exists() and dst.read_bytes() == data:
                    if not dry_run:                         # ecadlib's bytes (an earlier run that failed later on
                        written += _copy_recipe(to, src, dst)     # copied them): ecadlib's file
                        managed_files.update(_steps(to, rel))
                        managed_files.add(rel)
                    continue
                act = AN.replace_action(data, dst, overwrite_edits, save=not dry_run)
                if act == "keep":                               # Altium's build of these same bytes: leave it
                    if not dry_run:
                        managed_files.update(_steps(to, rel))
                        managed_files.add(rel)
                    continue
                if act == "edited":
                    notes.append(f"{rel} {AN.EDITED_NOTE}")
                    continue
                if (dst.exists() and rel not in managed_files and AN.state_of(dst) is None
                        and not overwrite_edits):               # a library of the company's own, not ecadlib's
                    problems.append(f"{rel} is already there and was not written by ecadlib: not replaced "
                                    "(rename the part's model, or --overwrite-edits replaces it)")
                    failed[rel] = problems[-1]
                    continue
                notes.append(f"{rel} {'updated' if dst.exists() else 'added'}")
                if not dry_run:
                    dst.parent.mkdir(parents=True, exist_ok=True)
                    tmp = dst.with_name(f"{dst.name}.{os.getpid()}.tmp")
                    tmp.write_bytes(data)
                    err = _readback(tmp, d, ref)
                    if err:
                        tmp.unlink()
                        problems.append(err)
                        failed[rel] = err
                        notes.pop()
                        continue
                    os.replace(tmp, dst)
                    # what Altium rebuilds the file from (ecadlib.altium_native): svn needs it too
                    written += _copy_recipe(to, src, dst)
                    managed_files.update(_steps(to, rel))       # its STEP files are ecadlib's too: --prune removes them
                else:
                    err = _readback(src, d, ref)
                    if err:
                        problems.append(err)
                        failed[rel] = err
                        continue
                written.append(rel)
            except PermissionError:                         # Windows: the file is open (locked) in Altium or another
                dst.with_name(f"{dst.name}.{os.getpid()}.tmp").unlink(missing_ok=True)   # program: this part is skipped, the others sync
                err = (f"{rel} could not be read or replaced: it is open in Altium or another program. "
                       "Close it there (or wait for Altium's build to finish), then run this again")
                problems.append(err)
                failed[rel] = err
                del notes[n_notes:]
        if problems:
            res["action"], res["detail"] = "error", "; ".join(problems + notes)
            continue
        want, t = plan_row(it, rules), res["table"]
        old_t = where.get(it["name"])
        cur = find(it["name"], old_t) if old_t else None
        if cur is None:
            hdr, rows = tables.setdefault(t, (list(want), []))
            hdr += [c for c in want if c not in hdr]
            rows.append(want)
            where[it["name"]] = t
            res["action"], res["detail"] = "added", "; ".join(notes)
            managed_rows.add(it["name"])
            continue
        new = dict(cur)
        for c in (param_name(c, rules) for c in it.get("cleared", [])):
            if c in new:
                new[c] = ""
        new.update(want)
        changed = [c for c in new if str(new.get(c, "")) != str(cur.get(c, ""))]
        if old_t != t:                                      # category changed: move the row
            tables[old_t][1].remove(cur)
            hdr, rows = tables.setdefault(t, (list(want), []))
            hdr += [c for c in new if c not in hdr]
            rows.append(new)
            where[it["name"]] = t
            notes.insert(0, f"moved from {old_t}")
        elif changed:
            hdr, rows = tables[t]
            hdr += [c for c in new if c not in hdr]
            rows[rows.index(cur)] = new
        if changed:
            notes.insert(0, "row: " + ", ".join(changed))
        res["action"] = "updated" if notes else "up_to_date"
        res["detail"] = "; ".join(notes)
        managed_rows.add(it["name"])

    # rows and model files the plan no longer has
    plan_names = {it["name"] for it in items}
    for t, (_, rows) in tables.items():
        for r in list(rows):
            pn = r.get("Part Number", "")
            if pn in plan_names:
                cur = find(pn, where[pn]) if pn in where else None
                if cur is None or r is cur:
                    continue
                res = {"name": pn, "action": "duplicate", "table": t, "kind": "row",     # a copy-paste in Excel: Altium
                       "detail": f"a second row for this Part Number (ecadlib keeps the one in {where[pn]} up to date)"}
                if prune and pn in managed_rows:                                     # lists the part twice
                    rows.remove(r)
                    res["action"] = "removed"
                else:
                    res["detail"] += "; kept" + (" (--prune deletes it)" if pn in managed_rows else "")
                results.append(res)
                continue
            res = {"name": pn, "action": "not_in_plan", "detail": "", "table": t, "kind": "row"}
            if pn not in managed_rows:
                res["detail"] = "not written by ecadlib: kept"
            elif prune:
                rows.remove(r)
                managed_rows.discard(pn)
                res["action"] = "removed"
            else:
                res["detail"] = "kept (--prune deletes it)"
            results.append(res)
    used = {f"{d}/{AL.file_stem(r.get(f'{w} Ref', ''))}.{d}" for _, (_, rows) in tables.items() for r in rows
            for _, d, w in KINDS if r.get(f"{w} Ref")}
    for _, d, _ in KINDS:
        for f in sorted((to / d).glob(f"*.{d}")) if (to / d).is_dir() else []:
            rel = f"{d}/{f.name}"
            if rel in used or rel in copied:
                continue
            res = {"name": rel, "action": "not_in_plan", "detail": "", "table": "", "kind": "file"}
            if rel not in managed_files:
                res["detail"] = "not written by ecadlib: kept"
            elif prune:
                res["action"] = "removed"
                removed.append(rel)
                managed_files.discard(rel)
                steps = _steps(to, rel)             # read from its recipe before that is deleted
                if not dry_run:
                    f.unlink()
                    recipe = f.with_name(f.name + AN.EXT)      # else Altium's rebuild would make a file put here later
                    if recipe.is_file():                       # from the recipe of the one pruned
                        recipe.unlink()
                prune_steps.update(steps)
            else:
                res["detail"] = "no row uses it; kept (--prune deletes it)"
            results.append(res)
    # the pruned footprints' STEP folders (ecadlib_3d/<stem>_<k>/): only files this tool wrote and no kept library
    # names; a folder is removed only once empty, so anything the user put there stays (and keeps the folder)
    keep = {s for _, d, _ in KINDS if d == "PcbLib" and (to / d).is_dir()
            for f in (to / d).glob(f"*.{d}") if f"{d}/{f.name}" not in removed for s in _steps(to, f"{d}/{f.name}")}
    for rel in sorted(prune_steps - keep):
        if rel not in managed_files:
            continue
        managed_files.discard(rel)
        if (to / rel).is_file():                # not in `removed`: svn never added them (svn delete would fail)
            results.append({"name": rel, "action": "removed", "detail": "STEP model of a pruned footprint", "table": "",
                            "kind": "3d"})
            if not dry_run:
                (to / rel).unlink()
                for folder in ((to / rel).parent, (to / rel).parent.parent):
                    if folder.name != "PcbLib" and folder.is_dir() and not any(folder.iterdir()):
                        folder.rmdir()
    managed_files |= {w for w in written if not w.endswith((AN.EXT, AN.STATE_EXT)) and "/ecadlib_3d/" not in w}
    had = {table_key(t) for t in old}
    tables = {t: v for t, v in tables.items() if v[1] or table_key(t) in had}  # never create an empty sheet

    # parts table, DbLib file, state: only rewritten when their content changes
    out = {"results": results, "written": written, "removed": removed, "dblib": str(lib), "mode": mode,
           "dry_run": dry_run}
    def put(rel, text=None, writer=None, data=None):
        if not dry_run:
            to.mkdir(parents=True, exist_ok=True)
            if writer:
                writer()
            else:
                (to / rel).write_bytes(data if data is not None else text.encode("utf-8"))
        written.append(rel)
    if tables and (tables != old or not xlsx.exists()):
        put(xlsx.name, writer=lambda: write_xlsx(xlsx, {**others, **tables}))   # other sheets: as they were
    if tables:
        ct = csv_text(tables)
        if not csvf.exists() or csvf.read_bytes().decode("utf-8", "replace") != ct:
            put(csvf.name, ct)

    url = svn_url or (svn_info(to, runner, which) if mode == "svndblib" else None)
    win = win_folder(win_dir)
    extra = (("SVNRepositoryURL", url or ""), ("FileRepositoryIsRelative", "0")) if mode == "svndblib" else ()
    links = dict(conn_str=excel_connection(win / xlsx.name), db_type="Microsoft Excel", db_path=xlsx.name,
                 relative=True, search_path=f"{url}/SchLib;{url}/PcbLib" if url else "", extra=extra)
    override = {"LibraryDatabaseType", "LibraryDatabasePath", "DatabasePathRelative"}
    if win_dir or legacy:                           # else keep a connection set in Altium's DbLib editor
        override.add("ConnectionString")
    if url:
        override |= {"SVNRepositoryURL", "LibrarySearchPath"}
    # keep the \r\n, and the encoding a DbLib saved by Altium's editor has (ANSI, or UTF-8 / UTF-16 with a BOM)
    cur_text, db_enc = read_ini_text(lib) if lib.exists() else ("", "utf-8")
    if tables:
        text = dblib_update(cur_text, {t: h for t, (h, _) in tables.items()}, links, override)
        if text != cur_text:
            try:
                data = text.encode(db_enc)
            except UnicodeEncodeError:                # a new table / column name the ANSI code page can't hold
                data = text.encode("utf-8-sig")
            put(lib.name, data=data)
        cur_text = text
    db = next((s for s in _sections(cur_text) if s[0] == "DatabaseLinks"), ["", []])
    url = url or _kv(db[1]).get("SVNRepositoryURL") or None      # the one already in the file
    if mode == "svndblib" and not url:
        out["warning"] = "no SVN repository URL (pass --svn-url, or sync into an svn working copy)"
    if tables and legacy:                           # the ASCII-named workbook replaces them
        for f in legacy:
            if f.suffix == ".xlsx" and not from_legacy:  # not what this run read: someone's edits may be in it
                results.append({"name": f.name, "action": "not_in_plan", "table": "", "kind": "file",
                                "detail": f"an older ecadlib's workbook, not merged ({fname}.xlsx is the one Altium "
                                          "reads): kept; move what you need from it, then delete it"})
                continue
            removed.append(f.name)
            results.append({"name": f.name, "action": "removed", "detail": f"renamed to {fname}{f.suffix}",
                            "table": "", "kind": "file"})
            if not dry_run:
                f.unlink()
    st = {"rows": sorted(managed_rows), "files": sorted(managed_files)}
    if st != {"rows": sorted(state.get("rows", [])), "files": sorted(state.get("files", []))}:
        put(STATE, json.dumps(st, indent=1) + "\n")
    out["svn_url"] = url
    if not dry_run:
        done = {r["name"] for r in results if r.get("action") in ("added", "updated", "up_to_date")}
        CG.record_publish(base, items, f"{mode}:{CG.project_key(base, to)}", ok=lambda it: it["name"] in done)
    return out


def summary(out):
    n = {}
    for r in out["results"]:
        n[r["action"]] = n.get(r["action"], 0) + 1
    return ", ".join(f"{v} {k}" for k, v in sorted(n.items()))


def main(argv=None):
    ap = argparse.ArgumentParser(description="Sync a plan.json into an Altium .DbLib / .SVNDbLib library folder")
    ap.add_argument("plan", help="plan.json from `ecadlib librarian` or `ecadlib altium365 plan`")
    ap.add_argument("--to", required=True, help="library folder (for --svn: an svn working copy)")
    ap.add_argument("--svn", action="store_true", help="write an .SVNDbLib (SVN Database Library) instead of a .DbLib")
    ap.add_argument("--name", default="ecadlib", help="base name of the .DbLib/.SVNDbLib, and (in ASCII) of the .xlsx and .csv")
    ap.add_argument("--prune", action="store_true", help="delete rows and model files (written by ecadlib) no longer in the plan")
    ap.add_argument("--win-dir", help=r"Windows folder the library lives in, for the Excel connection string (default C:\ecadlib\library; Altium finds the workbook relative to the DbLib either way)")
    ap.add_argument("--svn-url", help="repository URL of --to (default: from `svn info`)")
    ap.add_argument("--commit", action="store_true", help="svn add/commit the changed files (needs svn on PATH and a working copy)")
    ap.add_argument("-m", "--message", default="ecadlib dblib-sync", help="svn commit message")
    ap.add_argument("--dry-run", action="store_true", help="report only; write nothing")
    ap.add_argument("--overwrite-edits", action="store_true",
                    help="replace library files someone changed in Altium since ecadlib's copy was built there")
    ap.add_argument("-o", "--out", help="write the JSON report here")
    ap.add_argument("--rules", help="library rules file: house parameter names (default: ecadlib.toml found from the plan)")
    a = ap.parse_args(argv)

    try:
        out = sync(a.plan, a.to, "svndblib" if a.svn else "dblib", a.name, a.prune, a.win_dir, a.svn_url, a.dry_run,
                   rules=a.rules, overwrite_edits=a.overwrite_edits)
    except ValueError as e:                        # a wrong --name, a library folder that isn't ecadlib's kind
        print(f"ecadlib dblib-sync: {e}", file=sys.stderr)
        sys.exit(2)
    from . import messages as M, term as TM
    t = TM.Term()
    rows = []
    for r in out["results"]:
        st = {"added": "ready", "updated": "ready", "unchanged": "skipped", "held": "held", "error": "failed",
              "removed": "info"}.get(r["action"], "info")
        detail = M.explain(r["detail"].removeprefix("held: "))["text"] if r["action"] == "held" else M.clean(r["detail"])
        if r["action"] in ("added", "updated", "unchanged") and not t.verbose:
            detail = ""
        rows.append([t.status(st, r["action"]), r["table"], r["name"], detail])
    t.table(rows)
    if out.get("warning"):
        t.note(M.clean(out["warning"]), None)
    if a.dry_run and not t.quiet:
        t.wrap("would write " + (", ".join(out["written"]) or "nothing"), 0)
    if a.svn and not a.dry_run:
        s = svn_step(a.to, out["written"], out["removed"], a.commit, a.message)
        out["svn"] = s
        t.write(f"svn: {s.get('note') or ('committed' if s['ran'] else s.get('error', ''))}")
        if not s["ran"]:
            for c in s["commands"]:
                t.write(f"  {c}")
    written = [] if a.dry_run else [(out["dblib"], "open this in Altium (it lists the parts from the .xlsx next to it)")]
    if not a.dry_run and t.verbose:
        written += [(w, "") for w in out["written"]]
    t.summary(f"{summary(out)}" + (" (dry run: nothing written; " + M.plural(len(out["written"]), "file") + " would be)" if a.dry_run else "")
              + (f"; removed {', '.join(out['removed'])}" if out["removed"] else ""),
              written=written,
              next=[(f"ecadlib dblib-sync {a.plan} --to {a.to}", "write it for real")] if a.dry_run else
              ([(f"ecadlib review {a.plan}", "the held parts")] if any(r["action"] == "held" for r in out["results"]) else []))
    if a.out:
        Path(a.out).parent.mkdir(parents=True, exist_ok=True)
        Path(a.out).write_text(json.dumps(out, indent=2))
    errors = sum(r["action"] == "error" for r in out["results"])
    if errors or (out.get("svn") or {}).get("error"):
        sys.exit(f"{errors} part(s) failed" + (f"; svn: {out['svn']['error']}" if (out.get("svn") or {}).get("error") else ""))


if __name__ == "__main__":
    main()
